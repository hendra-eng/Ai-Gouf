"""
scripts/migrasi_coa_v1_ke_v2.py
===============================
Migrasi COA klien yang dibuat dari template industri LAMA (V1,
dataset/COA/COA_Industry.xlsx, kode pola 1-10001 ...) ke struktur template
V2 (dataset/COA/COA_Industry_V2.xlsx, sudah dimuat ke DB lewat migration 30:
management_coa_industry_template_accounts, kode 110001 / 120001 ...).

Keputusan bisnis 2026-10-24:
    - Yang dimigrasi HANYA klien template lama. Klien dengan COA asli (mis.
      SAU / NPI, hasil seed COA_Clients_GOUF) dilewati otomatis.
    - Kode yang tampil = kode V2 (kode V1 tidak disimpan sebagai Source
      Account Code). Source Account Code yang memang diisi user (kode sistem
      lama klien dari Import Legacy COA, beda dari kode V1) dipertahankan.
    - Akun V2 yang belum dimiliki klien ditambahkan (COA lengkap sesuai template).

Cara kerja per klien (1 transaksi, semua atau tidak sama sekali):
    1. Deteksi: >= AMBANG_V1 akun aktif klien punya kode yang ada di salah satu
       sheet template V1 -> klien template lama; template V2 = KBLI yang sama.
    2. Pemetaan per standard_account_code: akun lama (urut kode) dipasangkan
       berurutan ke akun V2 dengan standard code sama (urut template), mis.
       Bank Account 1..4 -> 120001..120004. Akun lama yang tidak kebagian
       pasangan (lebih banyak dari V2 / standard code tidak ada di V2) dapat
       kode baru = kode terbesar di blok V2 yang sama + 1 (blok: standard code
       sama -> account sub sama -> klasifikasi sama).
    3. acc_no diubah DI TEMPAT (id baris tetap -> semua relasi coa_id: saldo
       awal, Account Mapping, FS mapping, kategori aset/persediaan, dst. aman),
       2 tahap lewat kode sementara supaya tidak bentrok unique index.
       Akun yang dapat pasangan V2 ikut memakai head/sub/deskripsi/IFRS/saldo
       normal V2; namanya diganti nama V2 HANYA kalau masih sama dengan nama
       template V1 (nama yang sudah diubah user dipertahankan).
    4. Kode akun (teks) di tabel jurnal/transaksi klien ikut diganti
       (TABEL_KODE). Trigger kunci periode & kunci Account Mapping dimatikan
       SEMENTARA di dalam transaksi (ALTER TABLE ... DISABLE TRIGGER --
       transaksional, otomatis aktif lagi kalau rollback).
    5. Akun V2 yang belum ada ditambahkan.

Default DRY-RUN: tidak mengubah apa pun, cuma laporan + CSV pemetaan di
hasil_output/. Tambah --apply untuk menjalankan. BACKUP DATABASE DULU
(pg_dump) sebelum --apply di VPS -- ALTER TABLE mengunci tabel sebentar,
jalankan saat aplikasi tidak dipakai.

Cara pakai:
    cd backend
    venv\\Scripts\\python scripts\\migrasi_coa_v1_ke_v2.py                      # dry-run semua klien
    venv\\Scripts\\python scripts\\migrasi_coa_v1_ke_v2.py --client "PT ABC"     # dry-run 1 klien (nama / id)
    venv\\Scripts\\python scripts\\migrasi_coa_v1_ke_v2.py --client "PT ABC" --apply
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent / ".env")
except ImportError:
    pass

import openpyxl
from sqlalchemy import text

from db_client import engine

ROOT = Path(__file__).resolve().parent.parent.parent
FILE_V1 = ROOT / "dataset" / "COA" / "COA_Industry.xlsx"
FOLDER_HASIL = Path(__file__).resolve().parent.parent / "hasil_output"
SHEET_INDUSTRI = "COA_JENIS INDUSTRY"
AMBANG_V1 = 0.8   # porsi akun aktif klien yang kodenya ada di template V1

# Kolom semantik yang diambil dari akun V2 pasangannya.
KOLOM_V2 = ("account_classification", "account_head", "account_sub", "normal_balance", "description",
            "international_standard_group", "standard_account_code", "ifrs_taxonomy_reference", "ifrs_source")

# (tabel, kolom kode akun, SQL filter klien -- alias tabel = t, parameter :kid)
TABEL_KODE: List[Tuple[str, str, str]] = [
    ("asset_fixed_asset_depreciation_lines", "accum_account_code",
     "t.asset_id IN (SELECT id FROM asset_fixed_assets WHERE client_id = :kid)"),
    ("asset_fixed_asset_depreciation_lines", "expense_account_code",
     "t.asset_id IN (SELECT id FROM asset_fixed_assets WHERE client_id = :kid)"),
    ("asset_fixed_asset_journal_lines", "account_code",
     "t.journal_id IN (SELECT id FROM asset_fixed_asset_journals WHERE client_id = :kid)"),
    ("asset_inventory_journal_lines", "account_code",
     "t.journal_id IN (SELECT id FROM asset_inventory_journals WHERE client_id = :kid)"),
    ("asset_other_transaction_lines", "account_code",
     "t.journal_id IN (SELECT id FROM asset_other_transactions WHERE client_id = :kid)"),
    ("financial_transaction_bank_cash_transaction_lines", "account_code",
     "t.transaction_id IN (SELECT id FROM financial_transaction_bank_cash_transactions WHERE client_id = :kid)"),
    ("financial_transaction_bank_cash_transactions", "bank_account_code", "t.client_id = :kid"),
    ("financial_transaction_journal_entry_draft_lines", "account_code",
     "t.draft_id IN (SELECT id FROM financial_transaction_journal_entry_drafts WHERE management_client_id = :kid)"),
    ("financial_transaction_other_journal_lines", "account_code",
     "t.journal_id IN (SELECT id FROM financial_transaction_other_journals WHERE client_id = :kid)"),
    ("financial_transaction_purchase_transaction_lines", "account_code",
     "t.transaction_id IN (SELECT id FROM financial_transaction_purchase_transactions WHERE management_client_id = :kid)"),
    ("financial_transaction_purchase_transactions", "ap_account_code", "t.management_client_id = :kid"),
    ("financial_transaction_purchase_transactions", "tax_account_code", "t.management_client_id = :kid"),
    ("financial_transaction_sales_account_mappings", "piutang_account_code", "t.client_id = :kid"),
    ("financial_transaction_sales_account_mappings", "pendapatan_account_code", "t.client_id = :kid"),
    ("financial_transaction_sales_account_mappings", "ppn_account_code", "t.client_id = :kid"),
    ("financial_transaction_sales_account_mappings", "pph_account_code", "t.client_id = :kid"),
    ("planning_tax_ledger", "gl_account_code", "t.client_id = :kid"),
    ("planning_tax_payments", "bank_account_code", "t.client_id = :kid"),
]


def _teks(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _urut_kode(kode: str) -> Tuple:
    """Urutan natural kode akun: '1-10002' < '1-10010', '120002' < '120010'."""
    return tuple(int(x) if x.isdigit() else x for x in re.split(r"(\d+)", kode or "") if x != "")


# ============================================================
# BACA TEMPLATE
# ============================================================

def baca_template_v1(path: Path) -> Dict[str, Dict[str, Any]]:
    """KBLI -> {"sheet", "akun": {acc_no: (nama, std)}} dari Excel template V1."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[SHEET_INDUSTRI].iter_rows(values_only=True))
    header = [_teks(h) for h in rows[0]]
    i_kbli, i_sheet = header.index("KBLI CATEGORY"), header.index("TEMPLATE SHEET")
    hasil = {}
    for r in rows[1:]:
        kbli, sheet = _teks(r[i_kbli]), _teks(r[i_sheet])
        if not kbli or sheet not in wb.sheetnames:
            continue
        arows = list(wb[sheet].iter_rows(values_only=True))
        ah = [_teks(h) for h in arows[0]]
        i_no, i_nama, i_std = ah.index("ACC NO"), ah.index("ACCOUNT NAME"), ah.index("STANDARD ACCOUNT CODE")
        hasil[kbli] = {"sheet": sheet, "akun": {
            _teks(a[i_no]): (_teks(a[i_nama]), _teks(a[i_std]).lower()) for a in arows[1:] if a and _teks(a[i_no])
        }}
    return hasil


def baca_template_v2(conn) -> Dict[str, List[Dict[str, Any]]]:
    """KBLI -> akun template V2 (urut sort_order) dari DB."""
    rows = conn.execute(text("""
        SELECT t.kbli_category, a.* FROM management_coa_industry_template_accounts a
        JOIN management_coa_industry_templates t ON t.id = a.template_id
        ORDER BY t.kbli_category, a.sort_order
    """)).mappings().all()
    hasil: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for r in rows:
        hasil[r["kbli_category"]].append(dict(r))
    return hasil


# ============================================================
# RENCANA PER KLIEN
# ============================================================

def _kode_baru(blok: List[str], dipakai: set) -> Optional[str]:
    """Kode berikut setelah kode numerik terbesar di blok yang belum dipakai."""
    angka = [int(k) for k in blok if k.isdigit()]
    if not angka:
        return None
    n = max(angka) + 1
    while str(n) in dipakai:
        n += 1
    return str(n)


def rencana_klien(conn, klien: Dict[str, Any], v1: Dict[str, Dict[str, Any]],
                  v2: Dict[str, List[Dict[str, Any]]]) -> Dict[str, Any]:
    kid = klien["id"]
    semua = [dict(r) for r in conn.execute(text(
        "SELECT * FROM management_client_coa WHERE client_id = :kid"), {"kid": kid}).mappings().all()]
    aktif = [a for a in semua if a["deleted_at"] is None]
    hasil: Dict[str, Any] = {"klien": klien, "baris": [], "tambah": [], "alasan": None, "kbli": None}
    if not aktif:
        hasil["alasan"] = "COA kosong"
        return hasil

    # 1. Deteksi template V1 terbaik.
    skor = sorted(((sum(1 for a in aktif if a["acc_no"] in t["akun"]) / len(aktif), kbli) for kbli, t in v1.items()),
                  reverse=True)
    porsi, kbli = skor[0]
    kode_v2 = {a["acc_no"] for akun in v2.values() for a in akun}
    if sum(1 for a in aktif if a["acc_no"] in kode_v2) / len(aktif) >= AMBANG_V1:
        hasil["alasan"] = "sudah memakai struktur V2"
        return hasil
    if porsi < AMBANG_V1:
        hasil["alasan"] = f"COA asli klien (hanya {porsi:.0%} kode cocok template V1) -- tidak dimigrasi"
        return hasil
    if kbli not in v2:
        hasil["alasan"] = f"template V2 untuk KBLI {kbli} tidak ada di DB (jalankan migration 30)"
        return hasil
    hasil["kbli"], hasil["porsi"] = kbli, porsi
    nama_v1 = {k: n for k, (n, _) in v1[kbli]["akun"].items()}
    target = v2[kbli]

    # 2. Pasangkan per standard code (urut kode lama -> urut template V2).
    slot: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for t in target:
        slot[(t["standard_account_code"] or "").strip().lower()].append(t)
    lama_per_std: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for a in sorted(aktif, key=lambda x: _urut_kode(x["acc_no"])):
        lama_per_std[(a["standard_account_code"] or "").strip().lower()].append(a)

    v2_dipakai, rencana, sisa = set(), [], []
    for std, daftar in lama_per_std.items():
        pasangan = slot.get(std, []) if std else []
        for i, a in enumerate(daftar):
            if i < len(pasangan):
                rencana.append((a, pasangan[i]))
                v2_dipakai.add(pasangan[i]["acc_no"])
            else:
                sisa.append(a)

    terpakai = {t["acc_no"] for t in target} | {a["acc_no"] for a in semua}
    for a in sisa:
        std = (a["standard_account_code"] or "").strip().lower()
        blok = ([t["acc_no"] for t in slot.get(std, [])] if std else []) \
            or [t["acc_no"] for t in target if (t["account_sub"] or "") == (a["account_sub"] or "") and a["account_sub"]] \
            or [t["acc_no"] for t in target if t["account_classification"] == (a["account_classification"] or "").upper()]
        baru = _kode_baru(blok, terpakai)
        if baru is None:
            hasil["alasan"] = f"tidak bisa menentukan kode baru untuk akun {a['acc_no']} {a['account_name']}"
            return hasil
        terpakai.add(baru)
        rencana.append((a, {"acc_no": baru, "_kode_saja": True}))

    for a, t in rencana:
        kode_saja = t.get("_kode_saja", False)
        nama_baru = a["account_name"]
        if not kode_saja and a["account_name"].strip().lower() == nama_v1.get(a["acc_no"], "").strip().lower():
            nama_baru = t["account_name"]
        # Kode V1 tidak disimpan sebagai Source Account Code; source code asli user dipertahankan.
        src = a["source_acc_no"] if a["source_acc_no"] and a["source_acc_no"] != a["acc_no"] else None
        hasil["baris"].append({
            "akun": a, "v2": None if kode_saja else t, "acc_no_baru": t["acc_no"], "nama_baru": nama_baru,
            "source_baru": src, "status": "kode baru (di luar template)" if kode_saja else "dipetakan ke V2",
        })
    hasil["tambah"] = [t for t in target if t["acc_no"] not in v2_dipakai]

    # Validasi akhir: acc_no unik (termasuk baris terhapus) & kode tampil unik.
    final = [b["acc_no_baru"] for b in hasil["baris"]] + [t["acc_no"] for t in hasil["tambah"]]
    terhapus = {a["acc_no"] for a in semua if a["deleted_at"] is not None}
    bentrok = sorted((set(final) & terhapus) | {k for k in final if final.count(k) > 1})
    tampil = [(b["source_baru"] or b["acc_no_baru"]).lower() for b in hasil["baris"]] + [t["acc_no"].lower() for t in hasil["tambah"]]
    bentrok_tampil = sorted({k for k in tampil if tampil.count(k) > 1})
    if bentrok or bentrok_tampil:
        hasil["alasan"] = f"bentrok kode: {', '.join((bentrok + bentrok_tampil)[:10])} -- perbaiki manual dulu"
    return hasil


def hitung_referensi(conn, kid: str) -> Dict[str, int]:
    """Jumlah baris transaksi/jurnal per kode akun klien (untuk laporan)."""
    n: Dict[str, int] = defaultdict(int)
    for tabel, kolom, filter_ in TABEL_KODE:
        for kode, jml in conn.execute(text(
            f"SELECT t.{kolom}, count(*) FROM {tabel} t WHERE {filter_} AND t.{kolom} IS NOT NULL GROUP BY 1"
        ), {"kid": kid}).all():
            n[str(kode)] += jml
    return n


# ============================================================
# TERAPKAN
# ============================================================

def _trigger_kunci(conn) -> List[Tuple[str, str]]:
    tabel = sorted({t for t, _, _ in TABEL_KODE} | {"management_client_coa"})
    return [(r[0], r[1]) for r in conn.execute(text("""
        SELECT c.relname, g.tgname FROM pg_trigger g JOIN pg_class c ON c.oid = g.tgrelid
        WHERE NOT g.tgisinternal AND g.tgname LIKE 'trg_kunci_%' AND c.relname = ANY(:t)
    """), {"t": tabel}).all()]


def terapkan(conn, r: Dict[str, Any]) -> Dict[str, int]:
    kid = r["klien"]["id"]
    sekarang = datetime.now()
    pemicu = _trigger_kunci(conn)
    for tabel, nama in pemicu:
        conn.execute(text(f'ALTER TABLE {tabel} DISABLE TRIGGER "{nama}"'))

    # Tahap 1: kode sementara (hindari bentrok unique acc_no & kode tampil).
    for b in r["baris"]:
        conn.execute(text("UPDATE management_client_coa SET acc_no = :k, source_acc_no = NULL WHERE id = :id"),
                     {"k": f"~{b['akun']['id']}"[:50], "id": b["akun"]["id"]})
    # Tahap 2: kode final + atribut V2.
    for b in r["baris"]:
        isi = {"id": b["akun"]["id"], "acc_no": b["acc_no_baru"], "nama": b["nama_baru"],
               "src": b["source_baru"], "now": sekarang}
        set_v2 = ""
        if b["v2"] is not None:
            set_v2 = ", " + ", ".join(f"{k} = :{k}" for k in KOLOM_V2)
            isi.update({k: b["v2"].get(k) for k in KOLOM_V2})
        conn.execute(text(
            "UPDATE management_client_coa SET acc_no = :acc_no, account_name = :nama, source_acc_no = :src, "
            f"edited_at = :now{set_v2} WHERE id = :id"), isi)

    # Kode teks di tabel jurnal/transaksi (lewat tabel peta sementara).
    conn.execute(text("CREATE TEMP TABLE IF NOT EXISTS tmp_peta_coa (lama text PRIMARY KEY, baru text) ON COMMIT DROP"))
    conn.execute(text("DELETE FROM tmp_peta_coa"))
    peta = [{"lama": b["akun"]["acc_no"], "baru": b["acc_no_baru"]} for b in r["baris"] if b["akun"]["acc_no"] != b["acc_no_baru"]]
    if peta:
        conn.execute(text("INSERT INTO tmp_peta_coa (lama, baru) VALUES (:lama, :baru)"), peta)
    n_ref = 0
    for tabel, kolom, filter_ in TABEL_KODE:
        n_ref += conn.execute(text(
            f"UPDATE {tabel} t SET {kolom} = m.baru FROM tmp_peta_coa m WHERE t.{kolom} = m.lama AND {filter_}"
        ), {"kid": kid}).rowcount

    # Akun V2 yang belum dimiliki klien.
    kolom = ["acc_no", "account_name"] + list(KOLOM_V2)
    for t in r["tambah"]:
        conn.execute(text(
            f"INSERT INTO management_client_coa (client_id, client_code, is_active, created_at, {', '.join(kolom)}) "
            f"VALUES (:kid, :kode, true, :now, {', '.join(':' + k for k in kolom)})"
        ), {"kid": kid, "kode": r["klien"]["client_code"], "now": sekarang, **{k: t.get(k) for k in kolom}})

    for tabel, nama in pemicu:
        conn.execute(text(f'ALTER TABLE {tabel} ENABLE TRIGGER "{nama}"'))
    return {"akun": len(r["baris"]), "referensi": n_ref, "tambah": len(r["tambah"])}


# ============================================================
# MAIN
# ============================================================

def tulis_csv(path: Path, rencana: List[Dict[str, Any]], referensi: Dict[str, Dict[str, int]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["client", "coa_id", "status", "old_acc_no", "new_acc_no", "old_name", "new_name",
                    "standard_account_code", "transaction_rows_using_old_code"])
        for r in rencana:
            nama = r["klien"]["nama_client"]
            if r["alasan"]:
                w.writerow([nama, "", f"SKIPPED: {r['alasan']}", "", "", "", "", "", ""])
                continue
            ref = referensi.get(r["klien"]["id"], {})
            for b in r["baris"]:
                a = b["akun"]
                w.writerow([nama, a["id"], b["status"], a["acc_no"], b["acc_no_baru"], a["account_name"],
                            b["nama_baru"], a["standard_account_code"] or "", ref.get(a["acc_no"], 0)])
            for t in r["tambah"]:
                w.writerow([nama, "", "akun V2 ditambahkan", "", t["acc_no"], "", t["account_name"],
                            t["standard_account_code"] or "", 0])


def main() -> int:
    ap = argparse.ArgumentParser(description="Migrasi COA klien template V1 -> struktur template V2.")
    ap.add_argument("--client", action="append", help="Nama klien (nama_client) atau id. Bisa diulang. Default: semua.")
    ap.add_argument("--apply", action="store_true", help="Jalankan perubahan (default: dry-run).")
    ap.add_argument("--v1", default=str(FILE_V1), help="Path Excel template V1 (COA_Industry.xlsx).")
    args = ap.parse_args()

    v1 = baca_template_v1(Path(args.v1))
    mode = "APPLY" if args.apply else "DRY-RUN"
    print("=" * 70)
    print(f"Migrasi COA template V1 -> V2  [{mode}]")
    print("=" * 70)

    with engine.connect() as conn:
        v2 = baca_template_v2(conn)
        # Pengaman: template di DB harus sudah V2 (migration 30 dengan COA_Industry_V2.xlsx).
        kode_db = [a["acc_no"] for akun in v2.values() for a in akun]
        kode_v1 = {k for t in v1.values() for k in t["akun"]}
        if not kode_db or sum(1 for k in kode_db if k in kode_v1) / len(kode_db) >= AMBANG_V1:
            print("BERHENTI: template COA industri di database masih versi LAMA (V1) / kosong.")
            print("Jalankan dulu: python migrations/30-create_management_coa_industry_templates.py")
            print("(file dataset/COA/COA_Industry_V2.xlsx harus ada), lalu jalankan script ini lagi.")
            return 1
        klien = [dict(r) for r in conn.execute(text(
            "SELECT id::text AS id, nama_client, client_code, industry FROM management_clients ORDER BY nama_client"
        )).mappings().all()]
        if args.client:
            dicari = {c.strip().lower() for c in args.client}
            klien = [k for k in klien if k["id"].lower() in dicari or (k["nama_client"] or "").strip().lower() in dicari]
            if not klien:
                print("Klien tidak ditemukan.")
                return 1

        rencana, referensi = [], {}
        for k in klien:
            r = rencana_klien(conn, k, v1, v2)
            rencana.append(r)
            if r["alasan"]:
                print(f"- {k['nama_client']}: DILEWATI -- {r['alasan']}")
                continue
            referensi[k["id"]] = hitung_referensi(conn, k["id"])
            n_ref = sum(referensi[k["id"]].get(b["akun"]["acc_no"], 0) for b in r["baris"])
            n_luar = sum(1 for b in r["baris"] if b["v2"] is None)
            print(f"- {k['nama_client']}: template V1 KBLI {r['kbli']} ({r['porsi']:.0%} cocok) -> "
                  f"{len(r['baris'])} akun dipetakan ({n_luar} kode baru di luar template), "
                  f"{len(r['tambah'])} akun V2 ditambahkan, {n_ref} baris transaksi ikut diganti kodenya")
        conn.rollback()

        cap = datetime.now().strftime("%Y%m%d_%H%M%S")
        csv_path = FOLDER_HASIL / f"migrasi_coa_v2_{mode.lower()}_{cap}.csv"
        tulis_csv(csv_path, rencana, referensi)
        print(f"\nLaporan pemetaan: {csv_path}")

        if not args.apply:
            print("Dry-run: tidak ada data yang diubah. Tambahkan --apply untuk menjalankan.")
            return 0

        gagal = 0
        for r in rencana:
            if r["alasan"]:
                continue
            nama = r["klien"]["nama_client"]
            try:
                with conn.begin():
                    hasil = terapkan(conn, r)
                print(f"OK  {nama}: {hasil['akun']} akun diubah, {hasil['referensi']} baris transaksi, "
                      f"{hasil['tambah']} akun ditambahkan")
            except Exception as e:  # noqa: BLE001 -- 1 klien gagal tidak menghentikan klien lain
                gagal += 1
                print(f"GAGAL {nama}: {e}")
        return 1 if gagal else 0


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    sys.exit(main())
