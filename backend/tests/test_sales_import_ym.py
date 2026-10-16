"""
tests/test_sales_import_ym.py
=============================
Uji template Sales klien YM (PT Yoniso Manasikara): resep DIBACA LANGSUNG dari
migrations/sales/seed_template_ym_multi_format_sales.sql (bukan salinan), lalu dipakai
untuk mencocokkan + mengekstrak workbook buatan di memory yang meniru file asli
(INCOME LIST xlsx dan Pemasukan xls) beserta variasi yang mungkin dikirim pengguna.

Yang diuji: pencocokan lewat label header, pemilihan sheet per bulan ("YANG DI PAKAI"),
sheet tersembunyi, baris total, tanggal/angka berbentuk teks, rumus belum dihitung, valuta
asing tanpa kurs, nomor invoice ganda, dan konversi .xls -> .xlsx.
"""

import datetime as dt
import hashlib
import io
import json
import os
import re
from pathlib import Path

import openpyxl
import pytest

os.environ.setdefault("DATABASE_URL", "postgresql+psycopg2://u:p@localhost:5432/db")
os.environ.setdefault("JWT_SECRET_KEY", "test-only")

import modules.transactions.sales_import_v1 as si  # noqa: E402

SEED = Path(__file__).parent.parent / "migrations" / "sales" / "seed_template_ym_multi_format_sales.sql"
D = dt.datetime


def _varian_seed():
    teks = SEED.read_text(encoding="utf-8")
    return json.loads(re.search(r"\$json\$(.*?)\$json\$", teks, re.S).group(1))["variants"]


VARIAN = _varian_seed()
HDR_A = VARIAN[0]["header_columns"]
HDR_B = VARIAN[1]["header_columns"]


def _jalankan(content, varian=None, nama="file.xlsx"):
    si.dbc.list_sales_import_templates = lambda client_id=None, hanya_aktif=True: [
        {"id": "t1", "client_id": "c", "file_type": "Multi", "mapping_rules": {"variants": varian or VARIAN}}
    ]
    cocok = si._cocokkan_template(content, "Excel", "c")
    if cocok is None:
        return None, []
    template, sumber = cocok
    return template, si._ekstrak_dengan_template(sumber, template, nama)


def _ringkas(res):
    valid = [r for r in res if r["is_valid"]]
    return len(valid), len(res) - len(valid), round(sum(r["dpp"] for r in valid), 2)


def _simpan(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _baris_a(tgl, nama, total, paid="Cash", komponen=(0, 0, 0)):
    # A tanggal, B nama, C nights, D room, E dives, F extra, G eur, H rp, I paid, J diving, K room, L total
    return [tgl, nama, 2, "Bungalow", 3, komponen[2], None, None, paid, komponen[0], komponen[1], total]


def _wb_income_list(sheets):
    """sheets: {nama: (baris, opsi)}; opsi: state, hdr (baris header), header, total (bool), total_val."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Rekap").sheet_state = "visible"
    for nama, (baris, opsi) in sheets.items():
        ws = wb.create_sheet(nama)
        ws.sheet_state = opsi.get("state", "visible")
        baris_header = opsi.get("hdr", 10)
        ws.cell(1, 1, "North Bali Beach Cottages & Spa")
        for j, v in enumerate(opsi.get("header", HDR_A)):
            ws.cell(baris_header, j + 1, v)
        r = baris_header + 1
        for isi in baris:
            for j, v in enumerate(isi):
                if v is not None:
                    ws.cell(r, j + 1, v)
            r += 1
        if opsi.get("total", True):
            ws.cell(r + 3, 1, "Total")
            ws.cell(r + 3, 12, opsi.get("total_val", sum((x[11] or 0) for x in baris)))
            ws.cell(r + 5, 1, "Income NET Room")  # blok ringkasan: label di A, angka di B
            ws.cell(r + 5, 2, 4000000)
    return _simpan(wb)


# ---------------------------------------------------------------- seed

def test_hash_seed_konsisten_dengan_kode():
    for v in VARIAN:
        assert si._hash_baris(v["header_columns"]) == v["column_signature_hash"]
    teks = SEED.read_text(encoding="utf-8")
    hash_baris = hashlib.sha256(":".join(v["column_signature_hash"] for v in VARIAN).encode()).hexdigest()
    assert f"v_hash_baru   CONSTANT TEXT := '{hash_baris}'" in teks


# ---------------------------------------------------------------- varian #1 (xlsx)

def test_xlsx_dasar_dan_nama_dirapikan():
    c = _wb_income_list({"0826": ([_baris_a(D(2026, 8, 14), "Ines  & Elias", 1000), _baris_a(D(2026, 8, 15), "Maya", 500)], {})})
    t, res = _jalankan(c)
    assert _ringkas(res) == (2, 0, 1500.0)
    assert res[0]["nama_customer"] == "Ines & Elias" and res[0]["no_invoice"] == "YM-20260814-INES-ELIAS"
    assert res[0]["ppn"] == 0 and res[0]["total"] == res[0]["dpp"] and res[0]["cabang"] is None


def test_sheet_yang_dipakai_menang_dan_tersembunyi_dilewati():
    c = _wb_income_list({
        "0826": ([_baris_a(D(2026, 8, 14), "A", 1000)], {}),
        "0826 YANG DI PAKAI ": ([_baris_a(D(2026, 8, 14), "A", 700)], {}),
        "0725": ([_baris_a(D(2025, 7, 1), "B", 9999)], {"state": "hidden"}),
    })
    _, res = _jalankan(c)
    assert _ringkas(res) == (1, 0, 700.0)


def test_semua_sheet_data_tersembunyi_diperingatkan():
    c = _wb_income_list({"0826": ([_baris_a(D(2026, 8, 14), "A", 1000)], {"state": "hidden"})})
    t, res = _jalankan(c)
    assert t is not None and len(res) == 1 and not res[0]["is_valid"] and "tersembunyi" in res[0]["validation_notes"]


def test_header_bergeser_kolom_acak_label_beda_dan_nilai_teks():
    hdr = ["", "DATE OF INVOICE ", "name", "Total Rp.", "Part  Diving", "part r&b&s", "Extra Income", "Total", "xtra"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agustus 2026"
    ws.cell(3, 1, "kop")
    for j, v in enumerate(hdr):
        ws.cell(6, j + 1, v)
    ws.cell(7, 2, "2026-09-03"); ws.cell(7, 3, "Tamu Satu"); ws.cell(7, 5, 100); ws.cell(7, 6, 50); ws.cell(7, 7, 25); ws.cell(7, 8, 175)
    ws.cell(8, 2, "05/09/2026"); ws.cell(8, 3, "Tamu Dua"); ws.cell(8, 8, "1.500.000")
    ws.cell(9, 2, "7./09/2026"); ws.cell(9, 3, "Tamu Tiga"); ws.cell(9, 5, 10); ws.cell(9, 6, 20)  # Total kosong -> komponen
    ws.cell(10, 2, "09-09-26"); ws.cell(10, 3, "Tamu Empat"); ws.cell(10, 8, 5)
    _, res = _jalankan(_simpan(wb))
    assert _ringkas(res)[:2] == (4, 0) and res[2]["dpp"] == 30.0
    assert [r["tanggal"].day for r in res] == [3, 5, 7, 9]


def test_baris_total_hilang_aman_total_salah_diperingatkan():
    _, res = _jalankan(_wb_income_list({"0826": ([_baris_a(D(2026, 8, 14), "A", 1000)], {"total": False})}))
    assert _ringkas(res) == (1, 0, 1000.0)
    _, res = _jalankan(_wb_income_list({"0826": ([_baris_a(D(2026, 8, 14), "A", 1000)], {"total_val": 5000})}))
    assert _ringkas(res)[:2] == (1, 1)
    assert "tidak sama dengan jumlah" in next(r for r in res if not r["is_valid"])["validation_notes"]


def test_nomor_invoice_ganda_dan_akhiran():
    baris = [_baris_a(D(2026, 8, 14), "A", 1000), _baris_a(D(2026, 8, 14), "A", 1000), _baris_a(D(2026, 8, 14), "A", 300)]
    _, res = _jalankan(_wb_income_list({"0826": (baris, {})}))
    assert [r["no_invoice"] for r in res] == ["YM-20260814-A", "YM-20260814-A", "YM-20260814-A-2"]
    assert [r["is_duplicate_candidate"] for r in res] == [False, True, False]


def test_baris_bermasalah_tidak_hilang_diam_diam():
    baris = [_baris_a(D(2026, 8, 14), "A", -50), _baris_a("bukan tanggal", "B", 10), _baris_a(D(2026, 8, 1), "", 77), _baris_a(D(2026, 7, 30), "C", 5)]
    _, res = _jalankan(_wb_income_list({"0826": (baris, {"total": False})}))
    catatan = " | ".join((r["validation_notes"] or "") for r in res)
    assert "negatif" in catatan and "Tanggal tidak terbaca" in catatan and "tanpa nama tamu" in catatan
    ok = [r for r in res if r["is_valid"]]
    assert len(ok) == 1 and "di luar bulan sheet" in ok[0]["validation_notes"]


def test_sheet_tanpa_header_diperingatkan_sheet_kosong_diabaikan():
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Catatan")["A1"] = "halo"
    ws = wb.create_sheet("0826")
    for j, v in enumerate(HDR_A):
        ws.cell(10, j + 1, v)
    ws.cell(11, 1, D(2026, 8, 1)); ws.cell(11, 2, "A"); ws.cell(11, 12, 10)
    _, res = _jalankan(_simpan(wb))
    assert _ringkas(res)[:2] == (1, 1)


def test_rumus_belum_dihitung_jadi_invalid():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "0826"
    for j, v in enumerate(HDR_A):
        ws.cell(10, j + 1, v)
    ws.cell(11, 1, D(2026, 8, 1)); ws.cell(11, 2, "A"); ws.cell(11, 5, 3)
    ws.cell(11, 10, "=E11*330000"); ws.cell(11, 12, "=J11+K11+F11")
    _, res = _jalankan(_simpan(wb))
    assert _ringkas(res)[:2] == (0, 1) and "rumus" in res[0]["validation_notes"]


def test_file_lain_tidak_cocok():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "bukan laporan YM"
    t, res = _jalankan(_simpan(wb))
    assert t is None and res == []


# ---------------------------------------------------------------- varian #2 (xls / Total Rp + EUR)

def _wb_pemasukan(baris, nama_sheet="Jan 26"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nama_sheet
    for j, v in enumerate(HDR_B):
        ws.cell(7, j + 1, v)
    for i, isi in enumerate(baris):
        for j, v in enumerate(isi):
            if v is not None:
                ws.cell(8 + i, j + 1, v)
    return wb


def test_pemasukan_cash_euro_tanpa_kurs_invalid_dan_tanggal_teks_nyasar():
    # A tanggal, B nama, C nights, D room, E dive, F total rp, G extra, H total eur, I paid
    wb = _wb_pemasukan([
        ["12./02/2026", "Ann", 7, "Beachhouse", 12, 38517748, None, None, "Cc.Mandiri"],
        [D(2026, 2, 21), "Ivonne", 10, "Beachhouse", 25, None, None, 2053, "Cash"],
    ], "Feb 26")
    t, res = _jalankan(_simpan(wb))
    assert t["mapping_rules"]["valuta_asing"]["kolom"] == "total_eur"
    assert _ringkas(res) == (1, 1, 38517748.0) and res[0]["tanggal"] == dt.date(2026, 2, 12)
    assert "tanpa kurs" in res[1]["validation_notes"]


def test_pemasukan_kurs_diisi_dikonversi():
    varian = json.loads(json.dumps(VARIAN))
    varian[1]["mapping_rules"]["valuta_asing"]["kurs"] = 17000
    wb = _wb_pemasukan([[D(2026, 1, 6), "Leoni", 2, "Beachhouse", None, None, None, 100.5, "Cash"]])
    _, res = _jalankan(_simpan(wb), varian)
    assert _ringkas(res) == (1, 0, 1708500.0)


def test_pemasukan_baris_total_tanpa_label_tidak_jadi_invoice():
    wb = _wb_pemasukan([
        [D(2026, 1, 5), "Lana", 6, "Bungalow", 25, 100, None, None, "Cc.Mandiri"],
        [D(2026, 1, 5), "Stefan", None, None, None, 50, None, None, "Cc.Mandiri"],
        [None], [D(2026, 1, 1), None, 6, None, 25, 150, 0, None],  # total: nama kosong, nilai = jumlah
    ])
    _, res = _jalankan(_simpan(wb))
    assert _ringkas(res) == (2, 0, 150.0)


def test_xls_asli_dikonversi_dan_terbaca():
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    ws = book.add_sheet("Maret 26")
    for j, v in enumerate(HDR_B):
        ws.write(6, j, v)
    ws.write(7, 0, "27./03/2026"); ws.write(7, 1, "Linus"); ws.write(7, 5, 3071113)
    fmt = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    ws.write(8, 0, D(2026, 3, 9), fmt); ws.write(8, 1, "Kerstin"); ws.write(8, 5, 100)
    book.add_sheet("Sheet1")
    b = io.BytesIO()
    book.save(b)
    assert b.getvalue().startswith(si._OLE2_MAGIC)
    t, res = _jalankan(b.getvalue(), nama="Pemasukan.xls")
    assert t is not None and _ringkas(res) == (2, 0, 3071213.0)
    assert [r["tanggal"] for r in res] == [dt.date(2026, 3, 27), dt.date(2026, 3, 9)]


def test_pastikan_xlsx_tidak_mengubah_xlsx_dan_data_rusak():
    xlsx = _simpan(openpyxl.Workbook())
    assert si._pastikan_xlsx(xlsx) is xlsx
    rusak = si._OLE2_MAGIC + b"bukan excel"
    assert si._pastikan_xlsx(rusak) == rusak