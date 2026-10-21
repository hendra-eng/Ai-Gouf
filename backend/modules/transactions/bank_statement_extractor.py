"""
modules/transactions/bank_statement_extractor.py
================================================
Ekstraksi rekening koran (PDF / Excel / CSV) untuk Bank Feed Cash & Bank --
murni fungsi (tanpa DB), dipakai modules/transactions/bank_cash_v1.py.

Mengikuti skills/pdf-reading/SKILL.md:

  1. Jalur bertingkat untuk PDF:
       grid  -- pdfplumber page.extract_tables() (PDF dengan garis tabel)
       words -- posisi kata pdfplumber (PDF tanpa garis tabel, mis. BCA Tahapan)
       ocr   -- halaman tanpa lapisan teks (hasil scan) -> Tesseract, lalu
                hasil kata-nya diproses PERSIS seperti jalur words
     Excel/CSV lewat jalur 'table'. Kalau beberapa jalur menghasilkan baris,
     dipilih yang validasinya paling baik (bukan sekadar yang pertama jalan).
  2. Skema kolom TIDAK diasumsikan: header tabel dideteksi dari kata kunci
     (Tanggal/Date, Keterangan/Description, Debet/Debit, Kredit/Credit,
     Mutasi/Amount + penanda DB/CR, Saldo/Balance), format angka (1,234.56
     vs 1.234,56) & urutan tanggal dideteksi dari isi.
  3. Tahun diambil dari tanggal lengkap atau dari periode di header dokumen
     ("PERIODE : APRIL 2025", "01/04/2025 - 30/04/2025"), bukan ditebak.
  4. Validasi: saldo berjalan per baris (saldo sebelumnya + kredit - debit =
     saldo tercetak; bank yang mencetak saldo hanya di baris terakhir per hari
     divalidasi per segmen) + total/saldo resmi di ringkasan dokumen.
  5. Confidence PER BARIS: high (terverifikasi), medium (tidak ada data untuk
     verifikasi), low (tidak cocok / angka meragukan -> review manual).
  6. Hasil baca dokumen (bagian mahal: kata, tabel, OCR) di-cache per hash
     isi file, jadi preview lalu import file yang sama tidak membaca ulang.

Template ekstraksi (financial_transaction_bank_cash_import_templates):
signature = SHA-256 teks header tabel yang dinormalisasi. Kalau signature
cocok dengan template aktif klien, peran kolom & aturan angka/tanggal dari
template dipakai langsung (lihat ekstrak_rekening_koran).

Arah: debit = uang KELUAR dari rekening, credit = uang MASUK (sudut pandang
rekening koran, sama dengan tabel mutations).
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import threading
import time
from bisect import bisect_right
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from ..logging_config import get_module_logger

logger = get_module_logger("transactions_bank_statement_extractor")

NOL = Decimal("0")
TOLERANSI_SALDO = Decimal("0.01")
ROLES = ("date", "description", "debit", "credit", "amount", "dc_flag", "balance", "balance_before", "reference", "ignore")
ROLE_ANGKA = {"debit", "credit", "amount", "balance", "balance_before"}
ROLE_TEKS = {"description", "reference", "ignore", "date"}
METODE = ("grid", "words", "ocr", "table")

OCR_DPI = 300
OCR_WORKER_MAKS = int(os.environ.get("BANK_STATEMENT_OCR_WORKER_MAKS", "4"))
OCR_CONF_RENDAH = 75.0
MIN_KATA_LAPISAN_TEKS = 8   # halaman dengan kata lebih sedikit dari ini dianggap hasil scan

_CACHE_MAKS = 6
_CACHE_TTL_DETIK = 30 * 60
_cache_dokumen: "OrderedDict[str, Tuple[float, Dokumen]]" = OrderedDict()
_cache_lock = threading.Lock()


class EkstraksiGagal(Exception):
    """Dokumen tidak bisa dibaca sama sekali (terkunci password, rusak, scan tanpa OCR)."""

    def __init__(self, pesan: str, kode: str = "PARSE_ERROR"):
        super().__init__(pesan)
        self.kode = kode


# ============================================================
# 1) Kata kunci header & normalisasi
# ============================================================

def _norm(s: Any) -> str:
    s = str(s or "").upper().replace("\n", " ")
    s = re.sub(r"[^A-Z0-9/ ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_KUNCI_PERSIS: Dict[str, str] = {}
for _role, _daftar in {
    "date": ["TANGGAL", "TGL", "DATE", "TRANS DATE", "TRANSACTION DATE", "POSTING DATE", "POST DATE", "TGL TRANSAKSI",
             "TANGGAL TRANSAKSI", "TXN DATE", "TRX DATE", "TANGGAL POSTING", "BOOK DATE", "TGL TRANS", "TANGGAL TRANS"],
    "ignore": ["VALUE DATE", "TANGGAL VALUTA", "TGL VALUTA", "EFFECTIVE DATE", "VALUTA", "EFF DATE", "NO", "NO URUT"],
    "description": ["KETERANGAN", "URAIAN", "DESKRIPSI", "DESCRIPTION", "REMARKS", "REMARK", "TRANSACTION DESCRIPTION",
                    "TRANSACTION DETAILS", "DETAILS", "DETAIL", "URAIAN TRANSAKSI", "KETERANGAN TRANSAKSI", "NARRATIVE",
                    "PARTICULARS", "TRANSAKSI", "TRANSACTION", "BERITA", "INFORMASI", "DETAIL TRANSAKSI",
                    "DESK TRAN", "DESK TRANSAKSI", "DESKRIPSI TRANSAKSI"],
    "debit": ["DEBET", "DEBIT", "DB", "DR", "PENARIKAN", "WITHDRAWAL", "WITHDRAWALS", "MUTASI DEBET", "MUTASI DEBIT",
              "KELUAR", "UANG KELUAR", "DEBIT AMOUNT", "DEBET IDR", "DEBIT IDR", "PENGELUARAN", "KAS KELUAR"],
    "credit": ["KREDIT", "CREDIT", "CR", "KR", "SETORAN", "DEPOSIT", "DEPOSITS", "MUTASI KREDIT", "MASUK", "UANG MASUK",
               "CREDIT AMOUNT", "KREDIT IDR", "CREDIT IDR", "PEMASUKAN", "KAS MASUK"],
    "amount": ["MUTASI", "JUMLAH", "NOMINAL", "AMOUNT", "NILAI", "TRANSACTION AMOUNT", "JUMLAH TRANSAKSI", "MUTASI IDR"],
    "dc_flag": ["D/K", "D/C", "DB/CR", "CR/DB", "DK", "DC", "K/D", "C/D", "DEBIT/CREDIT", "DEBET/KREDIT", "D K", "DB CR"],
    "balance": ["SALDO", "BALANCE", "SALDO AKHIR", "ENDING BALANCE", "RUNNING BALANCE", "SALDO BERJALAN", "SALDO IDR",
                "BALANCE IDR"],
    "reference": ["CBG", "CABANG", "BRANCH", "REF", "NO REF", "REFERENSI", "REFERENCE", "NO REFERENSI", "REFF", "JOURNAL",
                  "NO JURNAL", "TELLER", "KODE", "CODE", "CHEQUE", "NO CEK", "CEK/BG", "SEQ", "KODE TRANSAKSI",
                  "TRANSACTION CODE", "REFERENCE NO", "REF NO", "JOURNAL NO", "USER ID",
                  # kolom kategori/akun di laporan kas tunai (mis. Laporan Kas Kasir SAU) -- ikut ke keterangan
                  "ACCOUNT", "AKUN", "KATEGORI", "CATEGORY"],
}.items():
    for _k in _daftar:
        _KUNCI_PERSIS[_k] = _role


def peran_header(teks: str) -> Optional[str]:
    """Peran 1 sel header, atau None kalau tidak dikenali."""
    n = _norm(teks)
    if not n:
        return None
    n = re.sub(r"\b(IDR|RP)\b", "", n).strip()
    if n in _KUNCI_PERSIS:
        return _KUNCI_PERSIS[n]
    kata = set(n.split())
    # Export core banking (mis. CSV BRI): SALDO_AWAL_MUTASI = saldo SEBELUM transaksi,
    # SALDO_AKHIR_MUTASI = saldo setelahnya (saldo berjalan biasa).
    if kata & {"SALDO", "BALANCE"} and kata & {"AWAL", "SEBELUM", "BEFORE", "OPENING", "PREVIOUS"}             and kata & {"MUTASI", "TRANSAKSI", "TRANSACTION", "SEBELUM", "BEFORE"}:
        return "balance_before"
    if kata & {"SALDO", "BALANCE"}:
        return "balance"
    debet = bool(kata & {"DEBET", "DEBIT"})
    kredit = bool(kata & {"KREDIT", "CREDIT"})
    if debet and kredit:
        return "dc_flag" if len(n) <= 14 else "amount"
    if debet:
        return "debit"
    if kredit:
        return "credit"
    # Laporan kas tunai: "KAS MASUK" / "UANG KELUAR" (sudut pandang kas: masuk = credit).
    if kata & {"MASUK", "PEMASUKAN"} and not kata & {"KELUAR", "PENGELUARAN"}:
        return "credit"
    if kata & {"KELUAR", "PENGELUARAN"} and not kata & {"MASUK", "PEMASUKAN"}:
        return "debit"
    if kata & {"VALUTA", "VALUE", "EFFECTIVE"}:
        return "ignore"
    if kata & {"TANGGAL", "TGL", "DATE"}:
        return "date"
    if kata & {"KETERANGAN", "URAIAN", "DESKRIPSI", "DESCRIPTION", "REMARK", "REMARKS", "NARRATIVE", "PARTICULARS"}:
        return "description"
    if kata & {"MUTASI", "AMOUNT", "JUMLAH", "NOMINAL"}:
        return "amount"
    if kata & {"REF", "REFERENSI", "REFERENCE", "CABANG", "BRANCH", "CBG"}:
        return "reference"
    return None


def _skor_header(roles: List[Optional[str]]) -> int:
    """>0 kalau deretan peran ini layak dianggap header tabel mutasi."""
    r = [x for x in roles if x]
    if "date" not in r or "description" not in r:
        return 0
    if not ({"amount", "debit", "credit"} & set(r)):
        return 0
    if len(r) < 3:
        return 0
    return len(r) * 10 + (5 if "balance" in r else 0) - (len(roles) - len(r))


def _rapikan_peran(roles: List[Optional[str]]) -> List[str]:
    """Peran final per kolom: tanggal/keterangan/saldo ganda -> kolom berikutnya diabaikan
    (mis. Tanggal Transaksi + Tanggal Valuta); kolom tak dikenal -> ignore."""
    hasil: List[str] = []
    sudah: set = set()
    for r in roles:
        r = r or "ignore"
        if r in {"date", "balance", "balance_before", "dc_flag", "amount"} and r in sudah:
            r = "ignore"
        if r == "description" and "description" in sudah:
            r = "reference"
        sudah.add(r)
        hasil.append(r)
    return hasil


def signature_header(headers: List[str]) -> str:
    """Kolom kosong di ujung kanan (lebar sheet Excel bisa berubah tiap bulan karena panel
    catatan di samping tabel) tidak ikut signature."""
    n = [_norm(h) for h in headers]
    while n and not n[-1]:
        n.pop()
    return hashlib.sha256("|".join(n).encode("utf-8")).hexdigest()


# ============================================================
# 2) Angka, tanggal, periode
# ============================================================

_FLAG_DEBIT = {"DB", "D", "DR", "DEBIT", "DEBET", "DBT"}
_FLAG_KREDIT = {"CR", "C", "K", "KR", "CREDIT", "KREDIT", "CRT"}
_RE_ANGKA_TOKEN = re.compile(
    r"^[-(]?(?:RP\.?|IDR)?\s*-?(?:\d{1,3}(?:[.,]\d{3})+|\d+)(?:[.,]\d{1,2})?\)?-?(?:\s*(?:DB|CR|DR|D|K|C|KR))?$", re.I
)
_RE_ANGKA_BEBAS = re.compile(r"-?\(?\d{1,3}(?:[.,]\d{3})+(?:[.,]\d{1,2})?\)?|-?\d+[.,]\d{2}(?!\d)")
_RE_EN = re.compile(r"^\d{1,3}(,\d{3})+(\.\d{1,2})?$|^\d+\.\d{2}$")
_RE_ID = re.compile(r"^\d{1,3}(\.\d{3})+(,\d{1,2})?$|^\d+,\d{2}$")


def mirip_angka(teks: str) -> bool:
    t = (teks or "").strip()
    return bool(t) and any(c.isdigit() for c in t) and bool(_RE_ANGKA_TOKEN.match(t.replace(" ", "")))


def deteksi_format_angka(sampel: List[str]) -> str:
    en = idn = 0
    for s in sampel:
        t = re.sub(r"[^\d.,]", "", s or "")
        if _RE_EN.match(t):
            en += 1
        elif _RE_ID.match(t):
            idn += 1
    return "id" if idn > en else "en"


def parse_angka(teks: Any, fmt: str = "en") -> Tuple[Optional[Decimal], Optional[str]]:
    """'1,500,000.00 DB' -> (Decimal('1500000.00'), 'debit'). Penanda arah dari
    suffix DB/CR, tanda minus, atau kurung. (None, None) kalau bukan angka."""
    if teks is None:
        return None, None
    if isinstance(teks, (int, float, Decimal)) and not isinstance(teks, bool):
        try:
            v = Decimal(str(teks))
        except InvalidOperation:
            return None, None
        return abs(v), ("debit" if v < 0 else None)
    s = str(teks).strip().upper().replace("\u00a0", " ")
    if not s or not any(c.isdigit() for c in s):
        return None, None
    arah = None
    m = re.search(r"(DB|DR|CR|KR|D|K|C)\s*$", s)
    if m and not s[: m.start()].strip().endswith(("RP", "IDR")):
        kode = m.group(1)
        arah = "debit" if kode in _FLAG_DEBIT else "credit"
        s = s[: m.start()]
    s = s.replace("RP", "").replace("IDR", "").replace(" ", "")
    if s.startswith("(") and s.endswith(")"):
        arah, s = "debit", s[1:-1]
    if s.endswith("-"):
        arah, s = "debit", s[:-1]
    if s.startswith("-"):
        arah, s = "debit", s[1:]
    if s.startswith("+"):
        s = s[1:]
    if not re.fullmatch(r"[\d.,]+", s):
        return None, None
    if fmt == "id":
        if "," not in s and re.fullmatch(r"\d+\.\d{1,2}", s):
            pass  # angka mentah sel Excel (titik desimal), bukan pemisah ribuan
        else:
            s = s.replace(".", "").replace(",", ".")
    else:
        if "." not in s and re.fullmatch(r"\d+,\d{1,2}", s):
            s = s.replace(",", ".")
        else:
            s = s.replace(",", "")
    if s.count(".") > 1 or not re.fullmatch(r"\d+(\.\d+)?", s):
        return None, None
    try:
        return Decimal(s).quantize(Decimal("0.01")), arah
    except InvalidOperation:
        return None, None


_BULAN = {
    "JAN": 1, "JANUARI": 1, "JANUARY": 1, "FEB": 2, "PEB": 2, "FEBRUARI": 2, "PEBRUARI": 2, "FEBRUARY": 2,
    "MAR": 3, "MARET": 3, "MARCH": 3, "APR": 4, "APRIL": 4, "MEI": 5, "MAY": 5, "JUN": 6, "JUNI": 6, "JUNE": 6,
    "JUL": 7, "JULI": 7, "JULY": 7, "AGU": 8, "AGS": 8, "AGT": 8, "AUG": 8, "AGUSTUS": 8, "AUGUST": 8,
    "SEP": 9, "SEPT": 9, "SEPTEMBER": 9, "OKT": 10, "OCT": 10, "OKTOBER": 10, "OCTOBER": 10,
    "NOV": 11, "NOPEMBER": 11, "NOVEMBER": 11, "DES": 12, "DEC": 12, "DESEMBER": 12, "DECEMBER": 12,
}
_POLA_BULAN = "|".join(sorted(_BULAN, key=len, reverse=True))
_RE_TGL_ISO = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
_RE_TGL_ANGKA = re.compile(r"^(\d{1,2})[-/.](\d{1,2})(?:[-/.](\d{2,4}))?(?![\d/])")
_RE_TGL_TEKS = re.compile(rf"^(\d{{1,2}})[\s\-/.]*({_POLA_BULAN})\b[\s\-/.,]*(\d{{2,4}})?", re.I)
_RE_TGL_TEKS_US = re.compile(rf"^({_POLA_BULAN})\b[\s\-/.]*(\d{{1,2}}),?[\s\-/.]*(\d{{4}})?", re.I)


def _tahun4(y: Optional[str]) -> Optional[int]:
    if not y:
        return None
    v = int(y)
    return v + 2000 if v < 100 else v


def pecah_tanggal(teks: Any) -> Optional[Tuple[int, int, Optional[int], str]]:
    """Komponen tanggal mentah -> (a, b, tahun|None, jenis): jenis 'num' (a/b
    belum pasti hari/bulan), 'dm' (hari, bulan pasti), atau 'iso'."""
    if isinstance(teks, datetime):
        return teks.day, teks.month, teks.year, "dm"
    if isinstance(teks, date):
        return teks.day, teks.month, teks.year, "dm"
    s = str(teks or "").strip().upper()
    if not s:
        return None
    m = _RE_TGL_ISO.match(s)
    if m:
        return int(m.group(3)), int(m.group(2)), int(m.group(1)), "dm"
    m = _RE_TGL_TEKS.match(s)
    if m:
        return int(m.group(1)), _BULAN[m.group(2).upper()], _tahun4(m.group(3)), "dm"
    m = _RE_TGL_TEKS_US.match(s)
    if m:
        return int(m.group(2)), _BULAN[m.group(1).upper()], _tahun4(m.group(3)), "dm"
    m = _RE_TGL_ANGKA.match(s)
    if m:
        return int(m.group(1)), int(m.group(2)), _tahun4(m.group(3)), "num"
    return None


def deteksi_urutan_tanggal(sampel: List[Any]) -> str:
    dmy = mdy = 0
    for s in sampel:
        p = pecah_tanggal(s)
        if not p or p[3] != "num":
            continue
        if p[0] > 12 >= p[1]:
            dmy += 1
        elif p[1] > 12 >= p[0]:
            mdy += 1
    return "mdy" if mdy > dmy else "dmy"


def _cari_periode(teks: str) -> Tuple[Optional[date], Optional[date]]:
    """Periode statement dari teks halaman pertama (lihat docstring modul)."""
    t = re.sub(r"\s+", " ", (teks or "").upper())

    def _d(dd, mm, yy):
        try:
            return date(_tahun4(yy), int(mm), int(dd))
        except (TypeError, ValueError):
            return None

    # "FROM : 2026-06-01 TO : 2026-06-30" (OCBC) -- ISO dicek lebih dulu supaya
    # "2026-06-01" tidak terbaca sebagai dd-mm-yy "26-06-01".
    sambung = r"\s*(?:-|S/?D|SAMPAI(?: DENGAN)?|HINGGA|TO|UNTIL)\s*:?\s*"
    m = re.search(rf"(\d{{4}})-(\d{{2}})-(\d{{2}}){sambung}(\d{{4}})-(\d{{2}})-(\d{{2}})", t)
    if m:
        return _d(m.group(3), m.group(2), m.group(1)), _d(m.group(6), m.group(5), m.group(4))
    m = re.search(rf"(\d{{1,2}})[/.-](\d{{1,2}})[/.-](\d{{2,4}}){sambung}(\d{{1,2}})[/.-](\d{{1,2}})[/.-](\d{{2,4}})", t)
    if m:
        return _d(m.group(1), m.group(2), m.group(3)), _d(m.group(4), m.group(5), m.group(6))
    m = re.search(rf"(\d{{1,2}})[ -]({_POLA_BULAN})[ -](\d{{2,4}}){sambung}(\d{{1,2}})[ -]({_POLA_BULAN})[ -](\d{{2,4}})", t)
    if m:
        return (_d(m.group(1), _BULAN[m.group(2)], m.group(3)), _d(m.group(4), _BULAN[m.group(5)], m.group(6)))
    # "1-30 JUNI 2026" (judul laporan kas tunai)
    m = re.search(rf"\b(\d{{1,2}})\s*-\s*(\d{{1,2}})\s+({_POLA_BULAN})\s+(\d{{4}})\b", t)
    if m:
        bln = _BULAN[m.group(3)]
        return _d(m.group(1), bln, m.group(4)), _d(m.group(2), bln, m.group(4))
    m = re.search(rf"(?:PERIODE|PERIOD|BULAN|MONTH)\s*:?\s*({_POLA_BULAN})\s*(\d{{4}})", t)
    if m:
        bln, th = _BULAN[m.group(1)], int(m.group(2))
        akhir = date(th + (bln == 12), bln % 12 + 1, 1)
        return date(th, bln, 1), date.fromordinal(akhir.toordinal() - 1)
    m = re.search(r"(?:PERIODE|PERIOD)[^0-9]{0,20}(\d{4})", t)
    if m:
        th = int(m.group(1))
        return date(th, 1, 1), date(th, 12, 31)
    return None, None


_BANK = [
    ("BCA", r"\bBCA\b|BANK CENTRAL ASIA|REKENING TAHAPAN|TAHAPAN"), ("Mandiri", r"\bMANDIRI\b"),
    ("BRI", r"\bBRI\b|BANK RAKYAT INDONESIA|\bBRITAMA\b|\bIBBIZ\b"), ("BNI", r"\bBNI\b|BANK NEGARA INDONESIA"), ("BSI", r"\bBSI\b|BANK SYARIAH INDONESIA"),
    ("CIMB Niaga", r"CIMB"), ("Permata", r"PERMATA"), ("Danamon", r"DANAMON"), ("OCBC", r"\bOCBC\b"), ("BTN", r"\bBTN\b"),
    ("Maybank", r"MAYBANK"), ("Panin", r"\bPANIN\b"), ("Bank Jago", r"\bJAGO\b"), ("SeaBank", r"SEABANK"),
    ("BJB", r"\bBJB\b"), ("Mega", r"BANK MEGA"), ("UOB", r"\bUOB\b"), ("HSBC", r"\bHSBC\b"), ("DBS", r"\bDBS\b"),
]


def deteksi_nama_bank(teks: str) -> Optional[str]:
    """Untuk label template & mapping bank -> akun COA (lihat bank_cash_v1._cari_mapping_bank).
    Logika ekstraksi TIDAK bergantung nama bank."""
    t = (teks or "").upper()
    for nama, pola in _BANK:
        if re.search(pola, t):
            return nama
    return None


_RE_LABEL_REKENING = re.compile(
    r"NO(?:MOR)?\.?\s*REK(?:ENING)?\.?|ACCOUNT\s*(?:NO\.?|NUMBER)|NO\.?\s*(?:OF\s*)?ACCOUNT|NOMOR\s*AKUN", re.I)


def deteksi_no_rekening(teks: str) -> Optional[str]:
    """Nomor rekening (digit saja) dari header dokumen: angka 6-20 digit pertama dalam
    ~80 karakter setelah label "No. Rekening" / "Account No." (label & nilai bisa berupa
    blok kolom: "Account No. Account Name Alias 1290011863475 ...")."""
    t = re.sub(r"\s+", " ", teks or "")
    for m in _RE_LABEL_REKENING.finditer(t):
        am = re.search(r"(?<![\d.,])(\d[\d-]{4,24}\d)(?![\d.,])", t[m.end(): m.end() + 80])
        if am:
            digit = re.sub(r"\D", "", am.group(1))
            if 6 <= len(digit) <= 20:
                return digit
    return None


# ============================================================
# 3) Baca dokumen (bagian mahal, di-cache per hash isi)
# ============================================================

@dataclass
class Halaman:
    no: int
    words: List[Dict[str, Any]]        # {text, x0, x1, top, bottom, conf?}
    tables: List[List[List[str]]]
    ocr: bool = False


@dataclass
class Dokumen:
    file_type: str                     # PDF / Excel / CSV
    halaman: List[Halaman] = field(default_factory=list)
    sheets: List[Tuple[str, List[List[str]]]] = field(default_factory=list)
    teks_awal: str = ""
    peringatan: List[str] = field(default_factory=list)
    ocr_dipakai: bool = False


def jenis_file(nama_file: str) -> Optional[str]:
    n = (nama_file or "").lower()
    if n.endswith(".pdf"):
        return "PDF"
    if n.endswith((".xlsx", ".xlsm", ".xls")):
        return "Excel"
    if n.endswith((".csv", ".txt")):
        return "CSV"
    return None


def _cmd_tesseract() -> Optional[str]:
    import shutil
    kandidat = [os.environ.get("TESSERACT_CMD"), shutil.which("tesseract"),
                r"C:\Program Files\Tesseract-OCR\tesseract.exe", r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                "/usr/bin/tesseract", "/usr/local/bin/tesseract", "/opt/homebrew/bin/tesseract"]
    for k in kandidat:
        if k and os.path.isfile(k):
            return k
    return None


def ocr_tersedia() -> bool:
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return _cmd_tesseract() is not None


def _ocr_halaman(gambar, skala: float, bahasa: str) -> List[Dict[str, Any]]:
    import pytesseract
    data = pytesseract.image_to_data(gambar, lang=bahasa, config="--psm 6 -c preserve_interword_spaces=1",
                                     output_type=pytesseract.Output.DICT)
    words = []
    for i, teks in enumerate(data["text"]):
        teks = (teks or "").strip()
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1.0
        if not teks or conf < 0:
            continue
        x, y, w, h = data["left"][i], data["top"][i], data["width"][i], data["height"][i]
        words.append({"text": teks, "x0": x / skala, "x1": (x + w) / skala, "top": y / skala, "bottom": (y + h) / skala, "conf": conf})
    return words


def _baca_pdf(isi: bytes) -> Dokumen:
    import pdfplumber
    try:
        from pdfminer.pdfdocument import PDFPasswordIncorrect
    except ImportError:  # pragma: no cover
        PDFPasswordIncorrect = Exception  # type: ignore

    dok = Dokumen(file_type="PDF")
    perlu_ocr: List[int] = []
    gambar_ocr: Dict[int, Any] = {}
    try:
        pdf = pdfplumber.open(io.BytesIO(isi), password="")
    except PDFPasswordIncorrect:
        raise EkstraksiGagal("This PDF is password-protected. Open it, save a copy without a password, then upload again.", "PDF_LOCKED")
    except Exception as e:
        raise EkstraksiGagal(f"The PDF could not be opened (the file may be damaged): {e}")
    with pdf:
        for no, page in enumerate(pdf.pages, 1):
            try:
                words = page.extract_words(keep_blank_chars=False, use_text_flow=False, x_tolerance=2, y_tolerance=2)
            except Exception:
                words = []
            words = [{"text": w["text"], "x0": float(w["x0"]), "x1": float(w["x1"]),
                      "top": float(w["top"]), "bottom": float(w["bottom"])} for w in words if (w.get("text") or "").strip()]
            tables: List[List[List[str]]] = []
            if len(words) >= MIN_KATA_LAPISAN_TEKS:
                # Jalur grid hanya untuk halaman yang memang bergaris (extract_tables mahal);
                # halaman tanpa garis antar-baris langsung lewat jalur words.
                if len(page.lines) + len(page.rects) < 6:
                    dok.halaman.append(Halaman(no=no, words=words, tables=tables))
                    continue
                try:
                    for tb in page.extract_tables() or []:
                        tables.append([[(c or "").strip() for c in (row or [])] for row in tb if row])
                except Exception:
                    pass
            else:
                perlu_ocr.append(no)
                try:
                    gambar_ocr[no] = page.to_image(resolution=OCR_DPI).original
                except Exception as e:
                    dok.peringatan.append(f"Page {no} could not be rendered for OCR: {e}")
            dok.halaman.append(Halaman(no=no, words=words, tables=tables))

    if perlu_ocr:
        if not ocr_tersedia():
            if len(perlu_ocr) == len(dok.halaman):
                raise EkstraksiGagal(
                    "This PDF is a scanned image (no text layer). OCR is needed, but the Tesseract OCR engine is not "
                    "installed on the server. Install Tesseract (with the 'ind' language pack) or set TESSERACT_CMD, "
                    "or upload the e-statement PDF/Excel/CSV downloaded from internet banking.", "OCR_UNAVAILABLE")
            dok.peringatan.append(f"{len(perlu_ocr)} page(s) have no text layer and were skipped (OCR engine not installed).")
        elif gambar_ocr:
            import pytesseract
            pytesseract.pytesseract.tesseract_cmd = _cmd_tesseract()
            try:
                bahasa = "ind+eng" if "ind" in pytesseract.get_languages(config="") else "eng"
            except Exception:
                bahasa = "eng"
            skala = OCR_DPI / 72.0
            urut = sorted(gambar_ocr)
            with ThreadPoolExecutor(max_workers=max(1, min(OCR_WORKER_MAKS, len(urut)))) as ex:
                hasil = list(ex.map(lambda n: _ocr_halaman(gambar_ocr[n], skala, bahasa), urut))
            for n, words in zip(urut, hasil):
                h = dok.halaman[n - 1]
                h.words, h.ocr = words, True
            dok.ocr_dipakai = True
            dok.peringatan.append(f"{len(urut)} scanned page(s) were read with OCR — check the flagged lines carefully.")

    if dok.halaman:
        dok.teks_awal = " ".join(w["text"] for w in sorted(dok.halaman[0].words, key=lambda w: (round(w["top"]), w["x0"])))
    return dok


def _sel_excel(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float, Decimal)):
        try:
            return format(Decimal(str(v)).quantize(Decimal("0.01")), "f")
        except InvalidOperation:
            return str(v)
    return str(v).strip()


def _baca_excel(isi: bytes, nama_file: str) -> Dokumen:
    dok = Dokumen(file_type="Excel")
    if nama_file.lower().endswith(".xls"):
        try:
            import pandas as pd
            semua = pd.read_excel(io.BytesIO(isi), sheet_name=None, header=None, dtype=object)
        except ImportError:
            raise EkstraksiGagal("Old Excel .xls files are not supported on this server. Save the file as .xlsx or CSV and upload again.")
        except Exception as e:
            raise EkstraksiGagal(f"The Excel file could not be read. Save it as .xlsx or CSV and try again ({e}).")
        for nama, df in semua.items():
            dok.sheets.append((str(nama), [[_sel_excel(None if (isinstance(v, float) and v != v) else v) for v in row]
                                           for row in df.itertuples(index=False)]))
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(isi), read_only=True, data_only=True)
        except Exception as e:
            raise EkstraksiGagal(f"The Excel file could not be read: {e}")
        try:
            for ws in wb.worksheets:
                dok.sheets.append((ws.title, [[_sel_excel(v) for v in row] for row in ws.iter_rows(values_only=True)]))
        finally:
            wb.close()
    dok.teks_awal = " ".join(" ".join(c for c in r if c) for _, rows in dok.sheets[:1] for r in rows[:25])
    return dok


def _baca_csv(isi: bytes) -> Dokumen:
    dok = Dokumen(file_type="CSV")
    teks = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            teks = isi.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    teks = teks or isi.decode("utf-8", errors="replace")
    contoh = teks[:8192]
    try:
        delim = csv.Sniffer().sniff(contoh, delimiters=",;\t|").delimiter
    except csv.Error:
        delim = max([",", ";", "\t", "|"], key=contoh.count)
    rows = [[(c or "").strip().strip("'") for c in r] for r in csv.reader(io.StringIO(teks), delimiter=delim)]
    dok.sheets.append(("CSV", rows))
    dok.teks_awal = " ".join(" ".join(c for c in r if c) for r in rows[:25])
    return dok


def baca_dokumen(isi: bytes, nama_file: str) -> Dokumen:
    """Baca file -> kata/tabel/sheet. Di-cache per SHA-256 isi (preview lalu import)."""
    kunci = hashlib.sha256(isi).hexdigest()
    sekarang = time.time()
    with _cache_lock:
        ada = _cache_dokumen.get(kunci)
        if ada and sekarang - ada[0] <= _CACHE_TTL_DETIK:
            _cache_dokumen.move_to_end(kunci)
            return ada[1]
    jenis = jenis_file(nama_file)
    if jenis == "PDF":
        dok = _baca_pdf(isi)
    elif jenis == "Excel":
        dok = _baca_excel(isi, nama_file)
    elif jenis == "CSV":
        dok = _baca_csv(isi)
    else:
        raise EkstraksiGagal("Only PDF, Excel (.xlsx/.xls) or CSV bank statements are supported.", "VALIDATION_ERROR")
    with _cache_lock:
        _cache_dokumen[kunci] = (sekarang, dok)
        _cache_dokumen.move_to_end(kunci)
        while len(_cache_dokumen) > _CACHE_MAKS:
            _cache_dokumen.popitem(last=False)
    return dok


# ============================================================
# 4) Dari dokumen -> deretan baris sel (per jalur)
# ============================================================

@dataclass
class BarisSel:
    cells: List[str]
    halaman: Optional[int] = None
    conf: Optional[float] = None       # confidence OCR terendah token angka di baris ini
    putus: bool = False                # jarak vertikal besar -> bukan sambungan baris sebelumnya
    teks_penuh: str = ""
    jarak: Optional[float] = None      # jalur words: celah vertikal ke baris visual sebelumnya (None = baris pertama halaman)
    tinggi: float = 0.0                # jalur words: tinggi baris median halaman


@dataclass
class TabelKandidat:
    metode: str
    headers: List[str]
    roles_terdeteksi: List[Optional[str]]
    baris: List[BarisSel]
    sheet_name: Optional[str] = None
    header_row_index: int = 1
    ringkasan_teks: List[str] = field(default_factory=list)   # baris di luar tabel (untuk ringkasan saldo/total)
    teks_atas: str = ""                # jalur table: teks baris di atas header (judul dokumen)

    @property
    def signature(self) -> str:
        return signature_header(self.headers)


def _baris_visual(words: List[Dict[str, Any]]) -> List[List[Dict[str, Any]]]:
    if not words:
        return []
    tinggi = sorted(w["bottom"] - w["top"] for w in words)
    tol = max(1.5, tinggi[len(tinggi) // 2] * 0.45)
    hasil: List[List[Dict[str, Any]]] = []
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        if hasil and abs(w["top"] - hasil[-1][0]["top"]) <= tol:
            hasil[-1].append(w)
        else:
            hasil.append([w])
    return [sorted(b, key=lambda w: w["x0"]) for b in hasil]


def _sel_header_dari_baris(baris: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Kata-kata 1 baris header -> sel header (kata berdekatan digabung, mis. 'TRANSACTION DATE')."""
    if not baris:
        return []
    tinggi = sorted(w["bottom"] - w["top"] for w in baris)[len(baris) // 2]
    batas = max(3.5, tinggi * 0.6)
    sel: List[Dict[str, Any]] = []
    for w in baris:
        if sel and w["x0"] - sel[-1]["x1"] <= batas:
            s = sel[-1]
            s["text"] += " " + w["text"]
            s["x1"] = max(s["x1"], w["x1"])
            s["kata"].append(w)
        else:
            sel.append({"text": w["text"], "x0": w["x0"], "x1": w["x1"], "kata": [w]})
    # Sel gabungan yang ternyata berisi >= 2 peran berbeda ("DEBET KREDIT") dipecah lagi per kata.
    hasil: List[Dict[str, Any]] = []
    for s in sel:
        peran_kata = {peran_header(w["text"]) for w in s["kata"]} - {None, "ignore"}
        if len(s["kata"]) > 1 and len(peran_kata) >= 2 and peran_header(s["text"]) in (None, "ignore", "amount", "dc_flag") \
                and not ({"debit", "credit"} <= peran_kata and peran_header(s["text"]) == "dc_flag"):
            hasil.extend({"text": w["text"], "x0": w["x0"], "x1": w["x1"], "kata": [w]} for w in s["kata"])
        else:
            hasil.append(s)
    return hasil


def _cari_header_halaman(lines: List[List[Dict[str, Any]]]) -> Optional[Tuple[int, int, List[Dict[str, Any]]]]:
    """(indeks baris header, indeks baris terakhir header, sel header) terbaik di 1 halaman."""
    terbaik = None
    for i, baris in enumerate(lines):
        if not any(peran_header(w["text"]) for w in baris):
            continue
        sel = _sel_header_dari_baris(baris)
        skor = _skor_header([peran_header(s["text"]) for s in sel])
        akhir = i
        # Header 2 baris: baris tepat di bawahnya tanpa angka/tanggal digabung ke sel yang bertumpuk.
        if i + 1 < len(lines):
            bawah = lines[i + 1]
            tinggi = max(w["bottom"] - w["top"] for w in baris)
            jarak = bawah[0]["top"] - max(w["bottom"] for w in baris)
            if jarak <= tinggi * 0.9 and not any(mirip_angka(w["text"]) or pecah_tanggal(w["text"]) for w in bawah):
                sel2 = [dict(s, kata=list(s["kata"])) for s in sel]
                for w in bawah:
                    tumpuk = [s for s in sel2 if w["x0"] < s["x1"] + 2 and w["x1"] > s["x0"] - 2]
                    if tumpuk:
                        s = tumpuk[0]
                        s["text"] += " " + w["text"]
                        s["x0"], s["x1"] = min(s["x0"], w["x0"]), max(s["x1"], w["x1"])
                    else:
                        sel2.append({"text": w["text"], "x0": w["x0"], "x1": w["x1"], "kata": [w]})
                sel2.sort(key=lambda s: s["x0"])
                skor2 = _skor_header([peran_header(s["text"]) for s in sel2])
                if skor2 > skor:
                    sel, skor, akhir = sel2, skor2, i + 1
        if skor > 0 and (terbaik is None or skor > terbaik[0]):
            terbaik = (skor, i, akhir, sel)
    if terbaik is None:
        return None
    return terbaik[1], terbaik[2], terbaik[3]


def _penanda_arah(teks: str) -> bool:
    return _norm(teks) in (_FLAG_DEBIT | _FLAG_KREDIT)


_RE_BAGIAN_TANGGAL = re.compile(
    rf"^(\d{{1,4}}[/.,-]?|\d{{1,2}}[/.-]\d{{1,2}}([/.-]\d{{2,4}})?|\d{{4}}[/.-]\d{{1,2}}[/.-]\d{{1,2}}|\d{{1,2}}:\d{{2}}(:\d{{2}})?|({_POLA_BULAN})[,.]?)$",
    re.I,
)


def _bagian_tanggal(teks: str, isi_sel: List[str]) -> bool:
    """Token ini bagian dari tanggal (mis. '01/04', atau '01' 'Apr' '2025' terpisah)?
    Teks lain yang kebetulan jatuh di zona kolom tanggal (keterangan yang mulai
    lebih kiri dari judul kolomnya, mis. BCA) dikembalikan ke keterangan."""
    return len(isi_sel) < 4 and bool(_RE_BAGIAN_TANGGAL.match(teks.strip()))


def _jalur_words(dok: Dokumen, metode: str) -> Optional[TabelKandidat]:
    """Jalur posisi kata (juga dipakai untuk kata hasil OCR)."""
    pakai_ocr = metode == "ocr"
    kolom: Optional[List[Dict[str, Any]]] = None
    roles: List[Optional[str]] = []
    headers: List[str] = []
    baris_out: List[BarisSel] = []
    ringkasan: List[str] = []
    ada_halaman = False

    for h in dok.halaman:
        if h.ocr != pakai_ocr or not h.words:
            continue
        ada_halaman = True
        lines = _baris_visual(h.words)
        hdr = _cari_header_halaman(lines)
        mulai = 0
        if hdr is not None:
            i_hdr, i_akhir, sel = hdr
            if kolom is None:
                kolom, headers = sel, [s["text"] for s in sel]
                roles = [peran_header(s["text"]) for s in sel]
            elif [_norm(s["text"]) for s in sel] == [_norm(x) for x in headers]:
                kolom = sel  # posisi kolom halaman ini (bisa sedikit bergeser antar halaman)
            ringkasan.extend(" ".join(w["text"] for w in b) for b in lines[:i_hdr])
            mulai = i_akhir + 1
        elif kolom is None:
            ringkasan.extend(" ".join(w["text"] for w in b) for b in lines)
            continue

        tinggi = sorted(w["bottom"] - w["top"] for w in h.words)[len(h.words) // 2]
        # Teks (rata kiri) -> batas di tengah celah antar judul kolom. Angka (rata kanan,
        # sering berakhir tepat sebelum judul kolom berikutnya) -> kolom = judul terakhir
        # yang dimulai sebelum tepi kanan angka itu.
        batas = [(kolom[i]["x1"] + kolom[i + 1]["x0"]) / 2 for i in range(len(kolom) - 1)]
        awal_kolom = [k["x0"] for k in kolom]
        peran_final = _rapikan_peran(roles)
        idx_teks = [i for i, r in enumerate(peran_final) if r not in ROLE_ANGKA and r != "dc_flag"]
        idx_desk = next((i for i, r in enumerate(peran_final) if r == "description"), idx_teks[0] if idx_teks else 0)
        idx_tgl = {i for i, r in enumerate(peran_final) if r == "date"}
        bawah_sebelumnya: Optional[float] = None
        for b in lines[mulai:]:
            teks_baris = " ".join(w["text"] for w in b)
            if re.search(r"\b(BERSAMBUNG|CONTINUED|HALAMAN BERIKUT|NEXT PAGE)\b", teks_baris.upper()):
                ringkasan.append(teks_baris)
                break
            cells: List[List[str]] = [[] for _ in kolom]
            conf_angka: List[float] = []
            terakhir_angka: Optional[Tuple[int, float]] = None
            terakhir_teks: Optional[Tuple[int, float]] = None
            for w in b:
                t = w["text"]
                # Kata yang menempel ke kata keterangan sebelumnya (jarak 1 spasi) tetap 1 frasa,
                # walau keterangannya melebar melewati batas kolom berikutnya (mis. Mandiri
                # "MCM InhouseTrf DARI SITI" yang menjorok ke zona kolom Reference No.).
                if terakhir_teks is not None and w["x0"] - terakhir_teks[1] <= max(3.0, tinggi * 0.5) \
                        and peran_final[terakhir_teks[0]] in ("description", "reference") \
                        and peran_final[bisect_right(batas, w["x0"])] in ("description", "reference"):
                    i = terakhir_teks[0]
                    cells[i].append(t)
                    terakhir_teks = (i, w["x1"])
                    continue
                if mirip_angka(t):
                    i = max(0, bisect_right(awal_kolom, w["x1"] - 0.5) - 1)
                    terakhir_angka = (i, w["x1"])
                    if w.get("conf") is not None:
                        conf_angka.append(w["conf"])
                elif _penanda_arah(t) and terakhir_angka is not None and w["x0"] - terakhir_angka[1] < 25 \
                        and peran_final[terakhir_angka[0]] in ROLE_ANGKA:
                    i = terakhir_angka[0]
                else:
                    i = bisect_right(batas, w["x0"])
                    # Kolom referensi rata kiri -> isinya mulai di judul kolomnya. Kata yang mulai
                    # jelas sebelum judul itu masih ekor keterangan (mis. Mandiri "MARIETTA/KKG  DR").
                    if 0 < i and peran_final[i] == "reference" and peran_final[i - 1] == "description" \
                            and w["x0"] < awal_kolom[i] - tinggi:
                        i -= 1
                    if peran_final[i] in ROLE_ANGKA or (peran_final[i] == "dc_flag" and not _penanda_arah(t)):
                        kiri = [j for j in idx_teks if j <= i]
                        i = kiri[-1] if kiri else idx_desk
                    if i in idx_tgl and not (_bagian_tanggal(t, cells[i]) and not any(cells[j] for j in range(i + 1, len(cells)))):
                        i = idx_desk
                if i in idx_tgl and mirip_angka(t) and not (_bagian_tanggal(t, cells[i])
                                                            and not any(cells[j] for j in range(i + 1, len(cells)))):
                    i = idx_desk
                cells[i].append(t)
                terakhir_teks = (i, w["x1"]) if peran_final[i] in ("description", "reference") else None
            jarak = None if bawah_sebelumnya is None else b[0]["top"] - bawah_sebelumnya
            putus = jarak is not None and jarak > tinggi * 3.2
            bawah_sebelumnya = max(w["bottom"] for w in b)
            baris_out.append(BarisSel(cells=[" ".join(c) for c in cells], halaman=h.no,
                                      conf=min(conf_angka) if conf_angka else None, putus=putus, teks_penuh=teks_baris,
                                      jarak=jarak, tinggi=tinggi))
    if not ada_halaman or kolom is None:
        return None
    peran_final = _rapikan_peran(roles)
    if _sel_rata_tengah(baris_out, peran_final):
        baris_out = _gabung_blok(baris_out)
    return TabelKandidat(metode=metode, headers=headers, roles_terdeteksi=roles, baris=baris_out, ringkasan_teks=ringkasan)


def _sel_rata_tengah(baris: List[BarisSel], roles: List[str]) -> bool:
    """Tabel tanpa garis yang selnya rata-tengah vertikal (mis. Mandiri Kopra): tanggal,
    keterangan & nominal 1 transaksi tercetak di baris visual BERBEDA (tanggal di atas,
    nominal di tengah, sisa keterangan di bawah). Ciri: mayoritas baris bertanggal tidak
    memuat nominal sama sekali."""
    idx_angka = [i for i, r in enumerate(roles) if r in ROLE_ANGKA]
    bertanggal = [b for b in baris if pecah_tanggal(_ambil(b.cells, roles, "date"))]
    if len(bertanggal) < 3 or not idx_angka:
        return False
    dengan_angka = sum(1 for b in bertanggal if any(b.cells[i] for i in idx_angka if i < len(b.cells)))
    return dengan_angka / len(bertanggal) < 0.5


def _gabung_blok(baris: List[BarisSel]) -> List[BarisSel]:
    """Gabungkan baris visual jadi 1 baris per transaksi: blok = baris-baris berdekatan
    (celah < ~1x tinggi baris); antar transaksi dipisah celah vertikal yang lebar."""
    hasil: List[BarisSel] = []
    blok: List[BarisSel] = []

    def _tutup():
        if not blok:
            return
        lebar = max(len(b.cells) for b in blok)
        cells = [" ".join(b.cells[i] for b in blok if i < len(b.cells) and b.cells[i]).strip() for i in range(lebar)]
        conf = [b.conf for b in blok if b.conf is not None]
        hasil.append(BarisSel(cells=cells, halaman=blok[0].halaman, conf=min(conf) if conf else None, putus=True,
                              teks_penuh=" ".join(b.teks_penuh for b in blok)))
        blok.clear()

    for b in baris:
        if blok and (b.halaman != blok[-1].halaman or b.jarak is None or b.jarak > max(b.tinggi, 1.0) * 1.2):
            _tutup()
        blok.append(b)
    _tutup()
    return hasil


def _cari_header_tabel(rows: List[List[str]], maks: int = 40) -> Optional[Tuple[int, List[Optional[str]]]]:
    terbaik = None
    for i, r in enumerate(rows[:maks]):
        roles = [peran_header(c) for c in r]
        skor = _skor_header(roles)
        if skor > 0 and (terbaik is None or skor > terbaik[0]):
            terbaik = (skor, i, roles)
    return None if terbaik is None else (terbaik[1], terbaik[2])


def _jalur_tabel(rows_per_bagian: List[Tuple[Optional[str], Optional[int], List[List[str]]]], metode: str) -> Optional[TabelKandidat]:
    """Jalur grid (tabel PDF) / table (sheet Excel, CSV). rows_per_bagian: (sheet, halaman, rows)."""
    hasil: Optional[TabelKandidat] = None
    lebar = 0
    for sheet, halaman, rows in rows_per_bagian:
        rows = [[re.sub(r"\s*\n\s*", " ", c or "").strip() for c in r] for r in rows]
        hdr = _cari_header_tabel(rows)
        if hdr is not None:
            i_hdr, roles = hdr
            if hasil is None:
                # Tabel grid yang 1 selnya berisi banyak tanggal (seluruh badan tabel dalam 1 sel) tidak bisa dipakai.
                i_tgl = roles.index("date")
                isi_tgl = [r[i_tgl] for r in rows[i_hdr + 1:i_hdr + 6] if i_tgl < len(r)]
                if metode == "grid" and any(len(re.findall(r"\d{1,2}[/.-]\d{1,2}", s)) > 1 for s in isi_tgl):
                    continue
                hasil = TabelKandidat(metode=metode, headers=list(rows[i_hdr]), roles_terdeteksi=roles, baris=[],
                                      sheet_name=sheet, header_row_index=i_hdr + 1)
                lebar = len(rows[i_hdr])
                hasil.ringkasan_teks.extend(" ".join(c for c in r if c) for r in rows[:i_hdr])
                hasil.teks_atas = " ".join(" ".join(c for c in r if c) for r in rows[:i_hdr])
            elif [_norm(c) for c in rows[i_hdr]] != [_norm(c) for c in hasil.headers]:
                if metode == "table":
                    break  # sheet lain dengan header berbeda -> bukan tabel yang sama
                continue
            badan = rows[i_hdr + 1:]
        elif hasil is not None and metode == "grid" and rows and len(rows[0]) == lebar:
            badan = rows  # tabel lanjutan di halaman berikutnya tanpa header
        else:
            continue
        kosong = [i for i, h in enumerate(hasil.headers) if not (h or "").strip()]
        for r in badan:
            r = (r + [""] * lebar)[:lebar]
            hasil.baris.append(BarisSel(cells=r, halaman=halaman, teks_penuh=" ".join(c for c in r if c)))
            # Panel ringkasan di samping tabel (kolom tanpa judul), mis. Laporan Kas Kasir:
            # "Saldo Awal Juni 2026 | 1274590096" di kolom J-K sejajar baris data.
            if metode == "table" and kosong:
                samping = " ".join(r[i] for i in kosong if r[i].strip())
                if samping:
                    hasil.ringkasan_teks.append(samping)
        if metode == "table":
            break  # 1 sheet cukup
    return hasil


def kandidat_tabel(dok: Dokumen) -> List[TabelKandidat]:
    """Semua jalur yang berhasil menemukan header tabel mutasi (urut: grid, words, ocr / table)."""
    hasil: List[TabelKandidat] = []
    if dok.file_type == "PDF":
        # Ringkasan resmi (saldo awal/akhir, total) bisa di mana saja di halaman, di luar tabel:
        # semua baris teks yang tidak diawali tanggal ikut dibaca sebagai kandidat ringkasan.
        teks_lepas = [" ".join(w["text"] for w in b) for h in dok.halaman for b in _baris_visual(h.words)]
        teks_lepas = [t for t in teks_lepas if not pecah_tanggal(t)]
        bagian = [(None, h.no, t) for h in dok.halaman if not h.ocr for t in h.tables]
        if bagian:
            k = _jalur_tabel(bagian, "grid")
            if k and k.baris:
                k.ringkasan_teks = teks_lepas
                hasil.append(k)
        for metode in ("words", "ocr"):
            k = _jalur_words(dok, metode)
            if k and k.baris:
                k.ringkasan_teks = teks_lepas
                hasil.append(k)
    else:
        k = _jalur_tabel([(nama, None, rows) for nama, rows in dok.sheets], "table")
        if k and k.baris:
            hasil.append(k)
    return hasil


# ============================================================
# 5) Rakit transaksi dari baris sel
# ============================================================

_RE_RINGKASAN = [
    ("opening", r"SALDO AWAL|OPENING BALANCE|BEGINNING BALANCE|SALDO SEBELUMNYA|PREVIOUS BALANCE|BALANCE B/F|BALANCE BROUGHT FORWARD"),
    ("closing", r"SALDO AKHIR|CLOSING BALANCE|ENDING BALANCE|SALDO PENUTUP|BALANCE C/F|BALANCE CARRIED FORWARD"),
    ("count_credit", r"NO\.? OF CREDITS?|JUMLAH TRANSAKSI KREDIT|JML\.? KREDIT"),
    ("count_debit", r"NO\.? OF DEBITS?|JUMLAH TRANSAKSI DEBET|JML\.? DEBET"),
    ("total_credit", r"MUTASI CR|MUTASI KREDIT|TOTAL MUTASI KREDIT|TOTAL TRANSAKSI KREDIT|TOTAL CREDIT TRANSACTIONS?|TOTAL AMOUNT CREDITED|"
                     r"TOTAL KREDIT|TOTAL CREDITS?|JUMLAH KREDIT|TOTAL SETORAN|TOTAL DEPOSITS?|TOTAL (?:UANG|KAS) MASUK"),
    ("total_debit", r"MUTASI DB|MUTASI DEBET|MUTASI DEBIT|TOTAL MUTASI DEBET|TOTAL TRANSAKSI DEBET|TOTAL DEBIT TRANSACTIONS?|TOTAL AMOUNT DEBITED|"
                    r"TOTAL DEBET|TOTAL DEBITS?|JUMLAH DEBET|TOTAL PENARIKAN|TOTAL WITHDRAWALS?|TOTAL (?:UANG|KAS) KELUAR"),
]
_RE_RINGKASAN_GABUNG = re.compile("|".join(f"(?P<{k}>{p})" for k, p in _RE_RINGKASAN))
_RE_TOKEN_NILAI = re.compile(r"^\(?-?[\d.,]*\d[\d.,]*\)?-?$")


def _baca_ringkasan_kolom(lines: List[str], fmt: str, ringkasan: Dict[str, Any]) -> None:
    """Ringkasan berbentuk blok kolom: 1 baris berisi >= 2 label tanpa angka, nilainya di
    baris (1-3 baris) di bawahnya dengan jumlah angka yang sama, urutan kiri-ke-kanan.
      BRI    : Saldo Awal | Total Transaksi Debet | Total Transaksi Kredit | Saldo Akhir
      Mandiri: Opening Balance | No. of Debit | Total Amount Debited"""
    norm = [re.sub(r"\s+", " ", (t or "").upper()) for t in lines]
    for i, t in enumerate(norm):
        label = [m.lastgroup for m in _RE_RINGKASAN_GABUNG.finditer(t)]
        if len(label) < 2 or _RE_ANGKA_BEBAS.search(t):
            continue
        for j in range(i + 1, min(i + 4, len(norm))):
            nilai = [x for x in norm[j].split() if _RE_TOKEN_NILAI.match(x)]
            if _RE_RINGKASAN_GABUNG.search(norm[j]) or len(nilai) != len(label):
                continue
            for kunci, x in zip(label, nilai):
                if kunci.startswith("count"):
                    if x.isdigit():
                        ringkasan.setdefault(kunci, int(x))
                    continue
                v, arah = parse_angka(x, fmt)
                if v is not None:
                    ringkasan.setdefault(kunci, -v if kunci in ("opening", "closing") and arah == "debit" else v)
            break


def _baca_ringkasan(teks: str, fmt: str, ringkasan: Dict[str, Any]) -> bool:
    """Isi ringkasan resmi (saldo awal/akhir, total & jumlah baris debit/kredit) dari 1 baris teks."""
    t = re.sub(r"\s+", " ", (teks or "").upper())
    cocok = list(_RE_RINGKASAN_GABUNG.finditer(t))
    if not cocok:
        return False
    ketemu = False
    for n, m in enumerate(cocok):
        segmen = t[m.end(): cocok[n + 1].start() if n + 1 < len(cocok) else len(t)]
        am = _RE_ANGKA_BEBAS.search(segmen)
        if not am:
            continue
        nilai, arah = parse_angka(am.group(0) + (" DB" if re.match(r"\s*(DB|DR)\b", segmen[am.end():]) else ""), fmt)
        if nilai is None:
            continue
        kunci = m.lastgroup
        if kunci.startswith("count"):
            continue  # jumlah baris inline dibaca dari angka setelah total (di bawah) / blok kolom
        if kunci in ("opening", "closing") and arah == "debit":
            nilai = -nilai
        ringkasan.setdefault(kunci, nilai)
        jumlah = re.match(r"\s*(?:DB|CR)?\s*(\d{1,5})\b(?![.,]\d)", segmen[am.end():])
        if jumlah and kunci in ("total_debit", "total_credit"):
            ringkasan.setdefault(kunci.replace("total", "count"), int(jumlah.group(1)))
        ketemu = True
    return ketemu


@dataclass
class Aturan:
    roles: List[str]
    amount_mode: str = "single"            # split / single
    default_direction: Optional[str] = None  # credit / debit / None (= dari saldo)
    number_format: str = "en"
    date_order: str = "dmy"

    def ke_dict(self, headers: List[str], metode: str, bank: Optional[str]) -> Dict[str, Any]:
        return {
            "columns": [{"header": h, "role": r} for h, r in zip(headers, self.roles)],
            "amount_mode": self.amount_mode, "default_direction": self.default_direction,
            "number_format": self.number_format, "date_order": self.date_order, "method": metode, "bank_name": bank,
        }


def aturan_dari_dict(d: Dict[str, Any], headers: List[str]) -> Optional[Aturan]:
    kolom = list(d.get("columns") or [])
    # Selisih jumlah kolom hanya boleh berupa kolom kosong di ujung kanan (lihat signature_header).
    if len(kolom) < len(headers) and not any(_norm(h) for h in headers[len(kolom):]):
        kolom += [{"header": h, "role": "ignore"} for h in headers[len(kolom):]]
    elif len(kolom) > len(headers) and not any(_norm(c.get("header")) for c in kolom[len(headers):]):
        kolom = kolom[:len(headers)]
    if len(kolom) != len(headers):
        return None
    roles = [(c.get("role") if c.get("role") in ROLES else "ignore") for c in kolom]
    return Aturan(roles=roles, amount_mode=d.get("amount_mode") if d.get("amount_mode") in ("split", "single") else
                  ("split" if "debit" in roles or "credit" in roles else "single"),
                  default_direction=d.get("default_direction") if d.get("default_direction") in ("credit", "debit") else None,
                  number_format="id" if d.get("number_format") == "id" else "en",
                  date_order=d.get("date_order") if d.get("date_order") in ("dmy", "mdy") else "dmy")


def _ambil(cells: List[str], roles: List[str], role: str, semua: bool = False) -> str:
    isi = [c for c, r in zip(cells, roles) if r == role and c]
    if not isi:
        return ""
    return " ".join(isi) if semua else isi[0]


def aturan_otomatis(tabel: TabelKandidat) -> Aturan:
    roles = _rapikan_peran(tabel.roles_terdeteksi)
    if "amount" in roles and ("debit" in roles) != ("credit" in roles):
        # 1 kolom debit/kredit saja + kolom mutasi -> anggap kolom tunggal
        roles = ["amount" if r in ("debit", "credit") else r for r in roles]
        roles = _rapikan_peran(roles)
    mode = "split" if ("debit" in roles and "credit" in roles) else "single"
    if mode == "single" and "amount" not in roles:
        roles = ["amount" if r in ("debit", "credit") else r for r in roles]
    sampel_angka = [c for b in tabel.baris for c, r in zip(b.cells, roles) if r in ROLE_ANGKA and c]
    fmt = deteksi_format_angka(sampel_angka)
    sampel_tgl = [_ambil(b.cells, roles, "date") for b in tabel.baris[:400]]
    urutan = deteksi_urutan_tanggal(sampel_tgl)
    arah_default = None
    if mode == "single":
        db = cr = 0
        for b in tabel.baris:
            teks = (_ambil(b.cells, roles, "amount") + " " + _ambil(b.cells, roles, "dc_flag")).strip()
            if not teks:
                continue
            _, arah = parse_angka(_ambil(b.cells, roles, "amount"), fmt)
            flag = _norm(_ambil(b.cells, roles, "dc_flag"))
            if flag in _FLAG_DEBIT or arah == "debit":
                db += 1
            elif flag in _FLAG_KREDIT or arah == "credit":
                cr += 1
        if db and not cr:
            arah_default = "credit"
        elif cr and not db:
            arah_default = "debit"
    return Aturan(roles=roles, amount_mode=mode, default_direction=arah_default, number_format=fmt, date_order=urutan)


@dataclass
class Baris:
    tanggal: Optional[date]
    keterangan: str
    debit: Decimal
    credit: Decimal
    saldo: Optional[Decimal]
    halaman: Optional[int]
    confidence: str = "medium"
    catatan: List[str] = field(default_factory=list)
    conf_ocr: Optional[float] = None
    arah_dari_saldo: bool = False

    def ke_dict(self) -> Dict[str, Any]:
        return {"date": self.tanggal.isoformat() if self.tanggal else None, "description": self.keterangan,
                "debit": float(self.debit), "credit": float(self.credit),
                "balance": None if self.saldo is None else float(self.saldo), "page": self.halaman,
                "confidence": self.confidence, "notes": self.catatan}


def _tanggal_final(p: Tuple[int, int, Optional[int], str], urutan: str, periode: Tuple[Optional[date], Optional[date]],
                   tahun_jalan: Optional[int], bulan_lalu: Optional[int]) -> Tuple[Optional[date], Optional[int]]:
    a, b, th, jenis = p
    hari, bulan = (a, b) if jenis == "dm" or urutan == "dmy" else (b, a)
    if th is None:
        mulai, akhir = periode
        if mulai and akhir and mulai.year != akhir.year:
            th = mulai.year if bulan >= mulai.month else akhir.year
        elif mulai:
            th = mulai.year
        elif tahun_jalan:
            th = tahun_jalan + (1 if bulan_lalu and bulan < bulan_lalu and bulan_lalu - bulan > 6 else 0)
        else:
            th = date.today().year
    try:
        return date(th, bulan, hari), th
    except ValueError:
        return None, th


_RE_AWAL_BARIS = re.compile(r"^(SALDO AWAL|OPENING BALANCE|BEGINNING BALANCE|SALDO SEBELUMNYA|PREVIOUS BALANCE|BALANCE B/F|BALANCE BROUGHT FORWARD)")
_RE_SALDO_PINDAHAN = re.compile(r"^(MODAL AWAL BULAN|SALDO BULAN LALU|SALDO PINDAHAN)|SALDO AKHIR .*(BULAN LALU|SEBELUMNYA)")


def rakit_baris(tabel: TabelKandidat, aturan: Aturan, periode: Tuple[Optional[date], Optional[date]]
                ) -> Tuple[List[Baris], Dict[str, Any], List[str]]:
    roles, fmt = aturan.roles, aturan.number_format
    ringkasan: Dict[str, Any] = {}
    peringatan: List[str] = []
    _baca_ringkasan_kolom(tabel.ringkasan_teks, fmt, ringkasan)
    for t in tabel.ringkasan_teks:
        _baca_ringkasan(t, fmt, ringkasan)

    mentah: List[Dict[str, Any]] = []
    sekarang: Optional[Dict[str, Any]] = None
    halaman_sekarang = None
    for b in tabel.baris:
        if b.halaman != halaman_sekarang:
            halaman_sekarang, sekarang = b.halaman, None
        if not any(c.strip() for c in b.cells):
            continue
        teks_tgl = _ambil(b.cells, roles, "date")
        p = pecah_tanggal(teks_tgl)
        if p is None:
            if _baca_ringkasan(b.teks_penuh, fmt, ringkasan):
                sekarang = None
                continue
            if sekarang is None or b.putus:
                sekarang = None
                continue
            # Blok ringkasan (label tanpa angka, nilainya di baris bawah -- mis. BRI "Saldo Awal
            # Total Transaksi Debet ...") langsung setelah transaksi terakhir: bukan sambungan keterangan.
            if len(_RE_RINGKASAN_GABUNG.findall(re.sub(r"\s+", " ", b.teks_penuh.upper()))) >= 2:
                sekarang = None
                continue
            # Baris sambungan: keterangan multi-baris; nominal/saldo yang tercetak di baris berikutnya.
            isi = _ambil(b.cells, roles, "description", semua=True)
            if isi:
                sekarang["desk"].append(isi)
            sekarang["ref"].extend(c for c, r in zip(b.cells, roles) if r == "reference" and c)
            for role in ("amount", "debit", "credit", "balance", "balance_before", "dc_flag"):
                isi = _ambil(b.cells, roles, role)
                if isi and not sekarang[role]:
                    sekarang[role] = isi
            if b.conf is not None:
                sekarang["conf"] = min(sekarang["conf"] or 100.0, b.conf)
            continue
        sekarang = {"p": p, "desk": [_ambil(b.cells, roles, "description", semua=True)],
                    "ref": [c for c, r in zip(b.cells, roles) if r == "reference" and c],
                    "amount": _ambil(b.cells, roles, "amount"), "debit": _ambil(b.cells, roles, "debit"),
                    "credit": _ambil(b.cells, roles, "credit"), "balance": _ambil(b.cells, roles, "balance"),
                    "balance_before": _ambil(b.cells, roles, "balance_before"),
                    "dc_flag": _ambil(b.cells, roles, "dc_flag"), "halaman": b.halaman, "conf": b.conf}
        mentah.append(sekarang)

    hasil: List[Baris] = []
    tahun_jalan = periode[0].year if periode[0] else None
    bulan_lalu: Optional[int] = None
    tanpa_nominal = tanggal_rusak = 0
    for m in mentah:
        tgl, tahun_jalan = _tanggal_final(m["p"], aturan.date_order, periode, tahun_jalan, bulan_lalu)
        # Spasi ganda (kolom lebar tetap di export CSV KlikBCA: "SWITCHING CR TRF        DR 002") dirapatkan.
        keterangan = re.sub(r"\s+", " ", " ".join(x for x in m["desk"] if x)).strip()
        # Per sel referensi: isi yang sudah ada di keterangan tidak diulang (kolom akun laporan kas,
        # REMARK_CUSTOM CSV BRI = potongan DESK_TRAN), "-" / nomor urut "0" dibuang.
        for x in m["ref"]:
            x = " ".join(tok for tok in x.split() if tok not in ("-", "--"))
            if x and not re.fullmatch(r"0+", x) and x.upper() not in keterangan.upper():
                keterangan = f"{keterangan} {x}".strip()
        saldo, arah_saldo = parse_angka(m["balance"], fmt)
        if saldo is not None and arah_saldo == "debit":
            saldo = -saldo
        catatan: List[str] = []
        debit = credit = NOL
        arah_dari_saldo = False
        if aturan.amount_mode == "split":
            d, _ = parse_angka(m["debit"], fmt)
            c, _ = parse_angka(m["credit"], fmt)
            debit, credit = d or NOL, c or NOL
            if m["debit"] and d is None or m["credit"] and c is None:
                catatan.append("Amount could not be read.")
        else:
            nilai, arah = parse_angka(m["amount"], fmt)
            flag = _norm(m["dc_flag"])
            if flag in _FLAG_DEBIT:
                arah = "debit"
            elif flag in _FLAG_KREDIT:
                arah = "credit"
            if m["amount"] and nilai is None:
                catatan.append("Amount could not be read.")
            if nilai is not None:
                if arah is None:
                    arah = aturan.default_direction
                if arah is None:
                    arah_dari_saldo = True
                    arah = "credit"
                if arah == "debit":
                    debit = nilai
                else:
                    credit = nilai
        # Baris saldo bawaan di awal laporan kas ("Modal Awal Bulan 1.274.590.096 | saldo
        # 1.274.590.096"): dicatat sebagai MASUK padahal saldo pindahan bulan lalu, bukan mutasi.
        if not hasil and saldo is not None and credit > 0 and debit == 0 and abs(credit - saldo) <= TOLERANSI_SALDO \
                and (_RE_AWAL_BARIS.search(keterangan.upper()) or _RE_SALDO_PINDAHAN.search(keterangan.upper())
                     or (ringkasan.get("opening") is not None and abs(ringkasan["opening"] - saldo) <= TOLERANSI_SALDO)):
            ringkasan.setdefault("opening_row", saldo)
            continue
        if debit == 0 and credit == 0:
            if _RE_AWAL_BARIS.match(keterangan.upper()) and saldo is not None:
                ringkasan.setdefault("opening_row", saldo)
                continue
            if not catatan:
                tanpa_nominal += 1
                continue
        if not hasil and m["balance_before"]:
            # Saldo sebelum transaksi pertama = saldo awal (jangkar validasi saldo berjalan).
            sebelum, arah_sebelum = parse_angka(m["balance_before"], fmt)
            if sebelum is not None:
                ringkasan.setdefault("opening_row", -sebelum if arah_sebelum == "debit" else sebelum)
        if tgl is None:
            tanggal_rusak += 1
            catatan.append("Date could not be read.")
        else:
            bulan_lalu = tgl.month
        if periode[0] and periode[1] and tgl and not (periode[0] <= tgl <= periode[1]):
            catatan.append("Date is outside the statement period.")
        hasil.append(Baris(tanggal=tgl, keterangan=keterangan[:2000], debit=debit, credit=credit, saldo=saldo,
                           halaman=m["halaman"], catatan=catatan, conf_ocr=m["conf"], arah_dari_saldo=arah_dari_saldo))
    if tanpa_nominal:
        peringatan.append(f"{tanpa_nominal} dated line(s) without any amount were skipped.")
    if tanggal_rusak:
        peringatan.append(f"{tanggal_rusak} line(s) have a date that could not be read.")
    return hasil, ringkasan, peringatan


# ============================================================
# 6) Validasi saldo & total -> confidence per baris
# ============================================================

def validasi(baris: List[Baris], ringkasan: Dict[str, Any]) -> Dict[str, Any]:
    opening = ringkasan.get("opening")
    if opening is None:
        opening = ringkasan.get("opening_row")
    elif ringkasan.get("opening_row") is not None and abs(ringkasan["opening_row"] - opening) > TOLERANSI_SALDO:
        ringkasan["opening_conflict"] = True

    ada_saldo = any(b.saldo is not None for b in baris)
    berjalan: Optional[Decimal] = opening
    segmen: List[Baris] = []
    titik_cek = titik_gagal = 0
    arah_ditebak = 0
    tanpa_jangkar: List[Baris] = []   # baris sebelum saldo pertama, tanpa saldo awal resmi

    def _tandai(seg: List[Baris], ok: bool, selisih: Optional[Decimal] = None):
        for x in seg:
            if ok:
                x.confidence = "high"
            else:
                x.confidence = "low"
                x.catatan.append(f"Running balance does not match the statement (difference {float(selisih or 0):,.2f})."
                                 if selisih is not None else "Running balance does not match the statement.")

    for b in baris:
        # Arah tanpa penanda (mutasi 1 kolom polos): tentukan dari pergerakan saldo kalau bisa.
        if b.arah_dari_saldo and b.saldo is not None and berjalan is not None:
            nilai = b.credit or b.debit
            if abs((b.saldo - berjalan) + nilai) <= TOLERANSI_SALDO:
                b.debit, b.credit = nilai, NOL
            arah_ditebak += 1
        if berjalan is not None:
            berjalan = berjalan + b.credit - b.debit
        segmen.append(b)
        if b.saldo is None:
            continue
        if berjalan is None:
            for x in segmen:
                x.catatan.append("No opening balance on the statement to verify this line against.")
            tanpa_jangkar.extend(segmen)
            berjalan = b.saldo
            segmen = []
            continue
        titik_cek += 1
        selisih = b.saldo - berjalan
        if abs(selisih) <= TOLERANSI_SALDO:
            _tandai(segmen, True)
        else:
            titik_gagal += 1
            _tandai(segmen, False, selisih)
        berjalan = b.saldo  # jangkar ulang: segmen berikutnya diperiksa terpisah
        segmen = []

    closing = ringkasan.get("closing")
    if segmen and closing is not None and berjalan is not None:
        titik_cek += 1
        selisih = closing - berjalan
        if abs(selisih) <= TOLERANSI_SALDO:
            _tandai(segmen, True)
        else:
            titik_gagal += 1
            _tandai(segmen, False, selisih)
        segmen = []

    total_debit = sum((b.debit for b in baris), NOL)
    total_credit = sum((b.credit for b in baris), NOL)
    cek: List[Dict[str, Any]] = []

    def _cek(label: str, resmi: Any, hitung: Any, tol: Decimal = TOLERANSI_SALDO):
        if resmi is None:
            return
        ok = abs(Decimal(str(resmi)) - Decimal(str(hitung))) <= tol
        cek.append({"label": label, "statement": float(resmi), "extracted": float(hitung), "ok": ok})

    _cek("Total debit (money out)", ringkasan.get("total_debit"), total_debit)
    _cek("Total credit (money in)", ringkasan.get("total_credit"), total_credit)
    _cek("Debit line count", ringkasan.get("count_debit"), sum(1 for b in baris if b.debit > 0), Decimal("0"))
    _cek("Credit line count", ringkasan.get("count_credit"), sum(1 for b in baris if b.credit > 0), Decimal("0"))
    if opening is not None and closing is not None:
        _cek("Closing balance (opening + credit - debit)", closing, opening + total_credit - total_debit)
    total_ok = bool(cek) and all(c["ok"] for c in cek)

    id_tanpa_jangkar = {id(x) for x in tanpa_jangkar}
    for b in baris:
        if b.confidence == "high" and b.catatan:
            continue
        if id(b) in id_tanpa_jangkar:
            continue   # tetap medium: wajar, saldo awal memang tidak tercetak
        if not ada_saldo and total_ok and not b.catatan:
            b.confidence = "high"   # tanpa kolom saldo: total resmi cocok persis
        if b.confidence != "low" and any(c.startswith(("Amount could not", "Date could not")) for c in b.catatan):
            b.confidence = "low"
        if b.confidence == "medium" and b.conf_ocr is not None and b.conf_ocr < OCR_CONF_RENDAH:
            b.confidence = "low"
            b.catatan.append(f"OCR is unsure about the amount ({b.conf_ocr:.0f}% confidence).")
        if b.confidence == "medium" and b.arah_dari_saldo:
            b.catatan.append("Direction (in/out) could not be determined; assumed money in.")
            b.confidence = "low"
        if b.confidence == "medium" and not b.catatan:
            b.catatan.append("No balance or statement total available to verify this line.")

    jml = {k: sum(1 for b in baris if b.confidence == k) for k in ("high", "medium", "low")}
    if not baris:
        status = "failed"
    elif (jml["low"] == 0 and jml["medium"] <= len(tanpa_jangkar) and jml["high"] > 0 and (titik_cek > 0 or total_ok)
          and not any(not c["ok"] for c in cek)):
        status = "verified"
    elif jml["high"] > 0:
        status = "partial"
    else:
        status = "unverified"
    return {
        "status": status, "checks": cek,
        "balance_checkpoints": titik_cek, "balance_mismatches": titik_gagal,
        "rows_high": jml["high"], "rows_medium": jml["medium"], "rows_low": jml["low"],
        "verified_pct": round(100.0 * jml["high"] / len(baris), 1) if baris else 0.0,
        "opening_balance": None if opening is None else float(opening),
        # Tanpa ringkasan resmi (export CSV): saldo akhir = saldo berjalan baris terakhir.
        "closing_balance": float(closing) if closing is not None else (
            float(baris[-1].saldo) if baris and baris[-1].saldo is not None else None),
        "total_debit": float(total_debit), "total_credit": float(total_credit),
        "opening_conflict": bool(ringkasan.get("opening_conflict")),
    }


def _skor_hasil(v: Dict[str, Any], n: int) -> float:
    if n == 0:
        return -1e9
    bonus = {"verified": 1000, "partial": 300, "unverified": 0, "failed": -1000}[v["status"]]
    return bonus + v["rows_high"] * 2 + v["rows_medium"] - v["rows_low"] * 3 - v["balance_mismatches"] * 5


# ============================================================
# 7) Titik masuk
# ============================================================

@dataclass
class HasilEkstraksi:
    file_type: str
    metode: str
    headers: List[str]
    signature: str
    aturan: Aturan
    baris: List[Baris]
    validasi: Dict[str, Any]
    peringatan: List[str]
    bank: Optional[str]
    periode: Tuple[Optional[date], Optional[date]]
    sheet_name: Optional[str]
    header_row_index: int
    template: Optional[Dict[str, Any]] = None      # template yang dipakai (cocok signature)
    mapping_rules: Dict[str, Any] = field(default_factory=dict)
    roles_terdeteksi: List[Optional[str]] = field(default_factory=list)
    ocr: bool = False
    no_rekening: Optional[str] = None               # nomor rekening dari header dokumen (untuk mapping bank)
    teks_judul: str = ""                            # teks awal dokumen (untuk kata kunci mapping bank)

    def ringkas(self, sertakan_baris: bool = True) -> Dict[str, Any]:
        return {
            "file_type": self.file_type, "method": self.metode, "bank_name": self.bank,
            "account_number": self.no_rekening,
            "period_start": self.periode[0].isoformat() if self.periode[0] else None,
            "period_end": self.periode[1].isoformat() if self.periode[1] else None,
            "columns": [{"header": h, "role": r, "detected_role": d}
                        for h, r, d in zip(self.headers, self.aturan.roles, self.roles_terdeteksi or [None] * len(self.headers))],
            "amount_mode": self.aturan.amount_mode, "default_direction": self.aturan.default_direction,
            "number_format": self.aturan.number_format, "date_order": self.aturan.date_order,
            "signature": self.signature, "validation": self.validasi, "warnings": self.peringatan,
            "template": None if not self.template else {
                "id": self.template.get("id"), "name": self.template.get("template_name"),
                "detected_by": self.template.get("detected_by"),
            },
            "rows": [b.ke_dict() for b in self.baris] if sertakan_baris else None,
            "row_count": len(self.baris),
        }


_KOLOM_REKENING = {"NOREK", "NO REK", "NO REKENING", "NOMOR REKENING", "REKENING", "ACCOUNT NO", "ACCOUNT NUMBER",
                   "ACC NO"}


def _no_rekening_kolom(tabel: TabelKandidat) -> Optional[str]:
    """Nomor rekening dari kolom tabel (export CSV tanpa judul, mis. BRI kolom NOREK)."""
    for i, h in enumerate(tabel.headers):
        if _norm(h) not in _KOLOM_REKENING:
            continue
        for b in tabel.baris[:20]:
            v = re.sub(r"\D", "", b.cells[i] if i < len(b.cells) else "")
            if 6 <= len(v) <= 20:
                return v
    return None


def _periode_dari_baris(baris: List[Baris]) -> Tuple[Optional[date], Optional[date]]:
    """Dokumen tanpa teks periode (export CSV) -> tanggal transaksi pertama s/d terakhir."""
    tgl = [b.tanggal for b in baris if b.tanggal]
    return (min(tgl), max(tgl)) if tgl else (None, None)


def _teks_semua(dok: Dokumen) -> str:
    """Teks seluruh halaman PDF."""
    return " ".join(w["text"] for h in dok.halaman for w in h.words)


def ekstrak_rekening_koran(isi: bytes, nama_file: str, templates: Optional[List[Dict[str, Any]]] = None,
                           override: Optional[Dict[str, Any]] = None) -> Optional[HasilEkstraksi]:
    """Ekstrak 1 file. templates: template AKTIF klien (dict kolom tabel template),
    override: mapping_rules dari modal review user (didahulukan di atas template).
    None kalau tidak ada tabel mutasi yang dikenali (pemanggil boleh mencoba parser lama)."""
    dok = baca_dokumen(isi, nama_file)
    kandidat = kandidat_tabel(dok)
    if not kandidat:
        return None
    periode = _cari_periode(dok.teks_awal)
    if dok.file_type == "PDF":
        # Nama bank kadang hanya tercetak di halaman terakhir (mis. BRI "Unit Kerja BANK BRI").
        bank_dok = deteksi_nama_bank(dok.teks_awal) or deteksi_nama_bank(_teks_semua(dok))
        no_rek = deteksi_no_rekening(dok.teks_awal)
        judul = dok.teks_awal
    else:
        # Excel/CSV: hanya judul di atas header -- isi tabel/panel samping (mis. "Pindah Dana
        # Ke BCA" di laporan kas) bukan identitas pemilik dokumen.
        judul = " ".join(k.teks_atas for k in kandidat)
        bank_dok, no_rek = deteksi_nama_bank(judul), deteksi_no_rekening(judul)
    per_sig = {t.get("column_signature_hash"): t for t in (templates or []) if t.get("file_type") == dok.file_type}

    if override and override.get("method") in METODE:
        dipilih = [k for k in kandidat if k.metode == override["method"]] or kandidat
    else:
        dipilih = kandidat
    cocok_template = [k for k in dipilih if k.signature in per_sig]
    if cocok_template and not override:
        dipilih = cocok_template

    terbaik: Optional[HasilEkstraksi] = None
    for k in dipilih:
        template = per_sig.get(k.signature)
        aturan = None
        sumber = override or ((template or {}).get("mapping_rules"))
        if sumber:
            aturan = aturan_dari_dict(sumber, k.headers)
        if aturan is None:
            aturan = aturan_otomatis(k)
        baris, ringkasan, peringatan = rakit_baris(k, aturan, periode)
        v = validasi(baris, ringkasan)
        periode_k = periode if periode[0] else _periode_dari_baris(baris)
        # Dokumen tanpa nama bank (mis. laporan kas tunai) -> label dari template yang cocok.
        bank = bank_dok or ((template or {}).get("mapping_rules") or {}).get("bank_name")
        h = HasilEkstraksi(
            file_type=dok.file_type, metode=k.metode, headers=k.headers, signature=k.signature, aturan=aturan,
            baris=baris, validasi=v, peringatan=list(dok.peringatan) + peringatan, bank=bank, periode=periode_k,
            sheet_name=k.sheet_name, header_row_index=k.header_row_index, template=template,
            roles_terdeteksi=k.roles_terdeteksi, ocr=k.metode == "ocr", no_rekening=no_rek or _no_rekening_kolom(k),
            teks_judul=judul[:2000],
        )
        h.mapping_rules = aturan.ke_dict(k.headers, k.metode, bank)
        if terbaik is None or _skor_hasil(h.validasi, len(h.baris)) > _skor_hasil(terbaik.validasi, len(terbaik.baris)):
            terbaik = h
    if terbaik is not None and terbaik.validasi.get("opening_conflict"):
        terbaik.peringatan.append("The opening balance line differs from the statement summary.")
    return terbaik
