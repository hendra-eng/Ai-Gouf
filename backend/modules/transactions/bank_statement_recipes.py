"""
modules/transactions/bank_statement_recipes.py
===============================================
RESEP khusus per bank untuk rekening koran / laporan kas yang TIDAK bisa dibaca
benar oleh parser generik (bank_statement_parser.py): PDF yang butuh posisi kata
(BCA: saldo hanya di baris terakhir tiap hari; Mandiri Kopra: sel rata-tengah
vertikal) dan Excel Kas Kasir (baris 'Modal Awal Bulan' = saldo awal, bukan mutasi).

Dipanggil dari bank_statement_parser.ekstrak_rekening_koran() lewat coba():
  - hanya jalan kalau ada template kandidat yang mapping_rules-nya berisi "recipe"
    (bca_pdf | bri_pdf | bni_pdf | bni_acctinfo_pdf | mandiri_pdf | permata_pdf | kas_kasir_xlsx) DAN header tabel file cocok
    dengan header_columns template itu;
  - mengembalikan dict `ex` berbentuk SAMA dengan hasil parser generik, jadi
    bank_cash_v1.py (preview/upload/mapping/hash/dedup) tidak perlu diubah;
  - kalau tidak ada resep yang cocok / resep gagal -> None -> parser generik jalan
    seperti biasa.

MURNI fungsi (tanpa FastAPI/DB). Diuji terhadap file contoh SAU Agustus 2026
(tests/test_bank_statement_recipes.py): jumlah baris, total debit/kredit, saldo awal
& akhir sama persis dengan ringkasan resmi di dokumen, saldo berjalan 0 selisih.

Konvensi arah (sama dengan frontend): debit = uang KELUAR, credit = uang MASUK.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Callable, Dict, List, Optional, Tuple

TWO = Decimal("0.01")
BULAN_ID = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6, "juli": 7,
    "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12,
}
BULAN_EN = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6, "jul": 7, "aug": 8,
    "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
AMOUNT_RE = re.compile(r"^\d{1,3}(?:,\d{3})*\.\d{2}$")


class ResepError(Exception):
    """Resep cocok tapi gagal membaca file -> pemanggil jatuh balik ke parser generik."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


# ============================================================
# Util umum
# ============================================================

def norm_header(teks: Any) -> str:
    """Normalisasi 1 sel header: huruf kecil, hanya a-z 0-9 (spasi, titik,
    garis miring, dst dibuang) -- 'T ANG G A L' dan 'Tanggal' sama-sama 'tanggal'."""
    return re.sub(r"[^a-z0-9]", "", str(teks or "").lower())




def _dec(s: str) -> Decimal:
    return Decimal(s.replace(",", ""))


def _is_amount(s: str) -> bool:
    return bool(AMOUNT_RE.match(s))



# ============================================================
# Baca PDF -> halaman berisi kata (text, x0, x1, top)
# ============================================================

def _halaman_pdf(content: bytes, hanya_halaman: Optional[int] = None, dedupe: bool = False) -> List[Dict[str, Any]]:
    import pdfplumber

    hasil = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for no, page in enumerate(pdf.pages, start=1):
            try:
                if dedupe:  # huruf tercetak ganda (lapisan teks dobel) -> buang duplikat sebelum jadi kata
                    page = page.dedupe_chars(tolerance=1)
                kata = page.extract_words(x_tolerance=1.5, y_tolerance=2)
            finally:
                page.flush_cache()
            hasil.append({"no": no, "kata": [
                {"text": w["text"], "x0": float(w["x0"]), "x1": float(w["x1"]), "top": float(w["top"])} for w in kata
            ]})
            if hanya_halaman and no >= hanya_halaman:
                break
    return hasil


def _baris(kata: List[Dict[str, Any]], toleransi: float = 3.0) -> List[List[Dict[str, Any]]]:
    urut = sorted(kata, key=lambda w: (w["top"], w["x0"]))
    out: List[List[Dict[str, Any]]] = []
    acuan: Optional[float] = None
    skrg: List[Dict[str, Any]] = []
    for w in urut:
        if acuan is None or w["top"] - acuan > toleransi:
            if skrg:
                out.append(sorted(skrg, key=lambda x: x["x0"]))
            skrg, acuan = [w], w["top"]
        else:
            skrg.append(w)
    if skrg:
        out.append(sorted(skrg, key=lambda x: x["x0"]))
    return out


def _teks(baris: List[Dict[str, Any]]) -> str:
    return " ".join(w["text"] for w in baris)


def _teks_halaman_norm(halaman: Dict[str, Any]) -> str:
    return norm_header(" ".join(_teks(b) for b in _baris(halaman["kata"])))


def _urutan_ada(teks_norm: str, kolom_norm: List[str]) -> bool:
    """Semua kolom header muncul BERURUTAN di teks halaman (toleran terhadap header 2 baris)."""
    pos = 0
    for k in kolom_norm:
        i = teks_norm.find(k, pos)
        if i < 0:
            return False
        pos = i + len(k)
    return True


def _cari_kata(kata: List[Dict[str, Any]], teks: str, x_maks: Optional[float] = None) -> Optional[Dict[str, Any]]:
    for w in sorted(kata, key=lambda w: (w["top"], w["x0"])):
        if w["text"] == teks and (x_maks is None or w["x0"] <= x_maks):
            return w
    return None


# ============================================================
# Hasil standar parser
# ============================================================

def _hasil(bank: str, metode: str) -> Dict[str, Any]:
    return {
        "bank_name": bank, "method": metode, "account_number": None, "period_start": None, "period_end": None,
        "opening_balance": None, "closing_balance": None, "summary": {}, "rows": [], "warnings": [],
    }


def _tambah(h: Dict[str, Any], tgl: date, jam: Optional[str], desk: str, debit: Decimal, kredit: Decimal,
            saldo: Optional[Decimal], ref: Optional[str], halaman: int, kategori: Optional[str] = None) -> None:
    desk = re.sub(r"\s+", " ", desk or "").strip()
    h["rows"].append({
        "line_no": len(h["rows"]) + 1, "date": tgl, "time": jam, "description": desk[:500] or None,
        "debit": debit, "credit": kredit, "balance_after": saldo, "reference": (ref or None),
        "category": kategori, "page_no": halaman, "confidence": None, "confidence_notes": [],
    })


# ============================================================
# BCA -- Rekening Giro (PDF)
# ============================================================

def _parse_bca(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("BCA", "words")
    p1 = _baris(pages[0]["kata"])
    tahun = bulan = None
    for b in p1:
        t = [w["text"] for w in b]
        if "NO." in t and "REKENING" in t:
            angka = [x for x in t if x.isdigit() and len(x) >= 6]
            if angka and not h["account_number"]:
                h["account_number"] = angka[-1]
        if "PERIODE" in t:
            sisa = [x for x in t[t.index("PERIODE") + 1:] if x != ":"]
            if len(sisa) >= 2:
                bulan = BULAN_ID.get(sisa[0].lower())
                tahun = int(sisa[1]) if sisa[1].isdigit() else None
    if not (bulan and tahun):
        raise ResepError("PERIOD_NOT_FOUND", "Baris 'PERIODE : <BULAN> <TAHUN>' tidak ditemukan di halaman 1 rekening BCA.")
    h["period_start"] = date(tahun, bulan, 1)
    h["period_end"] = date(tahun + (bulan == 12), bulan % 12 + 1, 1).fromordinal(date(tahun + (bulan == 12), bulan % 12 + 1, 1).toordinal() - 1)

    for pg in pages:
        baris = _baris(pg["kata"])
        header = next((b for b in baris if b and b[0]["text"] == "TANGGAL" and any(w["text"] == "SALDO" for w in b)), None)
        if not header:
            continue
        x_saldo = next(w["x0"] for w in header if w["text"] == "SALDO") - 5
        x_cbg = next((w["x0"] for w in header if w["text"] == "CBG"), 300) - 6
        x_mut = next((w["x0"] for w in header if w["text"] == "MUTASI"), 380) - 6
        sekarang: Optional[Dict[str, Any]] = None
        for b in baris:
            if b[0]["top"] <= header[0]["top"]:
                continue
            kata0 = b[0]
            if re.match(r"^\d\d/\d\d$", kata0["text"]) and kata0["x0"] < 70:
                if sekarang:
                    _selesai_bca(h, sekarang, tahun)
                    sekarang = None
                badan = [w for w in b[1:]]
                if len(badan) >= 2 and badan[0]["text"] == "SALDO" and badan[1]["text"] == "AWAL":
                    sal = [w for w in badan if _is_amount(w["text"])]
                    if sal:
                        h["opening_balance"] = _dec(sal[-1]["text"])
                    continue
                angka = [w for w in badan if _is_amount(w["text"]) and w["x0"] >= x_mut]
                db = any(w["text"] == "DB" and w["x0"] >= x_mut for w in badan)
                mut = next((w for w in angka if w["x0"] < x_saldo), None)
                sal = next((w for w in angka if w["x0"] >= x_saldo), None)
                if not mut:
                    continue
                desk = [w["text"] for w in badan if w["x0"] < x_cbg]
                cbg = [w["text"] for w in badan if x_cbg <= w["x0"] < x_mut]
                sekarang = {
                    "dd": kata0["text"], "desk": desk, "debit": db, "amt": _dec(mut["text"]),
                    "saldo": _dec(sal["text"]) if sal else None, "ref": " ".join(cbg), "page": pg["no"],
                }
            elif sekarang:
                teks = _teks(b)
                if re.match(r"^(SALDO AWAL|MUTASI (CR|DB)|SALDO AKHIR)\b", teks):
                    # ringkasan halaman terakhir
                    s = h["summary"]
                    ang = [w for w in b if _is_amount(w["text"])]
                    cnt = [w["text"] for w in b if w["text"].isdigit() and len(w["text"]) <= 6]
                    if teks.startswith("MUTASI CR") and ang:
                        s["total_credit"] = _dec(ang[0]["text"]); s["count_credit"] = int(cnt[-1]) if cnt else None
                    elif teks.startswith("MUTASI DB") and ang:
                        s["total_debit"] = _dec(ang[0]["text"]); s["count_debit"] = int(cnt[-1]) if cnt else None
                    elif teks.startswith("SALDO AKHIR") and ang:
                        s["closing_balance"] = _dec(ang[0]["text"])
                    elif teks.startswith("SALDO AWAL") and ang:
                        s["opening_balance"] = _dec(ang[0]["text"])
                    continue
                sekarang["desk"] += [w["text"] for w in b if 70 <= w["x0"] < x_cbg]
        if sekarang:
            _selesai_bca(h, sekarang, tahun)
            sekarang = None
    h["closing_balance"] = h["summary"].get("closing_balance")
    if h["opening_balance"] is None:
        h["opening_balance"] = h["summary"].get("opening_balance")
    return h


def _selesai_bca(h: Dict[str, Any], r: Dict[str, Any], tahun: int) -> None:
    dd, mm = r["dd"].split("/")
    tgl = date(tahun, int(mm), int(dd))
    debit = r["amt"] if r["debit"] else Decimal("0")
    kredit = Decimal("0") if r["debit"] else r["amt"]
    _tambah(h, tgl, None, " ".join(r["desk"]), debit, kredit, r["saldo"], r["ref"], r["page"])


# ============================================================
# BRI -- Britama Bisnis (PDF)
# ============================================================

def _parse_bri(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("BRI", "words")
    for b in _baris(pages[0]["kata"]):
        t = [w["text"] for w in b]
        if "No." in t and "Rekening" in t:
            angka = [x for x in t if x.isdigit() and len(x) >= 8]
            if angka and not h["account_number"]:
                h["account_number"] = angka[0]
        if "Periode" in t and "Transaksi" in t:
            tg = re.findall(r"\d\d/\d\d/\d\d", " ".join(t))
            if len(tg) == 2:
                h["period_start"] = datetime.strptime(tg[0], "%d/%m/%y").date()
                h["period_end"] = datetime.strptime(tg[1], "%d/%m/%y").date()

    for pg in pages:
        baris = _baris(pg["kata"])
        hd = next((b for b in baris if b and b[0]["text"] == "Tanggal" and any(w["text"] == "Uraian" for w in b)), None)
        if hd:
            c = lambda nama: next((w for w in hd if w["text"] == nama), None)
            x_teller = (c("Teller")["x0"] if c("Teller") else 305) - 8
            c_kredit = ((c("Kredit")["x0"] + c("Kredit")["x1"]) / 2) if c("Kredit") else 493
            c_saldo = ((c("Saldo")["x0"] + c("Saldo")["x1"]) / 2) if c("Saldo") else 590
            sekarang: Optional[Dict[str, Any]] = None
            for b in baris:
                if b[0]["top"] <= hd[0]["top"] + 12:
                    continue
                if re.match(r"^\d\d/\d\d/\d\d$", b[0]["text"]) and b[0]["x0"] < 60 and len(b) > 2 and re.match(r"^\d\d:\d\d:\d\d$", b[1]["text"]):
                    if sekarang:
                        _selesai_bri(h, sekarang)
                    amt = [w for w in b if _is_amount(w["text"]) and w["x0"] > x_teller]
                    d = k = s = None
                    for w in amt:
                        if w["x1"] < c_kredit:
                            d = _dec(w["text"])
                        elif w["x1"] < c_saldo:
                            k = _dec(w["text"])
                        else:
                            s = _dec(w["text"])
                    desk = [w["text"] for w in b[2:] if w["x0"] < x_teller]
                    ref = " ".join(w["text"] for w in b[2:] if x_teller <= w["x0"] < x_teller + 60 and not _is_amount(w["text"]))
                    sekarang = {"tgl": b[0]["text"], "jam": b[1]["text"], "desk": desk, "debit": d or Decimal("0"),
                                "kredit": k or Decimal("0"), "saldo": s, "ref": ref, "page": pg["no"]}
                elif sekarang:
                    teks = _teks(b)
                    if teks.startswith(("Created By", "Halaman", "Page ", "LAPORAN", "STATEMENT")) or b[0]["x0"] > x_teller:
                        continue
                    sekarang["desk"] += [w["text"] for w in b if w["x0"] < x_teller]
            if sekarang:
                _selesai_bri(h, sekarang)
        # ringkasan (halaman terakhir): 4 nominal berurutan di bawah header 'Saldo Awal'
        for i, b in enumerate(baris):
            ang = [w for w in b if _is_amount(w["text"])]
            if len(ang) == 4 and i > 0 and _teks(baris[i - 1]).startswith("Opening"):
                s = h["summary"]
                s["opening_balance"], s["total_debit"], s["total_credit"], s["closing_balance"] = (_dec(w["text"]) for w in ang)
    h["opening_balance"] = h["summary"].get("opening_balance")
    h["closing_balance"] = h["summary"].get("closing_balance")
    return h


def _selesai_bri(h: Dict[str, Any], r: Dict[str, Any]) -> None:
    tgl = datetime.strptime(r["tgl"], "%d/%m/%y").date()
    _tambah(h, tgl, r["jam"], " ".join(r["desk"]), r["debit"], r["kredit"], r["saldo"], r["ref"], r["page"])


# ============================================================
# Mandiri -- Kopra Account Statement (PDF)
# ============================================================

def _parse_mandiri(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("Mandiri", "words")
    b1 = _baris(pages[0]["kata"])
    for i, b in enumerate(b1):
        teks = _teks(b)
        if teks.startswith("Account No.") and i + 1 < len(b1):
            nxt = b1[i + 1]
            h["account_number"] = nxt[0]["text"] if nxt and nxt[0]["text"].isdigit() else None
        if teks.startswith("Period") and i + 1 < len(b1):
            m = re.findall(r"(\d\d) (\w{3}) (\d{4})", _teks(b1[i + 1]))
            if len(m) == 2:
                (d1, m1, y1), (d2, m2, y2) = m
                h["period_start"] = date(int(y1), BULAN_EN[m1.lower()], int(d1))
                h["period_end"] = date(int(y2), BULAN_EN[m2.lower()], int(d2))
        ang = [w for w in b if _is_amount(w["text"])]
        ints = [w["text"] for w in b if w["text"].isdigit() and w["x0"] > 150]
        if len(ang) == 2 and len(ints) == 1:
            s = h["summary"]
            if "opening_balance" not in s:
                s["opening_balance"], s["total_debit"], s["count_debit"] = _dec(ang[0]["text"]), _dec(ang[1]["text"]), int(ints[0])
            else:
                s["closing_balance"], s["total_credit"], s["count_credit"] = _dec(ang[0]["text"]), _dec(ang[1]["text"]), int(ints[0])
    h["opening_balance"] = h["summary"].get("opening_balance")
    h["closing_balance"] = h["summary"].get("closing_balance")

    pusat_baris: List[Dict[str, Any]] = []  # baris yang sedang dirakit lintas halaman
    for pg in pages:
        baris = _baris(pg["kata"])
        hd = next((b for b in baris if b and b[0]["text"] == "Posting" and any(w["text"] == "Remark" for w in b)), None)
        if not hd:
            continue
        w_ = lambda nama: next((w for w in hd if w["text"] == nama), None)
        x_remark = (w_("Remark")["x0"] if w_("Remark") else 120) - 4
        x_ref = (w_("Reference")["x0"] if w_("Reference") else 240) - 4
        c_debit = ((w_("Debit")["x0"] + w_("Debit")["x1"]) / 2) if w_("Debit") else 372
        x_debit = (w_("Debit")["x0"] if w_("Debit") else 362) - 45  # batas kanan kolom Reference No.
        c_credit = ((w_("Credit")["x0"] + w_("Credit")["x1"]) / 2) if w_("Credit") else 460
        c_bal = ((w_("Balance")["x0"] + w_("Balance")["x1"]) / 2) if w_("Balance") else 548
        top_hd = hd[0]["top"]
        pusat = []  # (y, debit, credit, balance, row_dict)
        rows_halaman: List[Dict[str, Any]] = []
        for b in baris:
            ang = [w for w in b if _is_amount(w["text"]) and w["x1"] >= c_debit]
            if len(ang) == 3 and b[0]["top"] > top_hd:
                r = {"y": b[0]["top"], "debit": _dec(ang[0]["text"]), "credit": _dec(ang[1]["text"]),
                     "balance": _dec(ang[2]["text"]), "tgl": [], "remark": [], "ref": [], "page": pg["no"]}
                rows_halaman.append(r)
        if not rows_halaman:
            continue

        def terdekat(y: float) -> Dict[str, Any]:
            return min(rows_halaman, key=lambda r: abs(r["y"] - y))

        for b in baris:
            for w in b:
                if w["top"] <= top_hd - 1 and w["x0"] >= x_remark - 5 and w["x0"] < x_ref:
                    # teks tumpahan di ATAS header = lanjutan remark baris terakhir halaman sebelumnya
                    if pusat_baris and w["top"] > 55:
                        pusat_baris[-1]["remark"].append(w)
                    continue
                if w["top"] <= top_hd + 6:
                    continue
                if w["x0"] < x_remark - 5:
                    if re.match(r"^(\d\d|[A-Z][a-z]{2}|\d{4},?|\d\d:\d\d:\d\d)$", w["text"]):
                        terdekat(w["top"])["tgl"].append(w)
                elif w["x0"] < x_ref:
                    terdekat(w["top"])["remark"].append(w)
                elif w["x0"] < x_debit:
                    terdekat(w["top"])["ref"].append(w)
        pusat_baris.extend(rows_halaman)

    for r in pusat_baris:
        tglw = sorted(r["tgl"], key=lambda w: (round(w["top"]), w["x0"]))
        teks_tgl = " ".join(w["text"] for w in tglw)
        m = re.search(r"(\d\d) ([A-Z][a-z]{2}) (\d{4}),? (\d\d:\d\d:\d\d)", teks_tgl)
        if not m:
            h["warnings"].append(f"Mandiri hal {r['page']}: tanggal tidak terbaca ('{teks_tgl}').")
            continue
        tgl = date(int(m.group(3)), BULAN_EN[m.group(2).lower()], int(m.group(1)))
        remark = " ".join(w["text"] for w in sorted(r["remark"], key=lambda w: (w["top"], w["x0"])))
        ref = " ".join(w["text"] for w in sorted(r["ref"], key=lambda w: (w["top"], w["x0"])))
        if ref.strip() == "-":
            ref = ""
        _tambah(h, tgl, m.group(4), remark, r["debit"], r["credit"], r["balance"], ref, r["page"])
    return h


# ============================================================
# Permata -- Transaction Inquiry, Permata eBusiness (PDF)
# ============================================================

_PM_NOTGL = re.compile(r"^(\d+?)(\d{2}-[A-Za-z]{3}-\d{4})$")  # kolom No. menempel ke Post Date: "101-Jun-2026" = No 1 + 01-Jun-2026
_PM_TGL = re.compile(r"^\d{2}-[A-Za-z]{3}-\d{4}$")
_PM_X_DEBIT_MAKS = 740.0   # nominal Debit rata kanan di x1 ~ 720, Credit di x1 ~ 830
_PM_X_DESK_MAKS = 645.0


def _pm_tgl(teks: str) -> date:
    d, b, y = teks.split("-")
    return date(int(y), BULAN_EN[b[:3].lower()], int(d))


def _pm_nominal(b: List[Dict[str, Any]]) -> List[Tuple[str, Decimal]]:
    """Nominal di sisi kanan baris -> [('D'|'K', nilai)] menurut posisi kolom."""
    out: List[Tuple[str, Decimal]] = []
    for w in b:
        if AMOUNT_RE.match(w["text"]) and w["x1"] > _PM_X_DESK_MAKS:
            out.append(("D" if w["x1"] <= _PM_X_DEBIT_MAKS else "K", _dec(w["text"])))
    return out


def _parse_permata(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Saldo TIDAK tercetak per baris: hanya 'Opening Ledger Balance', 'Ledger Balance per <tgl>' (= saldo
    sesudah semua mutasi sebelum tanggal itu), 'Closing Ledger Balance' dan baris 'Total' per tanggal.
    Saldo dipakai sebagai checkpoint di baris terakhir tiap tanggal; Total per tanggal dicocokkan dengan
    jumlah baris hasil baca."""
    h = _hasil("Permata", "words")
    sekarang: Optional[Dict[str, Any]] = None
    awal_grup = 0
    tot_d = tot_k = Decimal("0")
    cp_terakhir = 0
    no_terakhir = 0

    def selesai() -> None:
        nonlocal sekarang, no_terakhir
        r = sekarang
        sekarang = None
        if not r:
            return
        d = r["d"] or Decimal("0")
        k = r["k"] or Decimal("0")
        _tambah(h, r["tgl"], None, " ".join(r["desk"]), d, k, None, r["ref"] or r["cref"] or None, r["page"])
        no_terakhir = max(no_terakhir, r["no"])

    def checkpoint(saldo: Decimal) -> None:
        nonlocal cp_terakhir
        selesai()
        if h["rows"] and len(h["rows"]) > cp_terakhir:
            h["rows"][-1]["balance_after"] = saldo
            cp_terakhir = len(h["rows"])

    for pg in pages:
        mentah = _baris(pg["kata"])
        baris: List[List[Dict[str, Any]]] = []
        for b in mentah:  # label "Total" tercetak ~3pt lebih tinggi dari nominalnya -> sempat terpisah jadi 2 baris
            if baris and len(baris[-1]) == 1 and baris[-1][0]["text"] == "Total" and b[0]["top"] - baris[-1][0]["top"] <= 6:
                baris[-1] = sorted(baris[-1] + b, key=lambda w: w["x0"])
            else:
                baris.append(b)
        mulai = False
        top_header = 0.0
        for b in baris:
            t = [w["text"] for w in b]
            teks = " ".join(t)
            if pg["no"] == 1 and not h["account_number"] and t and t[0] == "Account":
                m = re.search(r"(\d{6,})/", teks)
                if m:
                    h["account_number"] = m.group(1)
            if pg["no"] == 1 and t and t[0] == "Period":
                tg = [x for x in t if _PM_TGL.match(x)]
                if len(tg) == 2:
                    h["period_start"], h["period_end"] = _pm_tgl(tg[0]), _pm_tgl(tg[1])
            if teks.startswith("Ineffective Balance"):  # bagian di bawah closing: bukan mutasi
                break
            if not mulai:
                if "Description" in t and "Debit" in t and "Credit" in t:
                    mulai, top_header = True, b[0]["top"]
                continue
            if b[0]["top"] <= top_header + 10:  # baris ke-2 header ("Code ... Ref No")
                continue
            x0 = b[0]["x0"]
            # --- baris struktur (kiri) ---
            if x0 < 60 and t[0] == "Opening" and len(t) > 3 and t[1] == "Ledger":
                a = _pm_nominal(b)
                if a:
                    h["opening_balance"] = a[-1][1]
                continue
            if x0 < 60 and t[0] == "Closing" and len(t) > 3 and t[1] == "Ledger":
                a = _pm_nominal(b)
                if a:
                    checkpoint(a[-1][1])
                    h["closing_balance"] = a[-1][1]
                continue
            if x0 < 60 and t[0] == "Ledger" and len(t) > 3 and t[1] == "Balance":
                a = _pm_nominal(b)
                if a:
                    checkpoint(a[-1][1])
                continue
            if x0 < 60 and t[0] == "Total":
                selesai()
                a = _pm_nominal(b)
                d = next((v for k, v in a if k == "D"), Decimal("0"))
                k_ = next((v for k, v in a if k == "K"), Decimal("0"))
                grp = h["rows"][awal_grup:]
                sd = sum((r["debit"] for r in grp), Decimal("0"))
                sk = sum((r["credit"] for r in grp), Decimal("0"))
                if abs(sd - d) > Decimal("0.01") or abs(sk - k_) > Decimal("0.01"):
                    tgl = grp[0]["date"] if grp else None
                    h["warnings"].append(f"Total per tanggal {tgl}: dokumen debit {d:,.2f} / kredit {k_:,.2f} vs hasil baca {sd:,.2f} / {sk:,.2f}.")
                tot_d += d
                tot_k += k_
                awal_grup = len(h["rows"])
                continue
            # --- baris transaksi ---
            m = _PM_NOTGL.match(t[0]) if x0 < 115 else None
            if m and len(b) > 1 and _PM_TGL.match(t[1]):
                selesai()
                ref = " ".join(w["text"] for w in b if 340 <= w["x0"] < 445)
                cref = " ".join(w["text"] for w in b if 445 <= w["x0"] < 500 and not AMOUNT_RE.match(w["text"]))
                sekarang = {
                    "no": int(m.group(1)), "tgl": _pm_tgl(m.group(2)), "ref": ref, "cref": cref, "page": pg["no"],
                    "desk": [w["text"] for w in b if w["x0"] >= 500 and w["x1"] <= _PM_X_DESK_MAKS and not AMOUNT_RE.match(w["text"])],
                    "d": None, "k": None,
                }
                for kolom, v in _pm_nominal(b):
                    if kolom == "D" and sekarang["d"] is None:
                        sekarang["d"] = v
                    elif kolom == "K" and sekarang["k"] is None:
                        sekarang["k"] = v
            elif sekarang:  # lanjutan deskripsi / nomor yang terbungkus
                sekarang["cref"] += "".join(w["text"] for w in b if 445 <= w["x0"] < 500)
                sekarang["desk"] += [w["text"] for w in b if w["x0"] >= 500 and w["x1"] <= _PM_X_DESK_MAKS and not AMOUNT_RE.match(w["text"])]
                for kolom, v in _pm_nominal(b):
                    if kolom == "D" and sekarang["d"] is None:
                        sekarang["d"] = v
                    elif kolom == "K" and sekarang["k"] is None:
                        sekarang["k"] = v
    selesai()
    h["summary"] = {"opening_balance": h["opening_balance"], "closing_balance": h["closing_balance"],
                    "total_debit": tot_d, "total_credit": tot_k}
    if no_terakhir and no_terakhir != len(h["rows"]):
        h["warnings"].append(f"Nomor urut terakhir di dokumen {no_terakhir} tetapi {len(h['rows'])} baris terbaca.")
    return h


# ============================================================
# Kas Kasir Finance (Excel)
# ============================================================

KAS_KOLOM = ["tanggal", "nojurnal", "keterangan", "account", "kasmasuk", "kaskeluar", "saldouangfisik"]


def _tgl_excel(v: Any) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _angka(v: Any) -> Decimal:
    if v is None or v == "":
        return Decimal("0")
    return Decimal(str(v))


def _cari_header_kas(rows: List[Tuple[Any, ...]]) -> Optional[Tuple[int, Dict[str, int]]]:
    for i, r in enumerate(rows[:20]):
        n = [norm_header(c) for c in r]
        if all(k in n for k in KAS_KOLOM):
            return i, {k: n.index(k) for k in KAS_KOLOM}
    return None


def _parse_kas_kasir(content: bytes, sheet_name: Optional[str]) -> Dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = None
    for nama in wb.sheetnames:
        if sheet_name and norm_header(nama) == norm_header(sheet_name):
            ws = wb[nama]
    ws = ws or wb[wb.sheetnames[0]]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    found = _cari_header_kas(rows)
    if not found:
        raise ResepError("HEADER_NOT_FOUND", "Header Kas Kasir (Tanggal / NO. JURNAL / KETERANGAN / ACCOUNT / KAS MASUK / KAS KELUAR / SALDO UANG FISIK) tidak ditemukan.")
    hi, idx = found
    h = _hasil("Kas Kasir", "table")
    # panel ringkasan samping (label di kolom J, nilai di kolom K) -- hanya baris di atas/di sekitar header
    for r in rows[hi + 1: hi + 12]:
        if len(r) > 10 and isinstance(r[9], str):
            lab = r[9].lower()
            if lab.startswith("saldo awal"):
                h["summary"]["opening_balance"] = _angka(r[10])
            elif lab.startswith("saldo akhir"):
                h["summary"]["closing_balance"] = _angka(r[10])
            elif lab.startswith("total uang masuk"):
                h["summary"]["total_credit"] = _angka(r[10])
            elif lab.startswith("total uang keluar"):
                h["summary"]["total_debit"] = _angka(r[10])
    tgls = []
    for no, r in enumerate(rows[hi + 1:], start=hi + 2):
        tgl = _tgl_excel(r[idx["tanggal"]]) if len(r) > idx["tanggal"] else None
        if not tgl:
            continue
        ket = str(r[idx["keterangan"]] or "").strip()
        masuk, keluar = _angka(r[idx["kasmasuk"]]), _angka(r[idx["kaskeluar"]])
        saldo = r[idx["saldouangfisik"]]
        saldo = _angka(saldo) if saldo is not None else None
        if h["opening_balance"] is None and norm_header(ket) == "modalawalbulan":
            h["opening_balance"] = masuk  # saldo awal, BUKAN mutasi
            continue
        tgls.append(tgl)
        no_jur = str(r[idx["nojurnal"]] or "").strip()
        acc = str(r[idx["account"]] or "").strip()
        _tambah(h, tgl, None, ket, keluar, masuk, saldo, no_jur or None, no, kategori=acc or None)
    if tgls:
        h["period_start"], h["period_end"] = min(tgls), max(tgls)
    if h["rows"]:
        h["closing_balance"] = next((r["balance_after"] for r in reversed(h["rows"]) if r["balance_after"] is not None), None)
    if h["opening_balance"] is None:
        h["opening_balance"] = h["summary"].get("opening_balance")
    return h


# ============================================================
# BNI -- Account Statement / Financial Statement BNIDirect (PDF)
# ============================================================
# Kolom: No | Post Date (tgl di atas, jam di bawah) | Branch | Journal No. | Description | Amount | Db/Cr | Balance
# Satu nominal 'Amount' + penanda 'D' (debit) / 'K' (kredit); saldo tercetak di tiap baris.
# KHUSUS file ini: kolom Amount tercetak dobel (lapisan teks ganda). Baris 'D' = angka yang sama persis dua kali,
# baris 'K' = angka asli ditimpa "0.00" (mis. "1.00" + "0.00" terbaca "01.00"; "67,428,216.00" terbaca
# "67,428,2106.00"). dedupe_chars biasa tidak cukup -> _halaman_pdf_bni membuang karakter '0' hantu
# kalau di posisi (x, y) yang sama ada dua karakter BERBEDA.


def _halaman_pdf_bni(content: bytes) -> List[Dict[str, Any]]:
    import pdfplumber
    from pdfplumber.utils import extract_words

    hasil = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for no, page in enumerate(pdf.pages, start=1):
            try:
                per_posisi: Dict[Tuple[float, float], List[Dict[str, Any]]] = {}
                for c in page.chars:
                    per_posisi.setdefault((round(c["x0"], 1), round(c["top"], 1)), []).append(c)
                bersih = []
                for cs in per_posisi.values():
                    if len({c["text"] for c in cs}) > 1:
                        cs = [c for c in cs if c["text"] != "0"] or cs
                    bersih.append(cs[0])
                kata = extract_words(bersih, x_tolerance=1.5, y_tolerance=2)
            finally:
                page.flush_cache()
            hasil.append({"no": no, "kata": [
                {"text": w["text"], "x0": float(w["x0"]), "x1": float(w["x1"]), "top": float(w["top"])} for w in kata
            ]})
    return hasil


def _parse_bni(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("BNI", "words")
    for b in _baris(pages[0]["kata"]):
        kiri = [w for w in b if w["x0"] < 350]  # kotak 'Catatan' di kanan tidak ikut
        t = _teks(kiri)
        m = re.match(r"^(\d{6,})\s*/", t)
        if m and not h["account_number"]:
            h["account_number"] = m.group(1)
        mp = re.findall(r"(\d\d) ([A-Za-z]{3}) (\d{4})", t)
        if len(mp) == 2 and "-" in t and h["period_start"] is None:
            (d1, m1, y1), (d2, m2, y2) = mp
            h["period_start"] = date(int(y1), BULAN_EN[m1.lower()], int(d1))
            h["period_end"] = date(int(y2), BULAN_EN[m2.lower()], int(d2))
        ang = [w for w in kiri if _is_amount(w["text"])]
        cnt = re.search(r"\((\d+)\)", t)
        s = h["summary"]
        if t.startswith("Beginning Balance") and ang:
            s["opening_balance"] = _dec(ang[0]["text"])
        elif t.startswith("Closing Balance") and ang:
            s["closing_balance"] = _dec(ang[0]["text"])
        elif t.startswith("Total Debit") and ang and cnt:
            s["total_debit"], s["count_debit"] = _dec(ang[0]["text"]), int(cnt.group(1))
        elif t.startswith("Total Credit") and ang and cnt:
            s["total_credit"], s["count_credit"] = _dec(ang[0]["text"]), int(cnt.group(1))
    h["opening_balance"] = h["summary"].get("opening_balance")
    h["closing_balance"] = h["summary"].get("closing_balance")

    for pg in pages:
        baris = _baris(pg["kata"])
        hd = next((b for b in baris if any(w["text"] == "Journal" for w in b) and any(w["text"] == "Db/Cr" for w in b)), None)
        if not hd:
            continue
        w_ = lambda nama: next((w for w in hd if w["text"] == nama), None)
        x_dc = (w_("Db/Cr")["x0"] if w_("Db/Cr") else 581) - 2
        x_jurnal = (w_("Journal")["x0"] if w_("Journal") else 227) - 6  # kolom Journal No. rata kiri ~ x 227
        x_desk = (w_("Description")["x0"] if w_("Description") else 280) - 3
        top_hd = hd[0]["top"]
        jangkar: List[Dict[str, Any]] = []
        for b in baris:
            if b[0]["top"] <= top_hd + 6:
                continue
            no_ = next((w for w in b if w["x0"] < 45 and w["text"].isdigit()), None)
            dc = next((w for w in b if w["text"] in ("D", "K") and x_dc <= w["x0"] < x_dc + 30), None)
            ang = [w for w in b if _is_amount(w["text"])]
            if no_ and dc and len(ang) == 2:
                nom = next(w for w in ang if w["x1"] < x_dc)
                sal = next(w for w in ang if w["x1"] > x_dc)
                jur = [w["text"] for w in b if x_jurnal <= w["x0"] < x_desk and w["text"].isdigit()]
                jangkar.append({"y": b[0]["top"], "dc": dc["text"], "nominal": _dec(nom["text"]), "saldo": _dec(sal["text"]),
                                "ref": " ".join(jur), "tgl": [], "jam": [], "desk": [], "page": pg["no"]})
        if not jangkar:
            continue

        def terdekat(y: float) -> Dict[str, Any]:
            return min(jangkar, key=lambda r: abs(r["y"] - y))

        for w in pg["kata"]:
            if w["top"] <= top_hd + 6:
                continue
            if w["x0"] < 100 and re.match(r"^\d\d/\d\d/\d{4}$", w["text"]):
                terdekat(w["top"])["tgl"].append(w["text"])
            elif w["x0"] < 100 and re.match(r"^\d\d\.\d\d\.\d\d$", w["text"]):
                terdekat(w["top"])["jam"].append(w["text"])
            elif x_desk <= w["x0"] < x_dc - 80:
                terdekat(w["top"])["desk"].append(w)
        for r in jangkar:
            if not r["tgl"]:
                h["warnings"].append(f"BNI hal {r['page']}: tanggal baris dengan saldo {r['saldo']:,.2f} tidak terbaca.")
                continue
            tgl = datetime.strptime(r["tgl"][0], "%d/%m/%Y").date()
            jam = r["jam"][0].replace(".", ":") if r["jam"] else None
            desk = " ".join(w["text"] for w in sorted(r["desk"], key=lambda w: (round(w["top"] / 3), w["x0"])))
            debit = r["nominal"] if r["dc"] == "D" else Decimal("0")
            kredit = r["nominal"] if r["dc"] == "K" else Decimal("0")
            _tambah(h, tgl, jam, desk, debit, kredit, r["saldo"], r["ref"], r["page"])
    return h


def _parse_bni_acctinfo(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("BNI", "words")
    s = h["summary"]

    # ---- kepala & ringkasan (semua halaman; ringkasan ada di bawah tabel halaman terakhir) ----
    for pg in pages:
        for b in _baris(pg["kata"]):
            t = _teks(b)
            ang = [w for w in b if _is_amount(w["text"])]
            if not h["account_number"]:
                m = re.search(r"Account No\.\s*(\d{6,})", t)  # satu baris dengan alamat klien di kiri
                if m:
                    h["account_number"] = m.group(1)
            mp = re.findall(r"(\d\d)-([A-Za-z]{3})-(\d\d)", t)
            if t.startswith("Period") and len(mp) == 2 and h["period_start"] is None:
                (d1, m1, y1), (d2, m2, y2) = mp
                h["period_start"] = date(2000 + int(y1), BULAN_EN[m1.lower()], int(d1))
                h["period_end"] = date(2000 + int(y2), BULAN_EN[m2.lower()], int(d2))
            if t.startswith("Ledger Balance") and ang and "opening_balance" not in s:
                s["opening_balance"] = _dec(ang[0]["text"])
            elif t.startswith("Ending Balance") and ang:
                s["closing_balance"] = _dec(ang[0]["text"])
            elif t.startswith("Total Debet") and ang:
                n = next((w["text"] for w in b if w["text"].isdigit()), None)
                s["total_debit"] = _dec(ang[0]["text"])
                s["count_debit"] = int(n) if n else None
            elif t.startswith("Total Credit") and ang:
                n = next((w["text"] for w in b if w["text"].isdigit()), None)
                s["total_credit"] = _dec(ang[0]["text"])
                s["count_credit"] = int(n) if n else None
    h["opening_balance"] = s.get("opening_balance")
    h["closing_balance"] = s.get("closing_balance")

    # ---- baris mutasi ----
    for pg in pages:
        baris = _baris(pg["kata"])
        hd = next((b for b in baris if any(w["text"] == "Journal" for w in b) and any(w["text"] == "DB/CR" for w in b)), None)
        if not hd:
            continue
        w_ = lambda nama: next((w for w in hd if w["text"] == nama), None)
        x_dc = (w_("DB/CR")["x0"] if w_("DB/CR") else 576)   # batas kolom Amount | DB/CR | Balance
        jr = w_("Journal")
        x_jurnal = (jr["x0"] if jr else 296) - 8
        x_desk = (jr["x1"] if jr else 333) + 1          # keterangan mulai tepat sesudah kolom Journal
        am = w_("Amount")
        x_amt = (am["x0"] if am else 516) - 12           # keterangan berhenti sebelum kolom Amount
        top_hd = hd[0]["top"]
        jangkar: List[Dict[str, Any]] = []
        for b in baris:
            if b[0]["top"] <= top_hd + 6:
                continue
            tg = next((w for w in b if w["x0"] < 60 and re.match(r"^\d\d/\d\d/\d{4}$", w["text"])), None)
            dc = next((w for w in b if w["text"] in ("D", "K") and x_dc - 2 <= w["x0"] < x_dc + 40), None)
            if not (tg and dc):
                continue
            jam = next((w["text"] for w in b if w["x0"] < 120 and re.match(r"^\d\d\.\d\d\.\d\d$", w["text"])), None)
            jur = [w["text"] for w in b if x_jurnal <= w["x0"] < x_desk and w["text"].isdigit()]
            jangkar.append({"y": b[0]["top"], "tgl": tg["text"], "jam": jam, "dc": dc["text"], "ref": " ".join(jur),
                            "nominal": None, "saldo": None, "desk": [], "page": pg["no"]})
        # nominal & saldo dicari lewat jendela vertikal: nominal panjang (mis. 150,000,000.00) tercetak ~6pt lebih
        # tinggi dari baris tanggal sehingga terpisah ke baris lain oleh _baris()
        for r in jangkar:
            dekat = [w for w in pg["kata"] if _is_amount(w["text"]) and abs(w["top"] - r["y"]) <= 9]
            nom = min((w for w in dekat if w["x1"] < x_dc), key=lambda w: abs(w["top"] - r["y"]), default=None)
            sal = min((w for w in dekat if w["x0"] > x_dc), key=lambda w: abs(w["top"] - r["y"]), default=None)
            if nom:
                r["nominal"] = _dec(nom["text"])
            if sal:
                r["saldo"] = _dec(sal["text"])   # baris terakhir tabel bisa tanpa saldo tercetak
        jangkar = [r for r in jangkar if r["nominal"] is not None]
        if not jangkar:
            continue
        for w in pg["kata"]:
            if w["top"] <= top_hd + 6 or not (x_desk <= w["x0"] < x_amt):
                continue
            if re.match(r"^\d\d/\d\d/\d{4}$", w["text"]) or _is_amount(w["text"]):
                continue
            min(jangkar, key=lambda r: abs(r["y"] - w["top"]))["desk"].append(w)
        for r in jangkar:
            tgl = datetime.strptime(r["tgl"], "%d/%m/%Y").date()
            jam = r["jam"].replace(".", ":") if r["jam"] else None
            desk = " ".join(w["text"] for w in sorted(r["desk"], key=lambda w: (round(w["top"] / 3), w["x0"])))
            debit = r["nominal"] if r["dc"] == "D" else Decimal("0")
            kredit = r["nominal"] if r["dc"] == "K" else Decimal("0")
            _tambah(h, tgl, jam, desk, debit, kredit, r["saldo"], r["ref"], r["page"])

    # baris terakhir tanpa saldo tercetak -> saldo = "Ending Balance" dokumen, HANYA kalau selisihnya persis cocok
    rows = h["rows"]
    if rows and rows[-1]["balance_after"] is None and h["closing_balance"] is not None:
        prev = next((r["balance_after"] for r in reversed(rows[:-1]) if r["balance_after"] is not None), h["opening_balance"])
        if prev is not None and prev - rows[-1]["debit"] + rows[-1]["credit"] == h["closing_balance"]:
            rows[-1]["balance_after"] = h["closing_balance"]
        else:
            h["warnings"].append("BNI: saldo baris terakhir tidak tercetak dan tidak cocok dengan Ending Balance dokumen.")
    return h


# ============================================================
# Cocokkan kandidat template + bentuk hasil (kompatibel parser generik)
# ============================================================

RESEP: Dict[str, Callable[..., Dict[str, Any]]] = {}  # diisi di bawah (butuh fungsi parser terdefinisi)


def _cocok_pdf(kandidat: List[Dict[str, Any]], p1: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    teks = _teks_halaman_norm(p1)
    for k in kandidat:
        kolom = [norm_header(c) for c in k.get("header_columns") or [] if norm_header(c)]
        if kolom and _urutan_ada(teks, kolom):
            return k
    return None


def _cocok_excel_sheet(kandidat: List[Dict[str, Any]], content: bytes) -> Optional[Tuple[Dict[str, Any], str]]:
    """-> (template, nama sheet tempat header ditemukan). Sheet bernama di rules.sheet_name dicoba dulu, lalu SEMUA
    sheet lain berurutan -- buku kas multi-sheet bulanan (AADL) bisa diawali 'Sheet1' kosong / sheet ringkasan, dan
    laporan Kas Kasir sheet 'Data'-nya tidak selalu yang pertama."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    for k in kandidat:
        sheet = (k.get("rules") or {}).get("sheet_name")
        urut = [n for n in wb.sheetnames if sheet and norm_header(n) == norm_header(sheet)]
        urut += [n for n in wb.sheetnames if n not in urut]
        want = [norm_header(c) for c in k.get("header_columns") or []]
        while want and not want[-1]:
            want.pop()
        if not want:
            continue
        for nama in urut:
            for r in wb[nama].iter_rows(min_row=1, max_row=20, values_only=True):
                have = [norm_header(c) for c in r]
                while have and not have[-1]:
                    have.pop()
                if have[:len(want)] == want:
                    return k, nama
    return None


def _cocok_excel(kandidat: List[Dict[str, Any]], content: bytes) -> Optional[Dict[str, Any]]:
    hit = _cocok_excel_sheet(kandidat, content)
    return hit[0] if hit else None


def _blob_pdf(p1: Dict[str, Any]) -> str:
    return "\n".join(_teks(b) for b in _baris(p1["kata"]))[:6000]


def _blob_excel(content: bytes, sheet: Optional[str] = None) -> str:
    """Teks kop sheet (judul di A2/A3 dst.) untuk pencocokan match_keyword bank mapping. `sheet` = sheet tempat
    header cocok; tanpa itu dipakai sheet pertama."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    bagian: List[str] = []
    for r in wb[sheet if sheet in wb.sheetnames else wb.sheetnames[0]].iter_rows(min_row=1, max_row=12, values_only=True):
        bagian.append(" ".join(str(c) for c in r if c not in (None, "")))
    return "\n".join(bagian)[:6000]


def _f(v: Any) -> Optional[float]:
    return None if v is None else round(float(v), 2)


def coba(content: bytes, filename: str, kandidat_template: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Titik masuk dari bank_statement_parser. Kembalikan `ex` (bentuk sama parser generik) atau None."""
    from . import bank_statement_parser as bsp  # lazy: hindari import melingkar

    kandidat = [k for k in (kandidat_template or []) if (k.get("rules") or {}).get("recipe") in RESEP]
    if not kandidat:
        return None
    nama = (filename or "").lower()
    if nama.endswith(".pdf"):
        tipe = "PDF"
        kandidat = [k for k in kandidat if str(k["rules"]["recipe"]).endswith("_pdf")]
        if not kandidat:
            return None
        p1 = _halaman_pdf(content, hanya_halaman=1)[0]
        tpl = _cocok_pdf(kandidat, p1)
        blob = _blob_pdf(p1)
    elif nama.endswith((".xlsx", ".xlsm")):
        tipe = "Excel"
        kandidat = [k for k in kandidat if str(k["rules"]["recipe"]).endswith("_xlsx")]
        if not kandidat:
            return None
        hit = _cocok_excel_sheet(kandidat, content)
        tpl = hit[0] if hit else None
        blob = _blob_excel(content, hit[1] if hit else None)
    else:
        return None
    if not tpl:
        return None

    rules = tpl["rules"]
    parser = RESEP[rules["recipe"]]
    h = parser(content, rules)  # PDF: parser membaca semua halaman sendiri
    if not h["rows"]:
        raise ResepError("NO_ROWS", "Tidak ada baris mutasi yang terbaca dari file ini.")

    records = [{
        "date": r["date"], "description": r["description"] or "", "debit": _f(r["debit"]) or 0.0, "credit": _f(r["credit"]) or 0.0,
        "balance": _f(r["balance_after"]), "page": r["page_no"], "notes": [], "_pasti": True,
    } for r in h["rows"]]
    opening = _f(h["opening_balance"])
    validasi = bsp.validasi_dan_confidence(records, opening)

    # cocokkan dengan ringkasan RESMI di dokumen (bila ada): total, jumlah baris, saldo awal/akhir
    s = h["summary"]
    td = round(sum(r["debit"] for r in records), 2)
    tk = round(sum(r["credit"] for r in records), 2)
    resmi = [
        ("Total debit sesuai ringkasan dokumen", s.get("total_debit"), td),
        ("Total kredit sesuai ringkasan dokumen", s.get("total_credit"), tk),
        ("Jumlah baris debit sesuai ringkasan dokumen", s.get("count_debit"), sum(1 for r in records if r["debit"] > 0)),
        ("Jumlah baris kredit sesuai ringkasan dokumen", s.get("count_credit"), sum(1 for r in records if r["credit"] > 0)),
        ("Saldo akhir sesuai ringkasan dokumen", s.get("closing_balance"), validasi["closing_balance"]),
    ]
    warnings: List[str] = list(h["warnings"])
    for label, stmt, hasil in resmi:
        if stmt is None or hasil is None:
            continue
        ok = abs(float(stmt) - float(hasil)) <= 0.015
        validasi["checks"].append({"label": label, "statement": float(stmt), "extracted": float(hasil), "ok": ok})
        if not ok:
            validasi["status"] = "failed"
            warnings.append(f"{label}: dokumen {float(stmt):,.2f} vs hasil baca {float(hasil):,.2f}.")

    kolom = [{"header": c, "role": next((x.get("role") for x in rules.get("columns") or [] if bsp._kompak(x.get("header", "")) == bsp._kompak(c)), "ignore"),
              "detected_role": bsp.deteksi_peran(c)} for c in tpl.get("header_columns") or [] if c]
    ada_d = any(c["role"] == "debit" for c in kolom)
    ada_k = any(c["role"] == "credit" for c in kolom)
    tgls = [r["date"] for r in records]
    return {
        "recognized": True, "file_type": tipe, "method": h["method"], "blob": blob, "filename": filename,
        "warnings": warnings, "records": records, "columns": kolom, "header_cells": list(tpl.get("header_columns") or []),
        "header_row_index": int(rules.get("header_row_index") or 1), "header_len": 1,
        "number_format": "en", "date_order": "dmy", "default_direction": None,
        "amount_mode": "split" if (ada_d and ada_k) else "single",
        "template": {"id": tpl["id"], "name": tpl["name"]}, "validation": validasi,
        "account_number": h["account_number"] or bsp.nomor_dari_aturan(rules, blob), "bank_name": rules.get("bank_name") or h["bank_name"],
        "period_start": min(tgls), "period_end": max(tgls),
    }


RESEP.update({
    "bca_pdf": lambda content, rules: _parse_bca(_halaman_pdf(content)),
    "bri_pdf": lambda content, rules: _parse_bri(_halaman_pdf(content)),
    "mandiri_pdf": lambda content, rules: _parse_mandiri(_halaman_pdf(content)),
    "permata_pdf": lambda content, rules: _parse_permata(_halaman_pdf(content, dedupe=True)),
    "bni_pdf": lambda content, rules: _parse_bni(_halaman_pdf_bni(content)),
    "bni_acctinfo_pdf": lambda content, rules: _parse_bni_acctinfo(_halaman_pdf_bni(content)),
    "kas_kasir_xlsx": lambda content, rules: _parse_kas_kasir(content, rules.get("sheet_name")),
})


# ============================================================
# RESEP TAMBAHAN klien AADL (PT Akshaya Ananda Dylan Lestari)
# (sebelumnya modul terpisah bank_statement_recipes_aadl.py -- kini digabung di file ini, tanpa import tambahan)
# ============================================================
# ===================================================
# RESEP tambahan untuk format file klien AADL (PT Akshaya Ananda Dylan Lestari) yang TIDAK terbaca oleh
# resep yang sudah ada di bank_statement_recipes.py (diuji: permata_pdf membaca 0 baris dari kedua PDF
# Permata di bawah, kas_kasir_xlsx menolak header buku kas AADL):
#
#   permata_me_pdf        Permata ME Saver -- "Transaction History" (PDF, gaya aplikasi mobile)
#                         * tanpa saldo & tanpa ringkasan; arah debit/kredit HANYA ada di WARNA nominal
#                           (merah = keluar, hijau = masuk); urutan dokumen = terbaru di atas.
#   permata_syariah_pdf   Permata Unit Usaha Syariah -- "TRANSACTION HISTORY / Financial" (PDF)
#                         * satu kolom Amount bertanda "-" untuk uang keluar; tanpa saldo per baris;
#                           ringkasan Opening/Total Debit/Total Credit/Closing di kepala dokumen;
#                           urutan dokumen = terbaru di atas; nominal rata-tengah vertikal terhadap baris.
#   kas_ledger_xlsx       Buku kas kecil Excel AADL (Rekapan Pembayaran Patty Cash / Kas d'Kappu):
#                         1 workbook = banyak sheet bulanan; kolom NO | TANGGAL | KETERANGAN | DEBET |
#                         CREDIT/KREDIT | SALDO. DEBET = uang MASUK (menambah saldo), CREDIT = uang KELUAR.
#                         Baris "Saldo Awal" bukan mutasi; sel tanggal kosong = ikut tanggal baris di atasnya.
#
#
#
# Bentuk hasil sama persis dengan parser lain di bank_statement_recipes.py (rows/summary/warnings/...),
# jadi coba() dan bank_cash_v1.py tidak perlu diubah. MURNI fungsi (tanpa FastAPI/DB).
# Konvensi arah: debit = uang KELUAR, credit = uang MASUK (sama dengan seluruh modul Bank Cash).

BULAN_PANJANG = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6, "july": 7, "august": 8,
    "september": 9, "october": 10, "november": 11, "december": 12,
    "januari": 1, "februari": 2, "maret": 3, "mei": 5, "juni": 6, "juli": 7, "agustus": 8,
    "oktober": 10, "desember": 12,
}


def _bulan(teks: str) -> Optional[int]:
    t = (teks or "").strip().lower()
    if t in BULAN_PANJANG:
        return BULAN_PANJANG[t]
    k = t[:3]
    for nama, no in BULAN_PANJANG.items():
        if nama[:3] == k and len(t) >= 3:
            return no
    return None


# ============================================================
# Permata ME Saver -- Transaction History (PDF, warna = arah)
# ============================================================

_ME_HEADER_TGL = re.compile(r"^(\d{1,2}) ([A-Za-z]+) (\d{4})$")
_ME_X_NOMINAL_MIN = 600.0     # kolom nominal ("Rp 701,053.00") rata kanan, x0 >= ~672
_ME_Y_FOOTER = 1120.0         # footer "Halaman/Page" & disclaimer mulai ~1139 (tinggi halaman 1242)


def _arah_warna(c: Any) -> Optional[str]:
    """merah (0.87, 0.05, 0.15) = keluar; hijau (0.00, 0.53, 0.45) = masuk."""
    try:
        r, g, b = (float(x) for x in list(c)[:3])
    except (TypeError, ValueError):
        return None
    if r > 0.6 and g < 0.3 and b < 0.4:
        return "keluar"
    if g > 0.4 and r < 0.3:
        return "masuk"
    return None


def _kata_berwarna(content: bytes) -> List[Dict[str, Any]]:
    import pdfplumber

    hasil: List[Dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for no, page in enumerate(pdf.pages, start=1):
            try:
                kata = page.extract_words(x_tolerance=1.5, y_tolerance=2, extra_attrs=["non_stroking_color"])
            finally:
                page.flush_cache()
            hasil.append({"no": no, "kata": [
                {"text": w["text"], "x0": float(w["x0"]), "x1": float(w["x1"]), "top": float(w["top"]),
                 "warna": w.get("non_stroking_color")} for w in kata
            ]})
    return hasil


def _parse_permata_me(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("Permata", "words")
    items: List[Dict[str, Any]] = []   # urutan dokumen (terbaru di atas)
    tak_berarah = 0
    for pg in pages:
        kata = pg["kata"]
        baris = _baris(kata)
        if not h["account_number"]:
            for b in baris:
                m = re.search(r"\b(\d{4})-(\d{4})-(\d{4})\b", _teks(b))
                if m:
                    h["account_number"] = "".join(m.groups())
                    break
        # judul tanggal: baris yang HANYA berisi "10 August 2026"
        heads = []
        for b in baris:
            m = _ME_HEADER_TGL.match(_teks(b).strip())
            if m and _bulan(m.group(2)):
                heads.append((b[0]["top"], date(int(m.group(3)), _bulan(m.group(2)), int(m.group(1)))))
        heads.sort()
        # nominal: kata nominal yang didahului "Rp" di kolom kanan
        nominal = []
        for w in kata:
            if w["x0"] >= _ME_X_NOMINAL_MIN and AMOUNT_RE.match(w["text"]) and w["top"] < _ME_Y_FOOTER:
                nominal.append(w)
        nominal.sort(key=lambda w: w["top"])
        if not nominal:
            continue
        for i, w in enumerate(nominal):
            sebelum = [d for (t, d) in heads if t < w["top"]]
            if not sebelum:
                h["warnings"].append(f"Halaman {pg['no']}: nominal {w['text']} tanpa judul tanggal -- dilewati.")
                continue
            tgl = sebelum[-1]
            top_head = max(t for (t, _) in heads if t < w["top"])
            batas_atas = top_head + 14.0 if i == 0 or nominal[i - 1]["top"] < top_head else (nominal[i - 1]["top"] + w["top"]) / 2
            head_berikut = [t for (t, _) in heads if t > w["top"]]
            batas_bawah_sec = min(head_berikut) if head_berikut else _ME_Y_FOOTER
            if i + 1 < len(nominal) and nominal[i + 1]["top"] < batas_bawah_sec:
                batas_bawah = (w["top"] + nominal[i + 1]["top"]) / 2
            else:
                batas_bawah = batas_bawah_sec
            desk = [x for x in kata
                    if batas_atas <= x["top"] < batas_bawah and x["x0"] < _ME_X_NOMINAL_MIN - 5
                    and not (x["text"] == "Rp" and x["x0"] >= _ME_X_NOMINAL_MIN - 40)]
            # "Rp" di kolom kanan bisa ikut kena -> buang kata "Rp" yang satu baris dengan nominal
            desk = [x for x in desk if not (x["text"] == "Rp" and abs(x["top"] - w["top"]) <= 3 and x["x0"] > 500)]
            teks_desk = " ".join(_teks(b) for b in _baris(desk))
            arah = _arah_warna(w.get("warna"))
            nilai = _dec(w["text"])
            if arah is None:
                tak_berarah += 1
                arah = "keluar"
            jam = re.search(r"\b(\d{2}:\d{2}:\d{2})\b", teks_desk)
            items.append({"tgl": tgl, "desk": teks_desk, "nilai": nilai, "arah": arah,
                          "jam": jam.group(1) if jam else None, "page": pg["no"]})
    if tak_berarah:
        h["warnings"].append(f"{tak_berarah} baris warna nominalnya tidak dikenali -- diasumsikan uang keluar, periksa manual.")
    # dokumen terbaru-di-atas -> balik jadi kronologis
    for it in reversed(items):
        d = it["nilai"] if it["arah"] == "keluar" else Decimal("0")
        k = it["nilai"] if it["arah"] == "masuk" else Decimal("0")
        _tambah(h, it["tgl"], it["jam"], it["desk"], d, k, None, None, it["page"])
    if h["rows"]:
        h["period_start"] = min(r["date"] for r in h["rows"])
        h["period_end"] = max(r["date"] for r in h["rows"])
    h["warnings"].append("Dokumen Permata ME Saver tidak memuat saldo/ringkasan: arah debit/kredit dibaca dari warna nominal "
                         "(merah = keluar, hijau = masuk) dan tidak bisa dicocokkan dengan saldo resmi.")
    return h


# ============================================================
# Permata Unit Usaha Syariah -- Transaction History (PDF)
# ============================================================

_SY_X_TGL_MAKS = 30.0      # tanggal transaksi mulai x ~ 24
_SY_X_DESK_MIN = 240.0     # kolom Description mulai x ~ 248
_SY_X_AMT_MIN = 500.0      # kolom Amount (rata kanan, x1 ~ 571); '-' penanda keluar tepat di kiri nominal
_SY_X_REF_MIN, _SY_X_REF_MAKS = 125.0, 240.0


def _parse_permata_syariah(pages: List[Dict[str, Any]]) -> Dict[str, Any]:
    h = _hasil("Permata Syariah", "words")
    s = h["summary"]
    items: List[Dict[str, Any]] = []
    for pg in pages:
        kata = pg["kata"]
        baris = _baris(kata)
        for b in baris:
            t = _teks(b)
            if not h["account_number"]:
                m = re.search(r"Account No\.\s*(\d{4}-\d{4}-\d{3,4})", t)
                if m:
                    h["account_number"] = m.group(1).replace("-", "")
            for lab, kunci in (("Opening Balance", "opening_balance"), ("Total Debit", "total_debit"),
                               ("Total Credit", "total_credit"), ("Closing Balance", "closing_balance")):
                m = re.search(re.escape(lab) + r"\s+(-?[\d,]+\.\d{2})", t)
                if m and kunci not in s:
                    s[kunci] = _dec(m.group(1))
        # jangkar baris: "dd Mon yyyy dd Mon yyyy" di kolom kiri
        jangkar = []
        for b in baris:
            kiri = [w for w in b if w["x0"] < 125.0]
            if len(kiri) >= 6 and kiri[0]["x0"] <= _SY_X_TGL_MAKS:
                tok = [w["text"] for w in kiri[:3]]
                if re.fullmatch(r"\d{2}", tok[0]) and _bulan(tok[1]) and re.fullmatch(r"\d{4}", tok[2]):
                    jangkar.append((b[0]["top"], date(int(tok[2]), _bulan(tok[1]), int(tok[0]))))
        jangkar.sort()
        for i, (top, tgl) in enumerate(jangkar):
            atas = top - 4.0
            bawah = jangkar[i + 1][0] - 4.0 if i + 1 < len(jangkar) else 10 ** 6
            area = [w for w in kata if atas <= w["top"] < bawah]
            # buang footer "Page 1 of 2"
            area = [w for w in area if not re.fullmatch(r"Page|of", w["text"]) or w["x0"] < _SY_X_DESK_MIN]
            nom = [w for w in area if AMOUNT_RE.match(w["text"]) and w["x0"] >= _SY_X_AMT_MIN - 40 and w["x1"] > 540]
            if not nom:
                h["warnings"].append(f"Halaman {pg['no']}: baris {tgl:%d/%m/%Y} tanpa nominal -- dilewati.")
                continue
            w = nom[0]
            minus = any(x["text"] == "-" and x["x0"] >= _SY_X_AMT_MIN - 40 and w["x0"] - 22 <= x["x0"] < w["x0"]
                        and abs(x["top"] - w["top"]) <= 3 for x in area)
            desk_w = [x for x in area if _SY_X_DESK_MIN <= x["x0"] < _SY_X_AMT_MIN - 40 and x is not w]
            desk = " ".join(_teks(bb) for bb in _baris(desk_w))
            ref_w = [x for x in area if _SY_X_REF_MIN <= x["x0"] < _SY_X_REF_MAKS and abs(x["top"] - top) <= 3]
            ref = _teks(sorted(ref_w, key=lambda x: x["x0"])).strip() or None
            if ref == "-":
                ref = None
            nilai = _dec(w["text"])
            items.append({"tgl": tgl, "desk": desk, "ref": ref, "debit": nilai if minus else Decimal("0"),
                          "kredit": Decimal("0") if minus else nilai, "page": pg["no"]})
    for it in reversed(items):   # dokumen terbaru-di-atas -> kronologis
        _tambah(h, it["tgl"], None, it["desk"], it["debit"], it["kredit"], None, it["ref"], it["page"])
    if h["rows"]:
        h["period_start"] = min(r["date"] for r in h["rows"])
        h["period_end"] = max(r["date"] for r in h["rows"])
        h["opening_balance"] = s.get("opening_balance")
        h["closing_balance"] = s.get("closing_balance")
        if s.get("closing_balance") is not None:   # saldo hanya tercetak di kepala -> jadikan checkpoint baris terakhir
            h["rows"][-1]["balance_after"] = s["closing_balance"]
    return h


# ============================================================
# Buku kas kecil Excel AADL (multi-sheet bulanan)
# ============================================================

_KAS_BLN_SHEET = ("jan", "feb", "mar", "apr", "mei", "may", "jun", "jul", "agu", "aug", "sep", "okt", "oct", "nov", "des", "dec")


def _sheet_bulanan(nama: str) -> bool:
    n = (nama or "").strip().lower()
    return bool(n) and not n.startswith("sheet") and n[:3] in _KAS_BLN_SHEET


def _tgl_apa_pun(v: Any) -> Optional[date]:
    d = _tgl_excel(v)
    if d:
        return d
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*", v)
        if m:
            th = int(m.group(3))
            th = th + 2000 if th < 100 else th
            try:
                return date(th, int(m.group(2)), int(m.group(1)))
            except ValueError:
                return None
    return None


def _nominal_sel(v: Any) -> Decimal:
    """Sel angka buku kas: kosong / spasi / teks non-angka = 0 (mis. sel ' ' di kolom CREDIT)."""
    if v is None or isinstance(v, bool):
        return Decimal("0")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    t = str(v).strip().replace(",", "")
    try:
        return Decimal(t) if t else Decimal("0")
    except Exception:  # noqa: BLE001
        return Decimal("0")


def _bulan_sheet(nama: str) -> Optional[date]:
    """'Desember 2025' -> 2025-12-01 (tanggal cadangan untuk baris tanpa tanggal di awal sheet)."""
    m = re.search(r"([A-Za-z]+)\s+(\d{4})", nama or "")
    if m and _bulan(m.group(1)):
        return date(int(m.group(2)), _bulan(m.group(1)), 1)
    return None


def _cari_header_ledger(rows: List[Tuple[Any, ...]]) -> Optional[Tuple[int, Dict[str, int]]]:
    for i, r in enumerate(rows[:20]):
        n = [norm_header(c) for c in r]
        if "tanggal" not in n or "saldo" not in n:
            continue
        masuk = n.index("debet") if "debet" in n else None
        keluar = n.index("credit") if "credit" in n else (n.index("kredit") if "kredit" in n else None)
        if masuk is None or keluar is None:
            continue
        tgl = n.index("tanggal")
        ket = n.index("keterangan") if "keterangan" in n else tgl + 1   # Omah: judul kolom keterangan = spasi kosong
        return i, {"tanggal": tgl, "keterangan": ket, "masuk": masuk, "keluar": keluar, "saldo": n.index("saldo")}
    return None


def _parse_kas_ledger(content: bytes, rules: Dict[str, Any]) -> Dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    h = _hasil(str(rules.get("bank_name") or "Kas Kecil"), "table")
    tgls: List[date] = []
    sheet_terbaca: List[str] = []
    saldo_akhir_sheet: Optional[Decimal] = None
    tanpa_tgl = luar_bulan = 0
    for nama in wb.sheetnames:
        if not _sheet_bulanan(nama):
            continue
        rows = [tuple(r) for r in wb[nama].iter_rows(values_only=True)]
        found = _cari_header_ledger(rows)
        if not found:
            continue
        hi, ix = found
        awal_sheet: Optional[Decimal] = None
        tgl_skrg: Optional[date] = None
        n_sheet = 0
        saldo_prev: Optional[Decimal] = None
        for no, r in enumerate(rows[hi + 1:], start=hi + 2):
            def sel(k: str) -> Any:
                return r[ix[k]] if len(r) > ix[k] else None
            ket = str(sel("keterangan") or "").strip()
            nk = norm_header(ket)
            if nk.startswith("saldoawal"):
                if sel("saldo") is not None:
                    awal_sheet = _nominal_sel(sel("saldo"))
                    saldo_prev = awal_sheet
                continue
            if nk.startswith("saldo") or nk.startswith(("total", "jumlah")):   # baris penutup (Saldo Akhir / Saldo Per 2025 / Saldo 2026 / Total)
                break
            masuk, keluar = _nominal_sel(sel("masuk")), _nominal_sel(sel("keluar"))
            if not ket or (masuk == 0 and keluar == 0):
                continue
            tg = _tgl_apa_pun(sel("tanggal"))
            if tg:
                tgl_skrg = tg
            if not tg and not tgl_skrg:
                tgl_skrg = _bulan_sheet(nama)       # baris pertama sheet tanpa tanggal -> tgl 1 bulan sheet
                if tgl_skrg:
                    tanpa_tgl += 1
            tg = tg or tgl_skrg
            if not tg:
                h["warnings"].append(f"Sheet '{nama}' baris {no}: tanpa tanggal -- dilewati.")
                continue
            bs = _bulan_sheet(nama)
            if bs and (tg.year, tg.month) != (bs.year, bs.month):
                luar_bulan += 1
            saldo = sel("saldo")
            saldo = _nominal_sel(saldo) if saldo is not None else None
            if not h["rows"] and h["opening_balance"] is None:
                h["opening_balance"] = awal_sheet
            elif n_sheet == 0 and saldo_akhir_sheet is not None and awal_sheet is not None \
                    and abs(awal_sheet - saldo_akhir_sheet) > Decimal("0.01"):
                h["warnings"].append(f"Sheet '{nama}': saldo awal {awal_sheet:,.2f} tidak sama dengan saldo akhir sheet sebelumnya {saldo_akhir_sheet:,.2f}.")
            n_sheet += 1
            tgls.append(tg)
            _tambah(h, tg, None, ket, keluar, masuk, saldo, None, no, kategori=nama)
            saldo_akhir_sheet = saldo if saldo is not None else saldo_akhir_sheet
        if n_sheet:
            sheet_terbaca.append(nama)
    if not h["rows"]:
        raise ResepError("NO_ROWS", "Tidak ada baris mutasi di sheet bulanan buku kas ini.")
    h["period_start"], h["period_end"] = min(tgls), max(tgls)
    if tanpa_tgl:
        h["warnings"].append(f"{tanpa_tgl} sheet punya baris pertama tanpa tanggal -- dipakai tanggal 1 bulan sheet.")
    if luar_bulan:
        h["warnings"].append(f"{luar_bulan} baris bertanggal di luar bulan sheet-nya (kemungkinan salah ketik tanggal) -- periksa.")
    h["closing_balance"] = next((r["balance_after"] for r in reversed(h["rows"]) if r["balance_after"] is not None), None)
    h["summary"]["sheets"] = sheet_terbaca
    return h


RESEP.update({
    "permata_me_pdf": lambda content, rules: _parse_permata_me(_kata_berwarna(content)),
    "permata_syariah_pdf": lambda content, rules: _parse_permata_syariah(_halaman_pdf(content)),
    "kas_ledger_xlsx": lambda content, rules: _parse_kas_ledger(content, rules),
})