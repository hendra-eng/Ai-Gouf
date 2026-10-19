"""
modules/transactions/bank_statement_parser.py
=============================================
Parser rekening koran / laporan kas (Excel, CSV, PDF ber-text-layer) untuk
Cash & Bank > Source Data. MURNI fungsi -- tanpa FastAPI/DB -- supaya bisa dites
sendiri. Dipakai bank_cash_v1.py.

Alur:
  1. baca file -> "sheet" berisi baris sel mentah (Excel/CSV: per sheet; PDF: tabel
     per halaman, fallback baca teks per baris kalau tabelnya tidak terdeteksi).
  2. cari baris header (butuh kolom tanggal + debit/kredit/jumlah), tiap header
     ditebak perannya (date/description/debit/credit/amount/dc_flag/balance/...).
  3. kalau ada template klien yang header_columns-nya sama, resep (peran kolom,
     format angka/tanggal, arah default) diambil dari template itu.
  4. ekstrak baris -> validasi saldo berjalan -> confidence per baris.

Konvensi arah (sama dengan frontend): debit = uang KELUAR dari rekening,
credit = uang MASUK ke rekening.

BATASAN yang disengaja (lihat catatan di bank_cash_v1.py): resep khusus per bank
yang butuh regex/posisi kata (mis. BCA PDF dengan saldo hanya di baris terakhir
tiap hari, BRI/Mandiri PDF) TIDAK dijalankan di sini -- parser ini generik.
PDF hasil scan (tanpa text layer) tidak didukung (butuh OCR).
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

TOL = 0.015  # toleransi selisih saldo (pembulatan sen)

BULAN = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "may": 5, "jun": 6, "jul": 7,
    "agu": 8, "agt": 8, "ags": 8, "aug": 8, "sep": 9, "okt": 10, "oct": 10, "nov": 11, "des": 12, "dec": 12,
}

BANK_DIKENAL = [
    "BCA", "BRI", "BNI", "MANDIRI", "CIMB", "PERMATA", "DANAMON", "BTN", "OCBC", "MAYBANK",
    "PANIN", "BSI", "JAGO", "BTPN", "MEGA", "SINARMAS", "UOB", "HSBC", "BPD BALI", "BANK BALI",
]


class ParserError(Exception):
    def __init__(self, message: str, code: str = "PARSE_ERROR"):
        super().__init__(message)
        self.message = message
        self.code = code


# ============================================================
# 1) Normalisasi teks / angka / tanggal
# ============================================================

def _teks(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return str(int(v))
    return str(v).replace("\xa0", " ").strip()


def _kompak(s: str) -> str:
    """huruf kecil, buang semua selain a-z0-9/ -- PDF BCA mencetak huruf terpisah spasi ('T ANG G A L')."""
    return re.sub(r"[^a-z0-9/]+", "", (s or "").lower())


_ROLE_RULES: List[Tuple[str, str, bool]] = [
    # (peran, regex pada teks kompak, exact?)
    ("balance_before", r"^(saldoawal|saldosebelum\w*|balancebefore|openingbalance)$", True),
    ("dc_flag", r"^(d/?k|db/?cr|cr/?db|dc|d/?c|k/?d|dr/?cr|cr/?dr|jenis\w*|tipe\w*|type|dk)$", True),
    ("balance", r"^(saldo\w*|balance\w*|runningbalance)$", True),
    ("debit", r"(debet|debit|keluar|withdraw|pengeluaran|^db$|^dr$)", False),
    ("credit", r"(kredit|credit|masuk|deposit|penerimaan|^cr$)", False),
    ("amount", r"^(jumlah|jml|amount|nominal|mutasi|nilai)$", True),
    ("date", r"(tanggal|^tgl|date|^tanggaltransaksi)", False),
    ("description", r"(keterangan|^ket$|uraian|deskripsi|description|remark|berita|narasi|particular|detail|transaksi|catatan)", False),
    ("reference", r"^(no|noref\w*|ref\w*|nobukti|nodok\w*|cbg|cabang|teller|voucher\w*)$", True),
]


def deteksi_peran(header: str) -> Optional[str]:
    k = _kompak(header)
    if not k:
        return None
    for peran, pola, exact in _ROLE_RULES:
        if (re.match if exact else re.search)(pola, k):
            return peran
    return None


def deteksi_format_angka(sampel: List[str]) -> str:
    id_n = en_n = 0
    for s in sampel:
        s = re.sub(r"[^\d.,]", "", s or "")
        if re.search(r"\d{1,3}(\.\d{3})+,\d{1,2}$", s) or re.search(r"^\d+,\d{2}$", s):
            id_n += 1
        elif re.search(r"\d{1,3}(,\d{3})+\.\d{1,2}$", s) or re.search(r"^\d+\.\d{2}$", s):
            en_n += 1
        elif re.search(r"\d{1,3}(\.\d{3})+$", s):
            id_n += 1
        elif re.search(r"\d{1,3}(,\d{3})+$", s):
            en_n += 1
    return "id" if id_n > en_n else "en"


_MARK_DEBIT = {"DB", "DR", "D", "DEBIT", "DEBET"}
_MARK_CREDIT = {"CR", "K", "C", "KR", "CREDIT", "KREDIT"}


def parse_angka(v: Any, fmt: str = "en") -> Tuple[Optional[float], Optional[str]]:
    """-> (nilai, marker 'D'/'C'/None). Marker dari sufiks 'DB'/'CR'. Nilai bisa negatif."""
    if v is None or isinstance(v, bool):
        return None, None
    if isinstance(v, (int, float)):
        f = float(v)
        return (None, None) if f != f else (f, None)
    s = _teks(v)
    if not s:
        return None, None
    marker = None
    m = re.search(r"(?<=[\d\s)])\s*([A-Za-z]{1,6})\.?$", s)
    if m and m.group(1).upper() in (_MARK_DEBIT | _MARK_CREDIT):
        marker = "D" if m.group(1).upper() in _MARK_DEBIT else "C"
        s = s[: m.start()].strip()
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    s = re.sub(r"(?i)\b(rp|idr|usd)\b\.?", "", s).replace(" ", "")
    if s.endswith("-"):
        neg, s = True, s[:-1]
    if s.startswith("-"):
        neg, s = True, s[1:]
    s = s.lstrip("+")
    if not re.search(r"\d", s) or re.search(r"[^\d.,]", s):
        return None, None
    if fmt == "id":
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        f = float(s)
    except ValueError:
        return None, None
    return (-f if neg else f), marker


def _valid_date(y: int, m: int, d: int) -> Optional[date]:
    if y < 100:
        y += 2000
    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_tanggal(v: Any, order: str = "dmy", year_hint: Optional[int] = None) -> Optional[date]:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _teks(v)
    if not s:
        return None
    m = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        return _valid_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r"^(\d{1,2})[-/. ](\d{1,2})[-/. ](\d{2,4})(?!\d)", s)
    if m:
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return _valid_date(y, a, b) if order == "mdy" else _valid_date(y, b, a)
    m = re.match(r"^(\d{1,2})[-/. ]([A-Za-z]{3,9})[-/. ,]*(\d{2,4})?", s)
    if m and m.group(2)[:3].lower() in BULAN:
        y = int(m.group(3)) if m.group(3) else year_hint
        return _valid_date(y, BULAN[m.group(2)[:3].lower()], int(m.group(1))) if y else None
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})$", s)
    if m and year_hint:
        a, b = int(m.group(1)), int(m.group(2))
        return _valid_date(year_hint, a, b) if order == "mdy" else _valid_date(year_hint, b, a)
    return None


# ============================================================
# 2) Baca file -> sheet berisi baris mentah
# ============================================================

_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def deteksi_tipe(filename: str) -> str:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    return {"xlsx": "Excel", "xlsm": "Excel", "xls": "Excel", "csv": "CSV", "pdf": "PDF", "txt": "CSV"}.get(ext, ext.upper() or "File")


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "iso-8859-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _baca_csv(content: bytes) -> List[Dict[str, Any]]:
    teks = _decode(content)
    sampel = "\n".join(teks.splitlines()[:30])
    try:
        delim = csv.Sniffer().sniff(sampel, delimiters=",;\t|").delimiter
    except csv.Error:
        delim = max(",;\t|", key=lambda d: sampel.count(d))
    rows = [list(r) for r in csv.reader(io.StringIO(teks), delimiter=delim)]
    return [{"name": "CSV", "rows": rows, "pages": [None] * len(rows)}]


def _xls_ke_xlsx(content: bytes) -> bytes:
    try:
        import openpyxl
        import xlrd

        buku = xlrd.open_workbook(file_contents=content)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sh in buku.sheets():
            ws = wb.create_sheet(title=sh.name[:31] or "Sheet")
            for r in range(sh.nrows):
                for c in range(sh.ncols):
                    t, v = sh.cell_type(r, c), sh.cell_value(r, c)
                    if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                        continue
                    if t == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate.xldate_as_datetime(v, buku.datemode)
                        except Exception:
                            pass
                    ws.cell(row=r + 1, column=c + 1, value=v)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except ImportError:
        raise ParserError("File .xls (Excel lama) butuh paket 'xlrd' -- simpan ulang sebagai .xlsx atau pip install xlrd.", "XLS_UNSUPPORTED")


def _baca_excel(content: bytes) -> List[Dict[str, Any]]:
    import openpyxl

    if content.startswith(_OLE2):
        content = _xls_ke_xlsx(content)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ParserError(f"File Excel tidak bisa dibuka: {e}", "BAD_FILE")
    sheets = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(_teks(c) for c in r)]
        if rows:
            sheets.append({"name": ws.title, "rows": rows, "pages": [None] * len(rows)})
    return sheets


_RE_BARIS_TGL = re.compile(
    r"^(\d{1,2}[/-]\d{1,2}(?:[/-]\d{2,4})?|\d{4}-\d{2}-\d{2}|\d{1,2}[ -][A-Za-z]{3,9}[ -]?\d{0,4})\s+(.*)$"
)
_RE_ANGKA_TOKEN = re.compile(r"^[(\-+]?\d[\d.,]*\d\)?(?:CR|DB|DR)?-?$|^\d$", re.I)


def _baris_dari_teks(teks_halaman: str) -> List[List[str]]:
    """Fallback PDF: 1 baris teks berawal tanggal -> [tanggal, keterangan, debit, kredit, jumlah, saldo]."""
    hasil: List[List[str]] = []
    for line in teks_halaman.splitlines():
        line = line.strip()
        if not line:
            continue
        m = _RE_BARIS_TGL.match(line)
        if not m:
            if hasil and not re.match(r"(?i)^(halaman|page|saldo awal|total)", line):
                hasil.append(["", line, "", "", "", ""])  # baris lanjutan keterangan
            continue
        tgl, sisa = m.groups()
        token = sisa.split()
        gabung: List[str] = []
        for tk in token:  # penanda 'DB'/'CR' yang terpisah spasi dari angkanya -> tempel ke angka
            if gabung and tk.upper() in ("DB", "CR", "DR") and _RE_ANGKA_TOKEN.match(gabung[-1]):
                gabung[-1] += tk.upper()
            else:
                gabung.append(tk)
        token = gabung
        angka: List[str] = []
        while token and len(angka) < 3 and _RE_ANGKA_TOKEN.match(token[-1]) and (
            len(angka) > 0 or re.search(r"[.,]\d{1,2}(?:CR|DB|DR)?$", token[-1], re.I) or len(token[-1]) >= 4
        ):
            angka.insert(0, token.pop())
        ket = " ".join(token)
        d = k = j = s = ""
        if len(angka) == 3:
            d, k, s = angka
        elif len(angka) == 2:
            j, s = angka
        elif len(angka) == 1:
            j = angka[0]
        hasil.append([tgl, ket, d, k, j, s])
    return hasil


_HEADER_TEKS_PDF = ["Tanggal", "Keterangan", "Debit", "Kredit", "Jumlah", "Saldo"]


def _baca_pdf(content: bytes) -> List[Dict[str, Any]]:
    try:
        import pdfplumber
    except ImportError:
        raise ParserError("Paket 'pdfplumber' belum terpasang di backend (pip install pdfplumber).", "PDF_UNSUPPORTED")
    rows: List[List[Any]] = []
    pages: List[Optional[int]] = []
    teks_total = []
    fallback_rows: List[List[Any]] = []
    fallback_pages: List[Optional[int]] = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for no, page in enumerate(pdf.pages, start=1):
                teks = page.extract_text() or ""
                teks_total.append(teks)
                tabel = []
                for setelan in ({}, {"vertical_strategy": "text", "horizontal_strategy": "text"}):
                    try:
                        tabel = page.extract_tables(setelan) if setelan else page.extract_tables()
                    except Exception:
                        tabel = []
                    if tabel:
                        break
                for t in tabel:
                    for r in t:
                        cells = [(c or "").replace("\n", " ").strip() for c in r]
                        if any(cells):
                            rows.append(cells)
                            pages.append(no)
                for r in _baris_dari_teks(teks):
                    fallback_rows.append(r)
                    fallback_pages.append(no)
    except ParserError:
        raise
    except Exception as e:
        raise ParserError(f"PDF tidak bisa dibaca: {e}", "BAD_FILE")
    if not any(t.strip() for t in teks_total):
        raise ParserError(
            "PDF ini tidak punya text layer (hasil scan). OCR belum terpasang di backend -- upload versi Excel/CSV atau PDF asli dari internet banking.",
            "PDF_SCANNED",
        )
    sheets = [{"name": "PDF", "rows": rows, "pages": pages, "teks": "\n".join(teks_total)}]
    sheets.append({
        "name": "PDF-teks", "rows": [_HEADER_TEKS_PDF] + fallback_rows, "pages": [None] + fallback_pages,
        "teks": "\n".join(teks_total), "fallback": True,
    })
    return sheets


def baca_sheets(content: bytes, filename: str) -> Tuple[str, List[Dict[str, Any]]]:
    tipe = deteksi_tipe(filename)
    if tipe == "Excel":
        return tipe, _baca_excel(content)
    if tipe == "PDF":
        return tipe, _baca_pdf(content)
    if tipe == "CSV":
        return tipe, _baca_csv(content)
    raise ParserError(f"Tipe file '{filename}' belum didukung. Gunakan PDF, Excel (.xlsx/.xls) atau CSV.", "UNSUPPORTED_FILE")


# ============================================================
# 3) Header & resep kolom
# ============================================================

def _cari_header(rows: List[List[Any]], maks: int = 80, kandidat: Optional[List[Dict[str, Any]]] = None) -> Optional[Tuple[int, int, List[str], List[Optional[str]]]]:
    punya_tpl = [_kompak_list([str(x) for x in (k.get("header_columns") or [])]) for k in (kandidat or [])]

    def nilai(cells: List[str]) -> Optional[List[Optional[str]]]:
        roles = [deteksi_peran(c) if c else None for c in cells]
        if _kompak_list(cells) in punya_tpl and punya_tpl:
            return roles
        ada_tgl = "date" in roles
        ada_uang = any(r in ("debit", "credit", "amount") for r in roles)
        return roles if ada_tgl and ada_uang and sum(1 for r in roles if r) >= 3 else None

    # Pass 1: header 1 baris di seluruh jendela -- diutamakan supaya baris kop ("Periode : ...") tidak ikut tergabung.
    for i, row in enumerate(rows[:maks]):
        cells = [_teks(c) for c in row]
        roles = nilai(cells)
        if roles:
            if i + 1 < len(rows):  # baris di bawahnya berisi sub-header Debit/Kredit (tanpa tanggal)? -> header 2 baris
                nxt = [_teks(c) for c in rows[i + 1]]
                sub = [deteksi_peran(c) for c in nxt if c]
                if nxt and ("debit" in sub or "credit" in sub) and not any(parse_tanggal(c) for c in nxt if c):
                    n = max(len(cells), len(nxt))
                    gabung = [((cells[j] if j < len(cells) else "") + " " + (nxt[j] if j < len(nxt) else "")).strip() for j in range(n)]
                    r2 = nilai(gabung)
                    if r2:
                        return i, 2, gabung, r2
            return i, 1, cells, roles
    # Pass 2: header 2 baris (mis. 'Mutasi' di atas 'Debit'/'Kredit'); baris pertama harus punya >= 2 sel terisi.
    for i, row in enumerate(rows[:maks]):
        if i + 1 >= len(rows):
            break
        cells = [_teks(c) for c in row]
        if sum(1 for c in cells if c) < 2:
            continue
        nxt = [_teks(c) for c in rows[i + 1]]
        n = max(len(cells), len(nxt))
        gabung = [((cells[j] if j < len(cells) else "") + " " + (nxt[j] if j < len(nxt) else "")).strip() for j in range(n)]
        roles = nilai(gabung)
        if roles:
            return i, 2, gabung, roles
    return None


def signature(header_cells: List[str]) -> str:
    isi = "|".join(_kompak(c) for c in header_cells if _kompak(c))
    return hashlib.sha256(isi.encode("utf-8")).hexdigest()


def _kompak_list(cells: List[str]) -> List[str]:
    return [_kompak(c) for c in cells if _kompak(c)]


def cocokkan_template(header_cells: List[str], kandidat: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """kandidat: [{id, name, header_columns, rules, ...}] -> kandidat yang header_columns-nya SAMA."""
    kita = _kompak_list(header_cells)
    for k in kandidat:
        if kita and _kompak_list([str(x) for x in (k.get("header_columns") or [])]) == kita:
            return k
    return None


def _peran_dari_rules(header_cells: List[str], rules: Dict[str, Any]) -> Optional[List[Optional[str]]]:
    kolom = rules.get("columns") or []
    if not kolom:
        return None
    peta = {_kompak(c.get("header", "")): c.get("role") for c in kolom if c.get("header")}
    roles: List[Optional[str]] = []
    for idx, h in enumerate(header_cells):
        r = peta.get(_kompak(h)) if _kompak(h) else None
        if r is None and len(kolom) == len(header_cells):
            r = kolom[idx].get("role")
        roles.append(None if r in (None, "ignore") else r)
    return roles


# ============================================================
# 4) Ekstraksi baris + validasi
# ============================================================

def _arah_dc(v: str) -> Optional[str]:
    s = (v or "").strip().upper().replace(".", "")
    if s in _MARK_DEBIT or s in ("DEBIT", "DEBET", "D/B"):
        return "D"
    if s in _MARK_CREDIT:
        return "C"
    return None


def _year_hint(blob: str, filename: str) -> Optional[int]:
    for sumber in (blob, filename):
        m = re.search(r"(?:periode|period)[^\n]{0,40}?(20\d{2})", sumber or "", re.I) or re.search(r"(20\d{2})", sumber or "")
        if m:
            return int(m.group(1))
    return None


def _ekstrak_sheet(
    rows: List[List[Any]], pages: List[Optional[int]], hdr_idx: int, hdr_len: int,
    roles: List[Optional[str]], fmt: str, order: str, default_dir: Optional[str], year: Optional[int],
) -> Dict[str, Any]:
    idx: Dict[str, List[int]] = {}
    for j, r in enumerate(roles):
        if r:
            idx.setdefault(r, []).append(j)

    def sel(row: List[Any], role: str) -> Any:
        for j in idx.get(role, []):
            if j < len(row) and _teks(row[j]):
                return row[j]
        return None

    records: List[Dict[str, Any]] = []
    saldo_awal: Optional[float] = None
    prev_bal: Optional[float] = None
    tanpa_arah = 0
    for ri in range(hdr_idx + hdr_len, len(rows)):
        row = rows[ri]
        page = pages[ri] if ri < len(pages) else None
        tgl = parse_tanggal(sel(row, "date"), order, year)
        desc = " ".join(_teks(row[j]) for j in idx.get("description", []) + idx.get("reference", []) if j < len(row) and _teks(row[j]))
        d_raw, k_raw, a_raw = sel(row, "debit"), sel(row, "credit"), sel(row, "amount")
        b_raw = sel(row, "balance")
        flag = _arah_dc(_teks(sel(row, "dc_flag")))
        d, d_mark = parse_angka(d_raw, fmt)
        k, k_mark = parse_angka(k_raw, fmt)
        a, a_mark = parse_angka(a_raw, fmt)
        bal, _ = parse_angka(b_raw, fmt)
        bef, _ = parse_angka(sel(row, "balance_before"), fmt)

        if re.match(r"(?i)^\s*(saldo\s*awal|opening balance|beginning balance)", desc or "") or (tgl is None and re.match(r"(?i)saldo\s*awal", _teks(row[0]) if row else "")):
            if bal is not None:
                saldo_awal, prev_bal = bal, bal
            elif a is not None:
                saldo_awal = prev_bal = a
            continue
        if tgl is None:
            if desc and d is None and k is None and a is None and records and not re.match(r"(?i)^\s*(total|jumlah|saldo akhir|ending|grand|halaman|page)", desc):
                records[-1]["description"] = (records[-1]["description"] + " " + desc).strip()  # baris lanjutan
            continue

        debit = abs(d) if d else 0.0
        kredit = abs(k) if k else 0.0
        catatan: List[str] = []
        pasti = True
        if not debit and not kredit and a is not None and a != 0:
            nilai = abs(a)
            arah = {"D": "debit", "C": "credit"}.get(flag or a_mark or "")
            if arah is None and a < 0:
                arah = "debit"
            if arah is None and bal is not None and prev_bal is not None:
                delta = bal - prev_bal
                if abs(delta - nilai) < TOL:
                    arah = "credit"
                elif abs(delta + nilai) < TOL:
                    arah = "debit"
            if arah is None and default_dir:
                arah = default_dir
            if arah is None:
                arah, pasti = "credit", False
                tanpa_arah += 1
                catatan.append("Arah debit/kredit tidak terdeteksi (diasumsikan kredit) -- cek manual")
            if arah == "debit":
                debit = nilai
            else:
                kredit = nilai
        if not debit and not kredit:
            if bal is not None:
                prev_bal = bal
            continue
        if bef is not None and saldo_awal is None and not records:
            saldo_awal = bef
        records.append({
            "date": tgl, "description": desc, "debit": round(debit, 2), "credit": round(kredit, 2),
            "balance": round(bal, 2) if bal is not None else None, "page": page, "notes": catatan, "_pasti": pasti,
        })
        if bal is not None:
            prev_bal = bal
    return {"records": records, "opening": saldo_awal, "tanpa_arah": tanpa_arah}


def validasi_dan_confidence(records: List[Dict[str, Any]], opening: Optional[float]) -> Dict[str, Any]:
    """Saldo berjalan: tiap baris ber-saldo = checkpoint. Semua baris sejak checkpoint sebelumnya
    di-'verified' bila cocok, di-'low' bila selisih. Baris tanpa checkpoint sesudahnya = 'medium'."""
    cek = bad = 0
    run = opening
    grup: List[Dict[str, Any]] = []
    if run is None and records and records[0]["balance"] is not None:
        r0 = records[0]
        run = round(r0["balance"] - r0["credit"] + r0["debit"], 2)
        opening = run
    for r in records:
        grup.append(r)
        if run is not None:
            run = round(run + r["credit"] - r["debit"], 2)
        if r["balance"] is not None:
            if run is not None:
                cek += 1
                if abs(run - r["balance"]) <= TOL:
                    for g in grup:
                        g["confidence"] = "high" if g["_pasti"] else "low"
                else:
                    bad += 1
                    for g in grup:
                        g["confidence"] = "low"
                        g["notes"].append(f"Saldo berjalan tidak cocok (hitung {run:,.2f} vs tercetak {r['balance']:,.2f})")
            else:
                for g in grup:
                    g["confidence"] = "medium"
            run = r["balance"]  # sinkron ke saldo tercetak, supaya 1 selisih tidak menular ke baris berikutnya
            grup = []
    for g in grup:
        g.setdefault("confidence", "medium")
        if not g["_pasti"]:
            g["confidence"] = "low"
    for r in records:
        r.setdefault("confidence", "medium")
        r.pop("_pasti", None)
    total_d = round(sum(r["debit"] for r in records), 2)
    total_k = round(sum(r["credit"] for r in records), 2)
    closing = next((r["balance"] for r in reversed(records) if r["balance"] is not None), None)
    checks = []
    if opening is not None and closing is not None:
        hitung = round(opening + total_k - total_d, 2)
        checks.append({"label": "Saldo akhir = saldo awal + kredit - debit", "statement": closing, "extracted": hitung, "ok": abs(hitung - closing) <= max(TOL, 0.01 * max(1, len(records)) / 100)})
    n_high = sum(1 for r in records if r["confidence"] == "high")
    n_med = sum(1 for r in records if r["confidence"] == "medium")
    n_low = sum(1 for r in records if r["confidence"] == "low")
    if bad or (checks and not all(c["ok"] for c in checks)):
        status = "failed"
    elif cek and n_med == 0 and n_low == 0:
        status = "verified"
    elif cek:
        status = "partial"
    else:
        status = "unverified"
    return {
        "status": status, "checks": checks, "balance_checkpoints": cek, "balance_mismatches": bad,
        "rows_high": n_high, "rows_medium": n_med, "rows_low": n_low,
        "verified_pct": round(100.0 * n_high / len(records), 1) if records else 0.0,
        "opening_balance": opening, "closing_balance": closing, "total_debit": total_d, "total_credit": total_k,
    }


def _blob(sheets: List[Dict[str, Any]], batas_baris: int = 40) -> str:
    bagian = []
    for sh in sheets:
        if sh.get("teks"):
            bagian.append(sh["teks"][:6000])
        else:
            for r in sh["rows"][:batas_baris]:
                bagian.append(" ".join(_teks(c) for c in r if _teks(c)))
    return "\n".join(bagian)


def deteksi_nomor_rekening(blob: str, filename: str) -> Optional[str]:
    m = re.search(r"(?i)(?:no\.?\s*(?:rek(?:ening)?|acc(?:ount)?)|nomor\s*rek(?:ening)?|account\s*(?:no|number)|rekening)\s*[:\-.]?\s*([0-9][0-9 .\-]{5,28}[0-9])", blob)
    if m:
        return re.sub(r"\D", "", m.group(1))
    m = re.search(r"(?<!\d)(\d{9,20})(?!\d)", filename)
    return m.group(1) if m else None


def deteksi_nama_bank(blob: str, filename: str) -> Optional[str]:
    atas = (filename + " " + blob[:1500]).upper()
    for b in BANK_DIKENAL:
        if re.search(rf"(?<![A-Z]){re.escape(b)}(?![A-Z])", atas):
            return b.title() if b not in ("BCA", "BRI", "BNI", "BTN", "OCBC", "BSI", "UOB", "HSBC", "BTPN", "CIMB") else b
    return None


def ekstrak_rekening_koran(
    content: bytes, filename: str, rules_user: Optional[Dict[str, Any]] = None,
    kandidat_template: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    tipe, sheets = baca_sheets(content, filename)
    blob = _blob(sheets)
    year = _year_hint(blob, filename)
    kandidat_template = kandidat_template or []

    def proses(grup: List[Dict[str, Any]], metode: str) -> Optional[Dict[str, Any]]:
        sh0 = h0 = None
        for sh in grup:
            h = _cari_header(sh["rows"], kandidat=kandidat_template)
            if h:
                sh0, h0 = sh, h
                break
        if sh0 is None:
            return None
        hi, hl, header_cells, roles_auto = h0
        template = cocokkan_template(header_cells, kandidat_template)
        rules = rules_user or (template or {}).get("rules") or {}
        roles = (_peran_dari_rules(header_cells, rules) if rules.get("columns") else None) or roles_auto

        sampel_angka = []
        for r in sh0["rows"][hi + hl: hi + hl + 80]:
            for j, role in enumerate(roles):
                if role in ("debit", "credit", "amount", "balance") and j < len(r) and isinstance(r[j], str):
                    sampel_angka.append(r[j])
        fmt = rules.get("number_format") or deteksi_format_angka(sampel_angka)
        order = rules.get("date_order") or "dmy"
        default_dir = rules.get("default_direction")

        semua: List[Dict[str, Any]] = []
        opening: Optional[float] = None
        tanpa_arah = 0
        for sh in grup:
            h = _cari_header(sh["rows"], kandidat=kandidat_template)
            if not h:
                continue
            hi2, hl2, cells2, roles2 = h
            if sh is sh0:
                r_pakai = roles
            else:
                r_pakai = (_peran_dari_rules(cells2, rules) if rules.get("columns") else None) or roles2
            hasil = _ekstrak_sheet(sh["rows"], sh["pages"], hi2, hl2, r_pakai, fmt, order, default_dir, year)
            if opening is None and hasil["opening"] is not None and not semua:
                opening = hasil["opening"]
            semua.extend(hasil["records"])
            tanpa_arah += hasil["tanpa_arah"]
            if tipe == "PDF":
                break
        return {
            "metode": metode, "template": template, "rules": rules, "roles": roles, "roles_auto": roles_auto,
            "header_cells": header_cells, "hi": hi, "hl": hl, "fmt": fmt, "order": order, "default_dir": default_dir,
            "records": semua, "opening": opening, "tanpa_arah": tanpa_arah,
        }

    tabel = [sh for sh in sheets if not sh.get("fallback")]
    teks = [sh for sh in sheets if sh.get("fallback")]
    hasil = None
    for grup, metode in ((tabel, {"Excel": "table", "CSV": "table"}.get(tipe, "grid")), (teks, "words")):
        if not grup:
            continue
        kandidat = proses(grup, metode)
        if kandidat is None:
            continue
        if hasil is None or len(kandidat["records"]) > len(hasil["records"]):
            hasil = kandidat
        if hasil["records"]:
            break  # tabel/grid sudah menghasilkan baris -> tidak perlu fallback teks

    if hasil is None:
        return {
            "recognized": False, "file_type": tipe, "method": "legacy", "blob": blob, "filename": filename,
            "warnings": ["Header kolom (tanggal + debit/kredit/jumlah) tidak ditemukan -- pilih peran kolom manual atau pakai file lain."],
            "records": [], "columns": [], "header_cells": [], "header_row_index": None,
            "number_format": "en", "date_order": "dmy", "default_direction": None, "amount_mode": "split",
            "template": None, "validation": validasi_dan_confidence([], None),
            "account_number": deteksi_nomor_rekening(blob, filename), "bank_name": deteksi_nama_bank(blob, filename),
            "period_start": None, "period_end": None,
        }

    warnings: List[str] = []
    semua, roles, roles_auto, header_cells = hasil["records"], hasil["roles"], hasil["roles_auto"], hasil["header_cells"]
    if hasil["tanpa_arah"]:
        warnings.append(f"{hasil['tanpa_arah']} baris tanpa petunjuk arah debit/kredit -- diasumsikan kredit, mohon dicek.")
    validasi = validasi_dan_confidence(semua, hasil["opening"])
    if not semua:
        warnings.append("Header ditemukan tapi tidak ada baris transaksi yang terbaca (cek format tanggal/angka).")
    elif validasi["status"] == "unverified":
        warnings.append("Tidak ada kolom saldo -- baris tidak bisa diverifikasi lewat saldo berjalan.")
    if hasil["metode"] == "words":
        warnings.append("Tabel PDF tidak terdeteksi; dibaca dari teks per baris (akurasi lebih rendah) -- cek hasil sebelum import.")

    tgls = [r["date"] for r in semua]
    cols = [{"header": c, "role": (roles[j] or "ignore") if j < len(roles) else "ignore", "detected_role": roles_auto[j] if j < len(roles_auto) else None}
            for j, c in enumerate(header_cells) if c]
    ada_d = any(c["role"] == "debit" for c in cols)
    ada_k = any(c["role"] == "credit" for c in cols)
    return {
        "recognized": True, "file_type": tipe, "method": hasil["metode"], "blob": blob, "filename": filename,
        "warnings": warnings, "records": semua, "columns": cols, "header_cells": header_cells,
        "header_row_index": hasil["hi"] + 1, "header_len": hasil["hl"],
        "number_format": hasil["fmt"], "date_order": hasil["order"], "default_direction": hasil["default_dir"],
        "amount_mode": "split" if (ada_d and ada_k) else "single",
        "template": hasil["template"], "validation": validasi,
        "account_number": deteksi_nomor_rekening(blob, filename), "bank_name": deteksi_nama_bank(blob, filename),
        "period_start": min(tgls) if tgls else None, "period_end": max(tgls) if tgls else None,
    }


def hash_baris(bank_coa_id: str, r: Dict[str, Any], urutan: int) -> str:
    bal = "" if r["balance"] is None else "%.2f" % r["balance"]
    isi = "%s|%s|%s|%.2f|%.2f|%s|%d" % (bank_coa_id, r["date"].isoformat(), (r["description"] or "").strip().lower(), r["debit"], r["credit"], bal, urutan)
    return hashlib.sha256(isi.encode("utf-8")).hexdigest()


def beri_hash(bank_coa_id: str, records: List[Dict[str, Any]]) -> None:
    """Isi r['row_hash'] -- kemunculan ke-n dari baris identik dalam 1 file diberi urutan, jadi 2 baris kembar yang sah tidak saling menimpa."""
    hitung: Dict[str, int] = {}
    for r in records:
        if r["date"] is None:
            r["row_hash"] = None
            continue
        dasar = f"{r['date']}|{(r['description'] or '').strip().lower()}|{r['debit']}|{r['credit']}|{r['balance']}"
        hitung[dasar] = hitung.get(dasar, 0) + 1
        r["row_hash"] = hash_baris(bank_coa_id, r, hitung[dasar])