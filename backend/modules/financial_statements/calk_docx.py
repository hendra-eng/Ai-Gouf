"""
modules/financial_statements/calk_docx.py
===========================================
Export CALK (Notes to Financial Statements) framework baru ke Word (.docx),
dari data yang SAMA dengan halaman CALK (statements_v1._susun_calk ->
mapped.susun_calk), jadi angka di Word = angka di layar = GL.

Pola styling mengikuti skills/docx-writing/SKILL.md: font di-set eksplisit
per run (_set_font), semua sel tabel lewat satu helper (_cell_text) dengan
garis atas/bawah untuk baris total. Dokumen dibuat di memori (tidak ditulis
ke disk), nama file download disanitasi di statements_v1.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any, Dict, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"
ABU = RGBColor(0x6B, 0x72, 0x80)


def _set_font(run, bold: bool = False, italic: bool = False, size: float = 10, color: Optional[RGBColor] = None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)


def _border(cell, sisi: str, gaya: Optional[str]):
    if not gaya:
        return
    tcpr = cell._element.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    el = OxmlElement(f"w:{sisi}")
    el.set(qn("w:val"), gaya)
    el.set(qn("w:sz"), "6")
    el.set(qn("w:color"), "000000")
    borders.append(el)


def _cell_text(cell, text: str, bold: bool = False, italic: bool = False, align: str = "left", size: float = 9.5,
               border_top: Optional[str] = None, border_bottom: Optional[str] = None, color: Optional[RGBColor] = None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    _set_font(p.add_run(text), bold=bold, italic=italic, size=size, color=color)
    _border(cell, "top", border_top)
    _border(cell, "bottom", border_bottom)


def _paragraf(doc, text: str, bold: bool = False, italic: bool = False, size: float = 10, color: Optional[RGBColor] = None,
              align=WD_ALIGN_PARAGRAPH.LEFT, after: float = 6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    _set_font(p.add_run(text), bold=bold, italic=italic, size=size, color=color)
    return p


def _angka(v: Optional[float]) -> str:
    if v is None:
        return "-"
    if abs(v) < 0.005:
        return "-"
    s = f"{abs(v):,.0f}".replace(",", ".")
    return f"({s})" if v < 0 else s


def _tgl(iso: Optional[str]) -> str:
    if not iso:
        return ""
    return date.fromisoformat(iso).strftime("%d %B %Y")


def buat_docx_calk(data: Dict[str, Any]) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)

    perusahaan = (data.get("client") or {}).get("company_name") or "Company"
    pemb = data.get("compare_period")
    _paragraf(doc, perusahaan.upper(), bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _paragraf(doc, "NOTES TO THE FINANCIAL STATEMENTS", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _paragraf(doc, f"For the period ended {_tgl(data['as_of'])}", italic=True, size=10, color=ABU,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _paragraf(doc, "(Expressed in Indonesian Rupiah, unless otherwise stated)", italic=True, size=9, color=ABU,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    for note in data.get("notes", []):
        _paragraf(doc, f"{note['no']}. {note['title'].upper()}", bold=True, size=11, after=4)
        if note.get("narrative"):
            for alinea in note["narrative"].split("\n"):
                if alinea.strip():
                    _paragraf(doc, alinea.strip(), size=10, after=4)
        for g in note.get("groups", []):
            tabel = doc.add_table(rows=1, cols=3 if pemb else 2)
            tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
            kepala = tabel.rows[0].cells
            _cell_text(kepala[0], g["label"], bold=True, border_bottom="single")
            _cell_text(kepala[1], _tgl(data["as_of"]), bold=True, align="right", border_bottom="single")
            if pemb:
                _cell_text(kepala[2], _tgl(pemb["end"]), bold=True, align="right", border_bottom="single")
            for r in g["rows"]:
                sel = tabel.add_row().cells
                nama = f"{r['name']} ({r['code']})"
                if r.get("override"):
                    nama += " *"
                _cell_text(sel[0], nama)
                _cell_text(sel[1], _angka(r["amount"]), align="right")
                if pemb:
                    _cell_text(sel[2], _angka(r.get("compare_amount")), align="right")
            sel = tabel.add_row().cells
            _cell_text(sel[0], "Total", bold=True, border_top="single", border_bottom="double")
            _cell_text(sel[1], _angka(g["total"]), bold=True, align="right", border_top="single", border_bottom="double")
            if pemb:
                _cell_text(sel[2], _angka(g.get("compare_total")), bold=True, align="right", border_top="single", border_bottom="double")
            for row in tabel.rows:
                row.cells[0].width = Cm(9.5)
            if any(r.get("override") for r in g["rows"]):
                _paragraf(doc, "* Amount overridden from the system-calculated value (see audit trail).", italic=True, size=8, color=ABU, after=2)
            _paragraf(doc, "", after=4)
        _paragraf(doc, "", after=6)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
