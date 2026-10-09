"""
migrations/30-create_management_coa_industry_templates.py
===========================================================
Template COA default per industri (KBLI 2020 kategori A..U) untuk onboarding
client baru. Menambahkan 2 tabel BARU + mengisinya dari
dataset/COA/COA_Industry.xlsx:

    management_coa_industry_templates          1 baris per industri, dari sheet
        "COA_JENIS INDUSTRY": KBLI CATEGORY, INDUSTRY (INDONESIA),
        INDUSTRY (ENGLISH) -> pilihan Industry di form client, TEMPLATE SHEET,
        jumlah akun, catatan.
    management_coa_industry_template_accounts  akun default tiap template, dari
        sheet yang disebut TEMPLATE SHEET (mis. "COA A AGRI"); kolom sama
        dengan management_client_coa. normal_balance diturunkan dari
        klasifikasi (generate_seed_coa.normal_balance, akun kontra dibalik).

Saat client baru dibuat dengan industry = salah satu INDUSTRY (ENGLISH),
akun template disalin ke management_client_coa client tsb
(db_client.salin_coa_template_industri, dipanggil create_management_client).

AMAN DIPANGGIL BERKALI-KALI (idempoten) -- data template di-upsert dari
Excel (industri by KBLI CATEGORY, akun by template + ACC NO); akun template
yang sudah tidak ada di Excel dihapus dari TEMPLATE (COA client yang sudah
terbentuk tidak disentuh). Jalankan ulang setelah Excel direvisi.

Cara pakai:
    cd backend
    venv\\Scripts\\python migrations\\30-create_management_coa_industry_templates.py
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    from dotenv import load_dotenv
    load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent / ".env")
except ImportError:
    pass

import openpyxl
from sqlalchemy import inspect, text
from sqlalchemy.exc import OperationalError, ProgrammingError

from _history import catat_history
from db_client import engine
from generate_seed_coa import HEADER_WAJIB, KOLOM_DB, normal_balance

FILE_EXCEL = Path(__file__).resolve().parent.parent.parent / "dataset" / "COA" / "COA_Industry.xlsx"
SHEET_INDUSTRI = "COA_JENIS INDUSTRY"
TABEL_TEMPLATE = "management_coa_industry_templates"
TABEL_AKUN = "management_coa_industry_template_accounts"

# header sheet industri -> kolom DB
KOLOM_INDUSTRI = {
    "KBLI CATEGORY": "kbli_category",
    "INDUSTRY (INDONESIA)": "industry_name_id",
    "INDUSTRY (ENGLISH)": "industry_name_en",
    "TEMPLATE SHEET": "template_sheet",
    "UNIVERSAL ACCOUNTS": "universal_accounts",
    "INDUSTRY-SPECIFIC ACCOUNTS": "industry_specific_accounts",
    "TOTAL TEMPLATE ACCOUNTS": "total_template_accounts",
    "RECOMMENDED USE": "recommended_use",
    "ACCOUNTING FRAMEWORK NOTE": "framework_note",
    "OFFICIAL SOURCE": "official_source",
}


def _teks(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def baca_excel():
    """[(data_industri, [akun, ...]), ...] -- error kalau struktur Excel tidak sesuai."""
    wb = openpyxl.load_workbook(FILE_EXCEL, read_only=True, data_only=True)
    rows = list(wb[SHEET_INDUSTRI].iter_rows(values_only=True))
    header = [_teks(h) for h in rows[0]]
    kurang = [h for h in KOLOM_INDUSTRI if h not in header]
    if kurang:
        raise ValueError(f"Sheet '{SHEET_INDUSTRI}' tidak punya kolom: {kurang}")
    idx = {h: header.index(h) for h in KOLOM_INDUSTRI}

    hasil = []
    for r in rows[1:]:
        if not _teks(r[idx["KBLI CATEGORY"]]):
            continue
        ind = {kolom: r[idx[h]] for h, kolom in KOLOM_INDUSTRI.items()}
        for k in ("kbli_category", "industry_name_id", "industry_name_en", "template_sheet", "recommended_use", "framework_note", "official_source"):
            ind[k] = _teks(ind[k])
        for k in ("universal_accounts", "industry_specific_accounts", "total_template_accounts"):
            ind[k] = int(ind[k]) if ind[k] not in (None, "") else None
        if ind["template_sheet"] not in wb.sheetnames:
            raise ValueError(f"TEMPLATE SHEET '{ind['template_sheet']}' ({ind['industry_name_en']}) tidak ada di workbook.")

        arows = list(wb[ind["template_sheet"]].iter_rows(values_only=True))
        aheader = [_teks(h) for h in arows[0]][:len(HEADER_WAJIB)]
        if aheader != HEADER_WAJIB:
            raise ValueError(f"Header sheet '{ind['template_sheet']}' tidak sesuai: {aheader}")
        akun, dilihat = [], set()
        for no, ar in enumerate(arows[1:], start=1):
            nilai = {k: _teks(v) for k, v in zip(KOLOM_DB, ar[:len(KOLOM_DB)])}
            if not nilai["acc_no"] or not nilai["account_name"]:
                continue
            if nilai["acc_no"] in dilihat:
                raise ValueError(f"ACC NO duplikat '{nilai['acc_no']}' di sheet '{ind['template_sheet']}'.")
            dilihat.add(nilai["acc_no"])
            nilai["account_classification"] = (nilai["account_classification"] or "").upper()
            nilai["normal_balance"] = normal_balance(nilai["account_classification"], (nilai["account_sub"] or "").upper())
            nilai["sort_order"] = no
            akun.append(nilai)
        hasil.append((ind, akun))
    return hasil


def _buat_tabel(connection, nama, orm_class) -> bool:
    if nama in inspect(connection).get_table_names():
        print(f"⏭️  Tabel '{nama}' sudah ada, skip.")
        return True
    try:
        from db_client import Base
        Base.metadata.create_all(bind=connection.engine, tables=[orm_class.__table__])
        connection.commit()
        print(f"✅ Tabel '{nama}' berhasil dibuat.")
        return True
    except (OperationalError, ProgrammingError) as e:
        connection.rollback()
        print(f"❌ Gagal buat tabel '{nama}': {e}")
        return False


def _muat_data(connection, data) -> bool:
    kolom_ind = list(KOLOM_INDUSTRI.values())
    kolom_akun = KOLOM_DB + ["normal_balance", "sort_order"]
    try:
        total_akun = 0
        for ind, akun in data:
            tpl_id = connection.execute(text(
                f"INSERT INTO {TABEL_TEMPLATE} ({', '.join(kolom_ind)}) "
                f"VALUES ({', '.join(':' + k for k in kolom_ind)}) "
                "ON CONFLICT (kbli_category) DO UPDATE SET "
                + ", ".join(f"{k} = EXCLUDED.{k}" for k in kolom_ind if k != "kbli_category")
                + ", edited_at = now() RETURNING id"
            ), ind).scalar()
            for a in akun:
                connection.execute(text(
                    f"INSERT INTO {TABEL_AKUN} (template_id, {', '.join(kolom_akun)}) "
                    f"VALUES (:template_id, {', '.join(':' + k for k in kolom_akun)}) "
                    "ON CONFLICT (template_id, acc_no) DO UPDATE SET "
                    + ", ".join(f"{k} = EXCLUDED.{k}" for k in kolom_akun if k != "acc_no")
                ), {**a, "template_id": tpl_id})
            connection.execute(text(
                f"DELETE FROM {TABEL_AKUN} WHERE template_id = :t AND NOT (acc_no = ANY(:nos))"
            ), {"t": tpl_id, "nos": [a["acc_no"] for a in akun]})
            total_akun += len(akun)
            print(f"   {ind['kbli_category']}  {ind['industry_name_en']:<42} {ind['template_sheet']:<24} {len(akun):>4} akun")
        connection.commit()
        print(f"✅ {len(data)} industri, {total_akun} akun template dimuat.")
        return True
    except (OperationalError, ProgrammingError) as e:
        connection.rollback()
        print(f"❌ Gagal memuat data template: {e}")
        return False


def main() -> int:
    print("=" * 60)
    print("🔄 MIGRATION: template COA per industri (COA_Industry.xlsx)")
    print("=" * 60)

    from db_client import ManagementCoaIndustryTemplate, ManagementCoaIndustryTemplateAccount

    try:
        data = baca_excel()
    except (FileNotFoundError, KeyError, ValueError) as e:
        print(f"❌ Gagal membaca {FILE_EXCEL.name}: {e}")
        return 1

    hasil = {}
    with engine.connect() as connection:
        hasil[f"buat tabel {TABEL_TEMPLATE}"] = _buat_tabel(connection, TABEL_TEMPLATE, ManagementCoaIndustryTemplate)
        hasil[f"buat tabel {TABEL_AKUN}"] = _buat_tabel(connection, TABEL_AKUN, ManagementCoaIndustryTemplateAccount)
        if all(hasil.values()):
            hasil["muat data template dari Excel"] = _muat_data(connection, data)

    print()
    print("=" * 60)
    print("📋 RINGKASAN MIGRATION")
    print("=" * 60)
    for k, v in hasil.items():
        print(f"{k:<70}: {'✅' if v else '❌'}")
    print("=" * 60)

    if all(hasil.values()):
        print("✅ SEMUA migration berhasil!")
        catat_history(Path(__file__).name)
        return 0
    print("⚠️  Ada migration yang gagal. Periksa error di atas.")
    return 1


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    sys.exit(main())
